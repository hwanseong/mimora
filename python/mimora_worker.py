from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import sys
import urllib.error
import urllib.request
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}
SUPPORTED_SCHEDULE_EXTENSIONS = {".xlsx", ".xlsm"}
RAG_ID_PREFIX = "RAG"
DEFAULT_EMBEDDING_PROVIDER = "local"
DEFAULT_LOCAL_EMBEDDING_MODEL = "bge-m3"
DEFAULT_OPENAI_EMBEDDING_MODEL = "text-embedding-3-small"
DEFAULT_EMBEDDING_MODEL = DEFAULT_LOCAL_EMBEDDING_MODEL
DEFAULT_OLLAMA_BASE_URL = "http://127.0.0.1:11434"
DEFAULT_SIMILARITY_THRESHOLD = 0.22
DEFAULT_WORKSPACE_SCORE_BONUS = 0.03
DEFAULT_MAX_CHUNKS_PER_DOCUMENT = 2
DEFAULT_SEARCH_CANDIDATE_COUNT = 50
DEFAULT_RRF_K = 60
DEFAULT_DENSE_RRF_WEIGHT = 1.0
DEFAULT_LEXICAL_RRF_WEIGHT = 1.25
DEFAULT_DENSE_ONLY_THRESHOLD = 0.55
TARGET_CHUNK_CHARS = 3600
CHUNK_OVERLAP_CHARS = 500
SCHEMA = """
CREATE TABLE IF NOT EXISTS rag_documents (
  rag_document_id TEXT PRIMARY KEY,
  original_filename TEXT NOT NULL,
  managed_file_path TEXT NOT NULL,
  file_type TEXT NOT NULL,
  file_size INTEGER NOT NULL,
  file_hash TEXT NOT NULL,
  security TEXT NOT NULL,
  status TEXT NOT NULL,
  created_at TEXT NOT NULL,
  indexed_at TEXT,
  embedding_provider TEXT,
  embedding_model TEXT,
  embedding_dimension INTEGER,
  chunk_count INTEGER NOT NULL DEFAULT 0
);
CREATE UNIQUE INDEX IF NOT EXISTS idx_rag_documents_file_hash
ON rag_documents(file_hash);
CREATE TABLE IF NOT EXISTS rag_document_workspaces (
  rag_document_id TEXT NOT NULL,
  workspace_id TEXT NOT NULL,
  PRIMARY KEY (rag_document_id, workspace_id),
  FOREIGN KEY (rag_document_id)
    REFERENCES rag_documents(rag_document_id)
    ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS rag_chunks (
  chunk_id TEXT PRIMARY KEY,
  rag_document_id TEXT NOT NULL,
  chunk_index INTEGER NOT NULL,
  text TEXT NOT NULL,
  heading TEXT,
  page INTEGER,
  section_index INTEGER,
  faiss_vector_id INTEGER NOT NULL UNIQUE,
  FOREIGN KEY (rag_document_id)
    REFERENCES rag_documents(rag_document_id)
    ON DELETE CASCADE
);
CREATE INDEX IF NOT EXISTS idx_rag_chunks_document_id
ON rag_chunks(rag_document_id);
CREATE INDEX IF NOT EXISTS idx_rag_chunks_faiss_vector_id
ON rag_chunks(faiss_vector_id);
CREATE TABLE IF NOT EXISTS rag_index_profile (
  index_name TEXT PRIMARY KEY,
  embedding_provider TEXT NOT NULL,
  embedding_model TEXT NOT NULL,
  embedding_dimension INTEGER NOT NULL,
  updated_at TEXT NOT NULL
);
"""


class WorkerError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedSection:
    section_index: int
    heading: str | None
    page: int | None
    text: str


@dataclass(frozen=True)
class RagChunk:
    chunk_id: str
    rag_document_id: str
    chunk_index: int
    text: str
    heading: str | None
    page: int | None
    section_index: int | None
    faiss_vector_id: int


@dataclass(frozen=True)
class ManagedCopyPaths:
    documents_root: Path
    document_directory: Path
    managed_file_path: Path


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def read_input() -> dict[str, Any]:
    raw_input = sys.stdin.buffer.read().decode("utf-8").strip()
    if not raw_input:
        return {}

    try:
        value = json.loads(raw_input)
    except json.JSONDecodeError as error:
        raise WorkerError("invalid_json", f"Invalid JSON input: {error}") from error

    if not isinstance(value, dict):
        raise WorkerError("invalid_input", "Worker input must be a JSON object.")

    return value


def write_success(data: Any) -> None:
    payload = json.dumps({"ok": True, "data": data}, ensure_ascii=False)
    sys.stdout.buffer.write(payload.encode("utf-8"))


def write_error(error: WorkerError) -> None:
    payload = json.dumps(
        {
            "ok": False,
            "error": {"code": error.code, "message": error.message},
        },
        ensure_ascii=False,
    )
    sys.stdout.buffer.write(payload.encode("utf-8"))


def get_storage_root(input_data: dict[str, Any]) -> Path:
    storage_root = input_data.get("storage_root")
    if not isinstance(storage_root, str) or not storage_root.strip():
        raise WorkerError("invalid_input", "storage_root is required.")
    return Path(storage_root)


def get_database_path(storage_root: Path) -> Path:
    return storage_root / "metadata.sqlite"


def get_index_path(storage_root: Path) -> Path:
    return storage_root / "index" / "internal.faiss"


def get_test_index_path(storage_root: Path) -> Path:
    return storage_root / "index" / "internal-test-vectors.json"


def get_rag_documents_root(storage_root: Path) -> Path:
    return storage_root / "documents"


def connect_database(storage_root: Path) -> sqlite3.Connection:
    storage_root.mkdir(parents=True, exist_ok=True)
    get_rag_documents_root(storage_root).mkdir(parents=True, exist_ok=True)
    (storage_root / "index").mkdir(parents=True, exist_ok=True)
    (storage_root / "temp").mkdir(parents=True, exist_ok=True)

    connection = sqlite3.connect(get_database_path(storage_root))
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.executescript(SCHEMA)
    ensure_schema_migrations(connection)
    ensure_fts_schema(connection)
    connection.commit()
    return connection


def ensure_schema_migrations(connection: sqlite3.Connection) -> None:
    columns = {
        row["name"]
        for row in connection.execute("PRAGMA table_info(rag_documents)").fetchall()
    }
    if "embedding_dimension" not in columns:
        connection.execute("ALTER TABLE rag_documents ADD COLUMN embedding_dimension INTEGER")


def ensure_fts_schema(connection: sqlite3.Connection) -> None:
    existing_columns = {
        row["name"]
        for row in connection.execute("PRAGMA table_info(rag_chunks_fts)").fetchall()
    }
    if existing_columns and "original_filename" not in existing_columns:
        connection.execute("DROP TABLE IF EXISTS rag_chunks_fts")

    try:
        connection.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS rag_chunks_fts
            USING fts5(
              chunk_id UNINDEXED,
              rag_document_id UNINDEXED,
              original_filename,
              heading,
              text,
              tokenize='trigram'
            )
            """
        )
    except sqlite3.OperationalError:
        try:
            connection.execute(
                """
                CREATE VIRTUAL TABLE IF NOT EXISTS rag_chunks_fts
                USING fts5(
                  chunk_id UNINDEXED,
                  rag_document_id UNINDEXED,
                  original_filename,
                  heading,
                  text
                )
                """
            )
        except sqlite3.OperationalError:
            return
    rebuild_fts_index_if_needed(connection)


def normalize_fts_text(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"[_\-.]+", " ", value)


def rebuild_fts_index_if_needed(connection: sqlite3.Connection) -> None:
    try:
        chunk_count = connection.execute("SELECT COUNT(*) AS count FROM rag_chunks").fetchone()["count"]
        fts_count = connection.execute("SELECT COUNT(*) AS count FROM rag_chunks_fts").fetchone()["count"]
    except sqlite3.OperationalError:
        return

    if int(chunk_count) == int(fts_count):
        return

    connection.execute("DELETE FROM rag_chunks_fts")
    rows = connection.execute(
        """
        SELECT rag_chunks.chunk_id,
               rag_chunks.rag_document_id,
               rag_documents.original_filename,
               rag_chunks.heading,
               rag_chunks.text
        FROM rag_chunks
        INNER JOIN rag_documents
          ON rag_documents.rag_document_id = rag_chunks.rag_document_id
        ORDER BY rag_chunks.faiss_vector_id
        """
    ).fetchall()
    connection.executemany(
        """
        INSERT INTO rag_chunks_fts (
          chunk_id, rag_document_id, original_filename, heading, text
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        [
            (
                row["chunk_id"],
                row["rag_document_id"],
                normalize_fts_text(row["original_filename"]),
                normalize_fts_text(row["heading"]),
                row["text"],
            )
            for row in rows
        ],
    )


def sha256_file(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_source_path(input_data: dict[str, Any]) -> Path:
    source_path = input_data.get("source_path")
    if not isinstance(source_path, str) or not source_path.strip():
        raise WorkerError("source_file_not_found", "Source file path is required.")

    path_value = Path(source_path)
    if not path_value.is_absolute():
        raise WorkerError("source_file_not_found", "Source file path must be absolute.")
    if not path_value.exists() or not path_value.is_file():
        raise WorkerError("source_file_not_found", "Source file was not found.")
    if path_value.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise WorkerError("unsupported_file_type", f"Unsupported file type: {path_value.suffix.lower()}")
    return path_value


def validate_security(input_data: dict[str, Any]) -> str:
    security = input_data.get("security")
    if security not in {"internal", "sensitive", "personal", "private"}:
        raise WorkerError("invalid_security", "security must be internal, sensitive, personal, or private.")
    return security


def validate_workspace_ids(input_data: dict[str, Any]) -> list[str]:
    workspace_ids = input_data.get("workspace_ids", [])
    if workspace_ids is None:
        return []
    if not isinstance(workspace_ids, list):
        raise WorkerError("invalid_workspace_ids", "workspace_ids must be a list.")

    result: list[str] = []
    seen: set[str] = set()
    for workspace_id in workspace_ids:
        if not isinstance(workspace_id, str):
            raise WorkerError("invalid_workspace_ids", "workspace_ids must contain strings.")
        normalized = workspace_id.strip()
        if normalized and normalized not in seen:
            seen.add(normalized)
            result.append(normalized)
    return result


def sanitize_filename(filename: str) -> str:
    return Path(filename).name.strip() or "document"


def row_to_document(connection: sqlite3.Connection, row: sqlite3.Row) -> dict[str, Any]:
    workspace_rows = connection.execute(
        """
        SELECT workspace_id
        FROM rag_document_workspaces
        WHERE rag_document_id = ?
        ORDER BY workspace_id
        """,
        (row["rag_document_id"],),
    ).fetchall()
    return {
        "ragDocumentId": row["rag_document_id"],
        "originalFilename": row["original_filename"],
        "managedFilePath": row["managed_file_path"],
        "fileType": row["file_type"],
        "fileSize": row["file_size"],
        "fileHash": row["file_hash"],
        "security": row["security"],
        "status": row["status"],
        "createdAt": row["created_at"],
        "indexedAt": row["indexed_at"],
        "embeddingProvider": row["embedding_provider"],
        "embeddingModel": row["embedding_model"],
        "embeddingDimension": row["embedding_dimension"],
        "chunkCount": row["chunk_count"],
        "workspaceIds": [workspace_row["workspace_id"] for workspace_row in workspace_rows],
    }


def get_document_by_hash(connection: sqlite3.Connection, file_hash: str) -> dict[str, Any] | None:
    row = connection.execute("SELECT * FROM rag_documents WHERE file_hash = ?", (file_hash,)).fetchone()
    return row_to_document(connection, row) if row else None


def get_document_row(connection: sqlite3.Connection, rag_document_id: str) -> sqlite3.Row:
    row = connection.execute(
        "SELECT * FROM rag_documents WHERE rag_document_id = ?",
        (rag_document_id,),
    ).fetchone()
    if not row:
        raise WorkerError("document_not_found", "RAG document was not found.")
    return row


def canonical_path(path_value: Path, *, must_exist: bool = True) -> Path:
    try:
        return path_value.resolve(strict=must_exist)
    except OSError as error:
        if must_exist:
            raise WorkerError(
                "unsafe_managed_path",
                "RAG 문서 저장 경로가 안전하지 않아 작업할 수 없습니다.",
            ) from error
        return path_value.resolve(strict=False)


def is_path_contained_by(child: Path, parent: Path) -> bool:
    child_value = os.path.normcase(str(child))
    parent_value = os.path.normcase(str(parent))
    try:
        return os.path.commonpath([child_value, parent_value]) == parent_value
    except ValueError:
        return False


def assert_managed_path_contained(child: Path, parent: Path) -> None:
    if not is_path_contained_by(child, parent):
        raise WorkerError(
            "unsafe_managed_path",
            "RAG 문서 저장 경로가 안전하지 않아 작업할 수 없습니다.",
        )


def resolve_managed_copy_paths(
    storage_root: Path,
    row: sqlite3.Row,
) -> ManagedCopyPaths:
    rag_document_id = str(row["rag_document_id"])
    documents_root = get_rag_documents_root(storage_root)
    documents_root.mkdir(parents=True, exist_ok=True)
    documents_root_canonical = canonical_path(documents_root)
    expected_directory = documents_root / rag_document_id
    expected_directory_canonical = canonical_path(
        expected_directory,
        must_exist=expected_directory.exists(),
    )

    assert_managed_path_contained(expected_directory_canonical, documents_root_canonical)

    managed_path_raw = Path(str(row["managed_file_path"]))
    if not managed_path_raw.is_absolute():
        raise WorkerError(
            "unsafe_managed_path",
            "RAG 문서 저장 경로가 안전하지 않아 작업할 수 없습니다.",
        )

    managed_path_canonical = canonical_path(
        managed_path_raw,
        must_exist=managed_path_raw.exists(),
    )
    expected_file_canonical = canonical_path(
        expected_directory / sanitize_filename(str(row["original_filename"])),
        must_exist=False,
    )

    if os.path.normcase(str(managed_path_canonical)) != os.path.normcase(str(expected_file_canonical)):
        raise WorkerError(
            "unsafe_managed_path",
            "RAG 문서 저장 경로가 안전하지 않아 작업할 수 없습니다.",
        )

    if os.path.normcase(str(managed_path_canonical.parent)) != os.path.normcase(str(expected_directory_canonical)):
        raise WorkerError(
            "unsafe_managed_path",
            "RAG 문서 저장 경로가 안전하지 않아 작업할 수 없습니다.",
        )

    assert_managed_path_contained(managed_path_canonical, expected_directory_canonical)

    return ManagedCopyPaths(
        documents_root=documents_root_canonical,
        document_directory=expected_directory_canonical,
        managed_file_path=managed_path_canonical,
    )


def generate_rag_document_id(connection: sqlite3.Connection) -> str:
    year = datetime.now(timezone.utc).year
    prefix = f"{RAG_ID_PREFIX}-{year}-"
    row = connection.execute(
        """
        SELECT rag_document_id
        FROM rag_documents
        WHERE rag_document_id LIKE ?
        ORDER BY rag_document_id DESC
        LIMIT 1
        """,
        (f"{prefix}%",),
    ).fetchone()
    next_number = 1
    if row:
        try:
            next_number = int(str(row["rag_document_id"]).split("-")[-1]) + 1
        except ValueError:
            next_number = 1
    return f"{prefix}{next_number:06d}"


def read_text_with_fallback(path_value: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp949"):
        try:
            return path_value.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise WorkerError("text_decode_failed", "Text file could not be decoded.")


def normalize_text(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n").replace("\r", "\n")).strip()


def strip_yaml_frontmatter(markdown: str) -> str:
    if not markdown.startswith("---"):
        return markdown
    match = re.match(r"^---[ \t]*\n.*?\n---[ \t]*(?:\n|$)", markdown, flags=re.DOTALL)
    return markdown[match.end():] if match else markdown


def sections_from_heading_text(text: str) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        body = normalize_text("\n".join(current_lines))
        if body:
            sections.append(ExtractedSection(len(sections), current_heading, None, body))

    for line in text.splitlines():
        heading_match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading_match:
            flush()
            current_heading = heading_match.group(2).strip()
            current_lines = []
            continue
        current_lines.append(line)
    flush()

    if not sections and normalize_text(text):
        sections.append(ExtractedSection(0, None, None, normalize_text(text)))
    return sections


def extract_markdown(path_value: Path) -> dict[str, Any]:
    body = strip_yaml_frontmatter(read_text_with_fallback(path_value))
    return {"sections": [section.__dict__ for section in sections_from_heading_text(body)]}


def extract_text(path_value: Path) -> dict[str, Any]:
    text = normalize_text(read_text_with_fallback(path_value))
    return {"sections": [ExtractedSection(0, None, None, text).__dict__] if text else []}


def extract_pdf(path_value: Path) -> dict[str, Any]:
    try:
        import pymupdf
    except Exception as error:
        raise WorkerError("dependency_missing", "PyMuPDF is required for PDF extraction.") from error

    sections: list[ExtractedSection] = []
    try:
        with pymupdf.open(str(path_value)) as document:
            for page_index in range(document.page_count):
                text = normalize_text(document.load_page(page_index).get_text("text") or "")
                if text:
                    sections.append(ExtractedSection(len(sections), None, page_index + 1, text))
    except Exception as error:
        raise WorkerError("extraction_failed", "PDF text extraction failed.") from error
    return {"sections": [section.__dict__ for section in sections]}


def docx_paragraphs_with_stdlib(path_value: Path) -> list[tuple[str | None, str]]:
    try:
        with zipfile.ZipFile(path_value) as archive:
            xml_text = archive.read("word/document.xml")
    except Exception as error:
        raise WorkerError("extraction_failed", "DOCX text extraction failed.") from error

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ElementTree.fromstring(xml_text)
    paragraphs: list[tuple[str | None, str]] = []
    for paragraph in root.findall(".//w:p", namespace):
        style_node = paragraph.find("./w:pPr/w:pStyle", namespace)
        style = style_node.attrib.get(f"{{{namespace['w']}}}val") if style_node is not None else None
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            paragraphs.append((style, text))
    return paragraphs


def extract_docx(path_value: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception:
        paragraphs = docx_paragraphs_with_stdlib(path_value)
    else:
        document = Document(str(path_value))
        paragraphs = [
            (paragraph.style.name if paragraph.style else None, paragraph.text.strip())
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                cell_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if cell_text:
                    paragraphs.append((None, cell_text))

    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        body = normalize_text("\n".join(current_lines))
        if body:
            sections.append(ExtractedSection(len(sections), current_heading, None, body))

    for style, text in paragraphs:
        if style and style.lower().replace(" ", "").startswith("heading"):
            flush()
            current_heading = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()

    if not sections:
        body = normalize_text("\n".join(text for _, text in paragraphs))
        if body:
            sections.append(ExtractedSection(0, None, None, body))
    return {"sections": [section.__dict__ for section in sections]}


def extract_document(path_value: Path) -> dict[str, Any]:
    extension = path_value.suffix.lower()
    if extension == ".pdf":
        return extract_pdf(path_value)
    if extension == ".docx":
        return extract_docx(path_value)
    if extension == ".md":
        return extract_markdown(path_value)
    if extension == ".txt":
        return extract_text(path_value)
    raise WorkerError("unsupported_file_type", f"Unsupported file type: {extension}")


def get_chunk_config(input_data: dict[str, Any] | None = None) -> tuple[int, int]:
    input_data = input_data or {}
    raw_chunk_size = input_data.get("chunk_size")
    raw_chunk_overlap = input_data.get("chunk_overlap")
    chunk_size = (
        int(raw_chunk_size)
        if isinstance(raw_chunk_size, int) and 1000 <= raw_chunk_size <= 12000
        else TARGET_CHUNK_CHARS
    )
    chunk_overlap = (
        int(raw_chunk_overlap)
        if isinstance(raw_chunk_overlap, int) and 0 <= raw_chunk_overlap < chunk_size
        else CHUNK_OVERLAP_CHARS
    )
    return chunk_size, chunk_overlap


def get_float_config(
    input_data: dict[str, Any],
    key: str,
    default_value: float,
    minimum: float,
    maximum: float,
) -> float:
    value = input_data.get(key)
    if not isinstance(value, (int, float)):
        return default_value
    return max(minimum, min(maximum, float(value)))


def get_int_config(
    input_data: dict[str, Any],
    key: str,
    default_value: int,
    minimum: int,
    maximum: int,
) -> int:
    value = input_data.get(key)
    if not isinstance(value, int):
        return default_value
    return max(minimum, min(maximum, int(value)))


def split_long_text(text: str, chunk_size: int = TARGET_CHUNK_CHARS, chunk_overlap: int = CHUNK_OVERLAP_CHARS) -> list[str]:
    normalized = normalize_text(text)
    if len(normalized) <= chunk_size:
        return [normalized] if normalized else []

    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n\s*\n", normalized) if paragraph.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
            overlap = current[-chunk_overlap:].strip() if chunk_overlap > 0 else ""
            current = f"{overlap}\n\n{paragraph}".strip() if overlap else paragraph
        else:
            start = 0
            while start < len(paragraph):
                end = min(len(paragraph), start + chunk_size)
                chunks.append(paragraph[start:end].strip())
                if end >= len(paragraph):
                    break
                start = max(end - chunk_overlap, start + 1)
            current = ""
    if current:
        chunks.append(current)
    return chunks


def next_vector_id(connection: sqlite3.Connection) -> int:
    row = connection.execute("SELECT MAX(faiss_vector_id) AS max_id FROM rag_chunks").fetchone()
    value = row["max_id"] if row else None
    return int(value) + 1 if value is not None else 1


def get_chunk_row_count(connection: sqlite3.Connection) -> int:
    row = connection.execute("SELECT COUNT(*) AS count FROM rag_chunks").fetchone()
    return int(row["count"]) if row else 0


def ensure_index_profile(
    connection: sqlite3.Connection,
    provider: str,
    model: str,
    dimension: int,
) -> None:
    row = connection.execute(
        "SELECT * FROM rag_index_profile WHERE index_name = 'internal'",
    ).fetchone()

    if row:
        matches = (
            row["embedding_provider"] == provider
            and row["embedding_model"] == model
            and int(row["embedding_dimension"]) == dimension
        )
        if not matches and get_chunk_row_count(connection) > 0:
            raise WorkerError(
                "embedding_profile_mismatch",
                "Embedding configuration changed. RAG documents need re-embedding.",
            )

    connection.execute(
        """
        INSERT INTO rag_index_profile (
          index_name,
          embedding_provider,
          embedding_model,
          embedding_dimension,
          updated_at
        )
        VALUES ('internal', ?, ?, ?, ?)
        ON CONFLICT(index_name) DO UPDATE SET
          embedding_provider = excluded.embedding_provider,
          embedding_model = excluded.embedding_model,
          embedding_dimension = excluded.embedding_dimension,
          updated_at = excluded.updated_at
        """,
        (provider, model, dimension, now_iso()),
    )


def assert_index_profile(
    connection: sqlite3.Connection,
    provider: str,
    model: str,
    dimension: int,
) -> None:
    row = connection.execute(
        "SELECT * FROM rag_index_profile WHERE index_name = 'internal'",
    ).fetchone()
    if not row:
        return

    if (
        row["embedding_provider"] != provider
        or row["embedding_model"] != model
        or int(row["embedding_dimension"]) != dimension
    ):
        raise WorkerError(
            "embedding_profile_mismatch",
            "Query embedding profile does not match the current RAG index.",
        )


def build_chunks(
    rag_document_id: str,
    extraction: dict[str, Any],
    first_vector_id: int,
    chunk_size: int = TARGET_CHUNK_CHARS,
    chunk_overlap: int = CHUNK_OVERLAP_CHARS,
) -> list[RagChunk]:
    raw_sections = extraction.get("sections")
    if not isinstance(raw_sections, list):
        raise WorkerError("extraction_failed", "Extractor returned an invalid section list.")

    chunks: list[RagChunk] = []
    vector_id = first_vector_id
    for raw_section in raw_sections:
        if not isinstance(raw_section, dict):
            continue
        text = raw_section.get("text")
        if not isinstance(text, str) or not text.strip():
            continue
        heading = raw_section.get("heading")
        page = raw_section.get("page")
        section_index = raw_section.get("section_index")
        for text_part in split_long_text(text, chunk_size, chunk_overlap):
            chunk_index = len(chunks)
            chunks.append(
                RagChunk(
                    chunk_id=f"{rag_document_id}-CH-{chunk_index + 1:06d}",
                    rag_document_id=rag_document_id,
                    chunk_index=chunk_index,
                    text=text_part,
                    heading=heading if isinstance(heading, str) and heading.strip() else None,
                    page=page if isinstance(page, int) else None,
                    section_index=section_index if isinstance(section_index, int) else None,
                    faiss_vector_id=vector_id,
                )
            )
            vector_id += 1
    return chunks


def deterministic_test_embeddings(texts: list[str], dimension: int = 16) -> list[list[float]]:
    vectors: list[list[float]] = []
    for text in texts:
        digest = hashlib.sha256(text.encode("utf-8")).digest()
        vector = [((digest[index % len(digest)] / 255.0) * 2.0) - 1.0 for index in range(dimension)]
        length = math.sqrt(sum(value * value for value in vector)) or 1.0
        vectors.append([value / length for value in vector])
    return vectors


class EmbeddingProvider:
    provider_name: str
    model_name: str
    dimension: int | None

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        raise NotImplementedError

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]


class TestEmbeddingProvider(EmbeddingProvider):
    provider_name = "test"

    def __init__(self, model_name: str):
        self.model_name = model_name or "deterministic-test"
        self.dimension = 16

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return deterministic_test_embeddings(texts, self.dimension or 16)


def post_json(url: str, payload: dict[str, Any], headers: dict[str, str] | None = None) -> dict[str, Any]:
    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", **(headers or {})},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        value = json.loads(response.read().decode("utf-8"))
    if not isinstance(value, dict):
        raise WorkerError("embedding_failed", "Embedding response was invalid.")
    return value


def get_json(url: str) -> dict[str, Any]:
    request = urllib.request.Request(url, method="GET")
    with urllib.request.urlopen(request, timeout=10) as response:
        value = json.loads(response.read().decode("utf-8"))
    if not isinstance(value, dict):
        raise WorkerError("ollama_unavailable", "Ollama response was invalid.")
    return value


def embed_with_openai(texts: list[str], model: str, api_key: str) -> list[list[float]]:
    if not api_key.strip():
        raise WorkerError("openai_api_key_missing", "OpenAI API Key is required for embedding.")

    request = urllib.request.Request(
        "https://api.openai.com/v1/embeddings",
        data=json.dumps({"model": model, "input": texts, "encoding_format": "float"}).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        print(error.read().decode("utf-8", errors="replace"), file=sys.stderr)
        raise WorkerError("embedding_failed", "OpenAI embedding request failed.") from error
    except Exception as error:
        raise WorkerError("embedding_failed", "OpenAI embedding request failed.") from error

    data = payload.get("data")
    if not isinstance(data, list):
        raise WorkerError("embedding_failed", "OpenAI embedding response was invalid.")

    vectors: list[list[float]] = []
    for item in sorted(data, key=lambda value: value.get("index", 0) if isinstance(value, dict) else 0):
        embedding = item.get("embedding") if isinstance(item, dict) else None
        if not isinstance(embedding, list) or not all(isinstance(value, (int, float)) for value in embedding):
            raise WorkerError("embedding_failed", "OpenAI embedding vector was invalid.")
        vectors.append([float(value) for value in embedding])
    if len(vectors) != len(texts):
        raise WorkerError("embedding_failed", "OpenAI embedding count did not match chunk count.")
    return vectors


class OpenAIEmbeddingProvider(EmbeddingProvider):
    provider_name = "openai"

    def __init__(self, model_name: str, api_key: str):
        self.model_name = model_name or DEFAULT_OPENAI_EMBEDDING_MODEL
        self.api_key = api_key
        self.dimension = None

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        vectors = embed_with_openai(texts, self.model_name, self.api_key)
        self.dimension = len(vectors[0]) if vectors else self.dimension
        return vectors


class LocalEmbeddingProvider(EmbeddingProvider):
    provider_name = "local"

    def __init__(self, model_name: str, ollama_base_url: str):
        self.model_name = model_name or DEFAULT_LOCAL_EMBEDDING_MODEL
        self.ollama_base_url = ollama_base_url.rstrip("/") or DEFAULT_OLLAMA_BASE_URL
        self.dimension = None

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        if not self.model_name.strip():
            raise WorkerError("embedding_model_unavailable", "Local embedding model is not configured.")

        try:
            payload = post_json(
                f"{self.ollama_base_url}/api/embed",
                {"model": self.model_name, "input": texts},
            )
        except urllib.error.URLError as error:
            raise WorkerError("ollama_unavailable", "Ollama is not reachable.") from error
        except Exception as error:
            raise WorkerError("local_embedding_failed", "Local embedding failed.") from error

        raw_embeddings = payload.get("embeddings")
        if raw_embeddings is None and isinstance(payload.get("embedding"), list):
            raw_embeddings = [payload.get("embedding")]

        if not isinstance(raw_embeddings, list):
            raise WorkerError("local_embedding_failed", "Ollama embedding response was invalid.")

        vectors: list[list[float]] = []
        for embedding in raw_embeddings:
            if not isinstance(embedding, list) or not all(isinstance(value, (int, float)) for value in embedding):
                raise WorkerError("local_embedding_failed", "Ollama embedding vector was invalid.")
            vectors.append([float(value) for value in embedding])

        if len(vectors) != len(texts):
            raise WorkerError("local_embedding_failed", "Ollama embedding count did not match input count.")

        self.dimension = len(vectors[0]) if vectors else self.dimension
        return vectors


def create_embedding_provider(input_data: dict[str, Any]) -> EmbeddingProvider:
    provider = str(input_data.get("embedding_provider") or DEFAULT_EMBEDDING_PROVIDER).strip().lower()
    model = str(input_data.get("embedding_model") or "").strip()

    if provider == "test":
        return TestEmbeddingProvider(model)
    if provider == "openai":
        return OpenAIEmbeddingProvider(model or DEFAULT_OPENAI_EMBEDDING_MODEL, str(input_data.get("openai_api_key") or ""))
    if provider == "local":
        return LocalEmbeddingProvider(
            model or DEFAULT_LOCAL_EMBEDDING_MODEL,
            str(input_data.get("ollama_base_url") or DEFAULT_OLLAMA_BASE_URL),
        )

    raise WorkerError("embedding_provider_unavailable", "Embedding provider is not available.")


def embed_texts(input_data: dict[str, Any], texts: list[str]) -> tuple[str, str, int | None, list[list[float]]]:
    provider = create_embedding_provider(input_data)
    if not texts:
        return provider.provider_name, provider.model_name, provider.dimension, []
    vectors = provider.embed_documents(texts)
    return provider.provider_name, provider.model_name, provider.dimension, vectors


def import_faiss_modules() -> tuple[Any, Any]:
    try:
        import faiss
        import numpy
    except Exception as error:
        raise WorkerError("dependency_missing", "faiss-cpu and numpy are required for FAISS indexing.") from error
    return faiss, numpy


def l2_normalize(vectors: list[list[float]], numpy: Any) -> Any:
    array = numpy.asarray(vectors, dtype="float32")
    norms = numpy.linalg.norm(array, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return array / norms


def add_faiss_vectors(storage_root: Path, chunks: list[RagChunk], vectors: list[list[float]]) -> None:
    if not chunks:
        return
    faiss, numpy = import_faiss_modules()
    normalized_vectors = l2_normalize(vectors, numpy)
    index_path = get_index_path(storage_root)
    if index_path.exists():
        index = faiss.read_index(str(index_path))
        if index.d != normalized_vectors.shape[1]:
            raise WorkerError("faiss_dimension_mismatch", "Existing FAISS index dimension does not match embedding model.")
    else:
        index = faiss.IndexIDMap2(faiss.IndexFlatIP(normalized_vectors.shape[1]))
    ids = numpy.asarray([chunk.faiss_vector_id for chunk in chunks], dtype="int64")
    index.add_with_ids(normalized_vectors, ids)
    faiss.write_index(index, str(index_path))


def remove_faiss_vectors(storage_root: Path, vector_ids: list[int]) -> None:
    index_path = get_index_path(storage_root)
    if not vector_ids or not index_path.exists():
        return
    faiss, numpy = import_faiss_modules()
    index = faiss.read_index(str(index_path))
    index.remove_ids(numpy.asarray(vector_ids, dtype="int64"))
    faiss.write_index(index, str(index_path))


def load_test_vectors(storage_root: Path) -> dict[str, list[float]]:
    path_value = get_test_index_path(storage_root)
    if not path_value.exists():
        return {}
    try:
        value = json.loads(path_value.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return {str(key): vector for key, vector in value.items() if isinstance(vector, list)}


def save_test_vectors(storage_root: Path, vectors: dict[str, list[float]]) -> None:
    get_test_index_path(storage_root).write_text(json.dumps(vectors), encoding="utf-8")


def add_vectors(storage_root: Path, provider: str, chunks: list[RagChunk], vectors: list[list[float]]) -> None:
    if provider == "test":
        stored = load_test_vectors(storage_root)
        for chunk, vector in zip(chunks, vectors):
            stored[str(chunk.faiss_vector_id)] = vector
        save_test_vectors(storage_root, stored)
        return
    add_faiss_vectors(storage_root, chunks, vectors)


def remove_vectors(storage_root: Path, vector_ids: list[int]) -> None:
    stored = load_test_vectors(storage_root)
    if stored:
        for vector_id in vector_ids:
            stored.pop(str(vector_id), None)
        save_test_vectors(storage_root, stored)
    try:
        remove_faiss_vectors(storage_root, vector_ids)
    except WorkerError as error:
        if error.code != "dependency_missing":
            raise


def delete_document_chunks(connection: sqlite3.Connection, storage_root: Path, rag_document_id: str) -> None:
    rows = connection.execute(
        "SELECT faiss_vector_id FROM rag_chunks WHERE rag_document_id = ?",
        (rag_document_id,),
    ).fetchall()
    remove_vectors(storage_root, [int(row["faiss_vector_id"]) for row in rows])
    try:
        connection.execute("DELETE FROM rag_chunks_fts WHERE rag_document_id = ?", (rag_document_id,))
    except sqlite3.OperationalError:
        pass
    connection.execute("DELETE FROM rag_chunks WHERE rag_document_id = ?", (rag_document_id,))


def insert_chunk_fts_rows(
    connection: sqlite3.Connection,
    chunks: list[RagChunk],
    original_filename: str,
) -> None:
    try:
        connection.executemany(
            """
            INSERT INTO rag_chunks_fts (
              chunk_id, rag_document_id, original_filename, heading, text
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            [
                (
                    chunk.chunk_id,
                    chunk.rag_document_id,
                    normalize_fts_text(original_filename),
                    normalize_fts_text(chunk.heading),
                    chunk.text,
                )
                for chunk in chunks
            ],
        )
    except sqlite3.OperationalError:
        pass


def import_document(input_data: dict[str, Any]) -> dict[str, Any]:
    storage_root = get_storage_root(input_data)
    source_path = validate_source_path(input_data)
    security = validate_security(input_data)
    workspace_ids = validate_workspace_ids(input_data)
    file_hash = sha256_file(source_path)
    file_size = source_path.stat().st_size
    file_type = source_path.suffix.lower()

    with connect_database(storage_root) as connection:
        duplicate = get_document_by_hash(connection, file_hash)
        if duplicate:
            return {"status": "duplicate", "document": duplicate, "duplicateOf": duplicate["ragDocumentId"]}

        rag_document_id = generate_rag_document_id(connection)
        document_directory = storage_root / "documents" / rag_document_id
        managed_file_path = document_directory / sanitize_filename(source_path.name)
        created_at = now_iso()
        try:
            document_directory.mkdir(parents=True, exist_ok=False)
            shutil.copy2(source_path, managed_file_path)
            with connection:
                connection.execute(
                    """
                    INSERT INTO rag_documents (
                      rag_document_id, original_filename, managed_file_path,
                      file_type, file_size, file_hash, security, status,
                      created_at, indexed_at, embedding_provider, embedding_model,
                      chunk_count
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, 'imported', ?, NULL, NULL, NULL, 0)
                    """,
                    (rag_document_id, source_path.name, str(managed_file_path), file_type, file_size, file_hash, security, created_at),
                )
                connection.executemany(
                    """
                    INSERT INTO rag_document_workspaces (rag_document_id, workspace_id)
                    VALUES (?, ?)
                    """,
                    [(rag_document_id, workspace_id) for workspace_id in workspace_ids],
                )
        except Exception as error:
            shutil.rmtree(document_directory, ignore_errors=True)
            if isinstance(error, sqlite3.Error):
                raise WorkerError("database_failed", "RAG document could not be saved.") from error
            raise WorkerError("copy_failed", "RAG document could not be copied.") from error

        return {"status": "imported", "document": row_to_document(connection, get_document_row(connection, rag_document_id))}


def list_documents(input_data: dict[str, Any]) -> list[dict[str, Any]]:
    storage_root = get_storage_root(input_data)
    with connect_database(storage_root) as connection:
        rows = connection.execute("SELECT * FROM rag_documents ORDER BY created_at DESC, rag_document_id DESC").fetchall()
        return [row_to_document(connection, row) for row in rows]


def delete_document(input_data: dict[str, Any]) -> dict[str, Any]:
    storage_root = get_storage_root(input_data)
    rag_document_id = input_data.get("rag_document_id")
    if not isinstance(rag_document_id, str):
        raise WorkerError("invalid_input", "rag_document_id is required.")

    with connect_database(storage_root) as connection:
        row = get_document_row(connection, rag_document_id)
        managed_paths = resolve_managed_copy_paths(storage_root, row)
        try:
            with connection:
                delete_document_chunks(connection, storage_root, rag_document_id)
                connection.execute("DELETE FROM rag_document_workspaces WHERE rag_document_id = ?", (rag_document_id,))
                connection.execute("DELETE FROM rag_documents WHERE rag_document_id = ?", (rag_document_id,))
            shutil.rmtree(managed_paths.document_directory, ignore_errors=True)
        except WorkerError:
            raise
        except sqlite3.Error as error:
            raise WorkerError("database_failed", "RAG document could not be deleted.") from error
        except Exception as error:
            raise WorkerError("delete_failed", "Managed RAG copy could not be deleted.") from error
    return {"ragDocumentId": rag_document_id, "deleted": True}


def replace_document(input_data: dict[str, Any]) -> dict[str, Any]:
    storage_root = get_storage_root(input_data)
    source_path = validate_source_path(input_data)
    rag_document_id = input_data.get("rag_document_id")
    if not isinstance(rag_document_id, str):
        raise WorkerError("invalid_input", "rag_document_id is required.")

    new_hash = sha256_file(source_path)
    with connect_database(storage_root) as connection:
        current_row = get_document_row(connection, rag_document_id)
        if current_row["file_hash"] == new_hash:
            return {"status": "unchanged", "document": row_to_document(connection, current_row)}

        duplicate = get_document_by_hash(connection, new_hash)
        if duplicate and duplicate["ragDocumentId"] != rag_document_id:
            return {"status": "duplicate", "document": row_to_document(connection, current_row), "duplicateOf": duplicate["ragDocumentId"]}

        managed_paths = resolve_managed_copy_paths(storage_root, current_row)
        document_directory = managed_paths.document_directory
        new_managed_file_path = document_directory / sanitize_filename(source_path.name)
        try:
            document_directory.mkdir(parents=True, exist_ok=True)
            for child in document_directory.iterdir():
                if child.is_file():
                    child.unlink()
            shutil.copy2(source_path, new_managed_file_path)
            with connection:
                delete_document_chunks(connection, storage_root, rag_document_id)
                connection.execute(
                    """
                    UPDATE rag_documents
                    SET original_filename = ?, managed_file_path = ?, file_type = ?,
                        file_size = ?, file_hash = ?, status = 'imported',
                        indexed_at = NULL, embedding_provider = NULL,
                        embedding_model = NULL, embedding_dimension = NULL,
                        chunk_count = 0
                    WHERE rag_document_id = ?
                    """,
                    (
                        source_path.name,
                        str(new_managed_file_path),
                        source_path.suffix.lower(),
                        source_path.stat().st_size,
                        new_hash,
                        rag_document_id,
                    ),
                )
        except WorkerError:
            raise
        except sqlite3.Error as error:
            raise WorkerError("database_failed", "RAG document could not be replaced.") from error
        except Exception as error:
            raise WorkerError("replace_failed", "Managed RAG copy could not be replaced.") from error
        return {"status": "replaced", "document": row_to_document(connection, get_document_row(connection, rag_document_id))}


def index_document(input_data: dict[str, Any]) -> dict[str, Any]:
    storage_root = get_storage_root(input_data)
    rag_document_id = input_data.get("rag_document_id")
    if not isinstance(rag_document_id, str):
        raise WorkerError("invalid_input", "rag_document_id is required.")

    with connect_database(storage_root) as connection:
        row = get_document_row(connection, rag_document_id)
        provider = str(input_data.get("embedding_provider") or DEFAULT_EMBEDDING_PROVIDER).strip().lower()
        if row["security"] == "private" and provider != "local":
            raise WorkerError("private_requires_local_embedding", "Private RAG documents require local embedding.")

        managed_file_path = Path(row["managed_file_path"])
        if not managed_file_path.exists() or not managed_file_path.is_file():
            raise WorkerError("managed_copy_not_found", "Managed RAG copy was not found.")

        try:
            with connection:
                connection.execute("UPDATE rag_documents SET status = 'indexing' WHERE rag_document_id = ?", (rag_document_id,))

            extraction = extract_document(managed_file_path)
            chunk_size, chunk_overlap = get_chunk_config(input_data)
            chunks = build_chunks(
                rag_document_id,
                extraction,
                next_vector_id(connection),
                chunk_size,
                chunk_overlap,
            )
            provider, model, dimension, vectors = embed_texts(input_data, [chunk.text for chunk in chunks])
            if vectors and dimension is None:
                dimension = len(vectors[0])

            with connection:
                delete_document_chunks(connection, storage_root, rag_document_id)
                chunks = build_chunks(
                    rag_document_id,
                    extraction,
                    next_vector_id(connection),
                    chunk_size,
                    chunk_overlap,
                )
                ensure_index_profile(connection, provider, model, dimension or 0)
                add_vectors(storage_root, provider, chunks, vectors)
                connection.executemany(
                    """
                    INSERT INTO rag_chunks (
                      chunk_id, rag_document_id, chunk_index, text, heading,
                      page, section_index, faiss_vector_id
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            chunk.chunk_id,
                            chunk.rag_document_id,
                            chunk.chunk_index,
                            chunk.text,
                            chunk.heading,
                            chunk.page,
                            chunk.section_index,
                            chunk.faiss_vector_id,
                        )
                        for chunk in chunks
                    ],
                )
                insert_chunk_fts_rows(connection, chunks, str(row["original_filename"]))
                connection.execute(
                    """
                    UPDATE rag_documents
                    SET status = 'indexed', indexed_at = ?, embedding_provider = ?,
                        embedding_model = ?, embedding_dimension = ?, chunk_count = ?
                    WHERE rag_document_id = ?
                    """,
                    (now_iso(), provider, model, dimension, len(chunks), rag_document_id),
                )
        except WorkerError as error:
            with connection:
                delete_document_chunks(connection, storage_root, rag_document_id)
                connection.execute("UPDATE rag_documents SET status = 'failed' WHERE rag_document_id = ?", (rag_document_id,))
            raise error
        except Exception as error:
            with connection:
                delete_document_chunks(connection, storage_root, rag_document_id)
                connection.execute("UPDATE rag_documents SET status = 'failed' WHERE rag_document_id = ?", (rag_document_id,))
            raise WorkerError("indexing_failed", "RAG indexing failed.") from error

        return {"status": "indexed", "document": row_to_document(connection, get_document_row(connection, rag_document_id))}


def workspace_ids_for_document(connection: sqlite3.Connection, rag_document_id: str) -> list[str]:
    rows = connection.execute(
        "SELECT workspace_id FROM rag_document_workspaces WHERE rag_document_id = ? ORDER BY workspace_id",
        (rag_document_id,),
    ).fetchall()
    return [row["workspace_id"] for row in rows]


def get_document_scope(
    document_workspace_ids: list[str],
    query_workspace_ids: list[str],
    include_global: bool,
) -> str | None:
    if not document_workspace_ids:
        return "global" if include_global else None
    if not query_workspace_ids:
        return None
    return "workspace" if set(document_workspace_ids).intersection(query_workspace_ids) else None


def document_matches_security(document_security: str, query_security: str) -> bool:
    if document_security == "private":
        return query_security == "private"
    return query_security != "private"


def chunk_result_from_row(
    connection: sqlite3.Connection,
    row: sqlite3.Row,
    dense_score: float | None,
    dense_rank: int | None,
    lexical_score: float | None,
    lexical_rank: int | None,
    hybrid_score: float,
    adjusted_score: float,
    scope: str,
) -> dict[str, Any]:
    return {
        "score": adjusted_score,
        "rawScore": dense_score if dense_score is not None else 0.0,
        "denseScore": dense_score,
        "denseRank": dense_rank,
        "lexicalScore": lexical_score,
        "lexicalRank": lexical_rank,
        "hybridScore": hybrid_score,
        "adjustedScore": adjusted_score,
        "scope": scope,
        "ragDocumentId": row["rag_document_id"],
        "filename": row["original_filename"],
        "workspaceIds": workspace_ids_for_document(connection, row["rag_document_id"]),
        "security": row["security"],
        "chunkId": row["chunk_id"],
        "chunkIndex": row["chunk_index"],
        "heading": row["heading"],
        "page": row["page"],
        "text": row["text"],
    }


def search_test_vectors(storage_root: Path, query_vector: list[float], candidate_limit: int) -> list[tuple[int, float]]:
    scored: list[tuple[int, float]] = []
    for raw_id, vector in load_test_vectors(storage_root).items():
        score = sum(float(left) * float(right) for left, right in zip(query_vector, vector))
        scored.append((int(raw_id), score))
    return sorted(scored, key=lambda item: item[1], reverse=True)[:candidate_limit]


def search_faiss_vectors(storage_root: Path, query_vector: list[float], candidate_limit: int) -> list[tuple[int, float]]:
    faiss, numpy = import_faiss_modules()
    index_path = get_index_path(storage_root)
    if not index_path.exists():
        return []
    index = faiss.read_index(str(index_path))
    scores, ids = index.search(l2_normalize([query_vector], numpy), candidate_limit)
    results: list[tuple[int, float]] = []
    for vector_id, score in zip(ids[0], scores[0]):
        if int(vector_id) >= 0:
            results.append((int(vector_id), float(score)))
    return results


def quote_fts_value(value: str) -> str:
    escaped = value.replace('"', '""')
    return f'"{escaped}"'


def build_fts_queries(query: str) -> list[str]:
    normalized = normalize_text(query)
    if not normalized:
        return []

    raw_terms = [
        term.lower()
        for term in re.findall(r"[0-9A-Za-z_\uac00-\ud7a3]+", normalized, flags=re.UNICODE)
        if len(term) >= 2
    ]
    stop_terms = {
        "\ubb34\uc5c7",
        "\ubb34\uc5c7\uc778\uac00",
        "\uc5b4\ub5bb\uac8c",
        "\uc54c\ub824\uc918",
        "\uc54c\ub824\uc8fc\uc138\uc694",
        "\uad00\ub828",
        "\uad00\ub828\ub41c",
    }
    particle_suffixes = [
        "\uc785\ub2c8\uae4c",
        "\uc778\uac00",
        "\uc5d0\uc11c",
        "\uc73c\ub85c",
        "\uc740",
        "\ub294",
        "\uc774",
        "\uac00",
        "\uc744",
        "\ub97c",
        "\uc758",
        "\uc5d0",
        "\uc640",
        "\uacfc",
        "\ub3c4",
        "\ub9cc",
        "\ub85c",
    ]

    terms: list[str] = []
    seen_terms: set[str] = set()
    for raw_term in raw_terms:
        candidates = [raw_term]
        for suffix in particle_suffixes:
            if raw_term.endswith(suffix) and len(raw_term) > len(suffix) + 1:
                candidates.append(raw_term[: -len(suffix)])
                break
        for candidate in candidates:
            if candidate in stop_terms or len(candidate) < 2 or candidate in seen_terms:
                continue
            seen_terms.add(candidate)
            terms.append(candidate)

    queries = [quote_fts_value(term) for term in terms]
    if terms:
        queries.append(" OR ".join(quote_fts_value(term) for term in terms))

    result: list[str] = []
    seen: set[str] = set()
    for item in queries:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def search_lexical_vectors(
    connection: sqlite3.Connection,
    query: str,
    candidate_limit: int,
) -> list[tuple[int, float]]:
    best_by_vector_id: dict[int, tuple[int, float]] = {}
    for fts_query in build_fts_queries(query):
        try:
            rows = connection.execute(
                """
                SELECT rag_chunks.faiss_vector_id AS faiss_vector_id,
                       bm25(rag_chunks_fts) AS lexical_score
                FROM rag_chunks_fts
                INNER JOIN rag_chunks
                  ON rag_chunks.chunk_id = rag_chunks_fts.chunk_id
                WHERE rag_chunks_fts MATCH ?
                ORDER BY lexical_score ASC
                LIMIT ?
                """,
                (fts_query, candidate_limit),
            ).fetchall()
        except sqlite3.OperationalError:
            continue

        for rank, row in enumerate(rows, start=1):
            vector_id = int(row["faiss_vector_id"])
            lexical_score = float(row["lexical_score"])
            current = best_by_vector_id.get(vector_id)
            if current is None or rank < current[0]:
                best_by_vector_id[vector_id] = (rank, lexical_score)

    return [
        (vector_id, lexical_score)
        for vector_id, (_, lexical_score) in sorted(
            best_by_vector_id.items(),
            key=lambda item: (item[1][0], item[1][1]),
        )
    ][:candidate_limit]


def rrf_score(rank: int | None, weight: float) -> float:
    if rank is None:
        return 0.0
    return weight / (DEFAULT_RRF_K + rank)


def search_documents(input_data: dict[str, Any]) -> list[dict[str, Any]]:
    storage_root = get_storage_root(input_data)
    query = input_data.get("query")
    if not isinstance(query, str) or not query.strip():
        raise WorkerError("invalid_input", "query is required.")

    workspace_ids = validate_workspace_ids({"workspace_ids": input_data.get("workspace_ids", [])})
    include_global = input_data.get("include_global") is not False
    top_k = input_data.get("top_k")
    top_k_value = max(1, min(20, int(top_k) if isinstance(top_k, int) else 5))
    query_security = str(input_data.get("security") or "internal")
    requested_provider = str(input_data.get("embedding_provider") or DEFAULT_EMBEDDING_PROVIDER).strip().lower()
    if query_security == "private" and requested_provider != "local":
        raise WorkerError("private_external_embedding_blocked", "Private RAG search requires local embedding.")
    similarity_threshold = get_float_config(
        input_data,
        "similarity_threshold",
        DEFAULT_SIMILARITY_THRESHOLD,
        -1.0,
        1.0,
    )
    dense_only_threshold = max(
        similarity_threshold,
        get_float_config(
            input_data,
            "dense_only_threshold",
            DEFAULT_DENSE_ONLY_THRESHOLD,
            -1.0,
            1.0,
        ),
    )
    workspace_score_bonus = get_float_config(
        input_data,
        "workspace_score_bonus",
        DEFAULT_WORKSPACE_SCORE_BONUS,
        0.0,
        0.2,
    )
    max_chunks_per_document = get_int_config(
        input_data,
        "max_chunks_per_document",
        DEFAULT_MAX_CHUNKS_PER_DOCUMENT,
        1,
        10,
    )
    configured_candidate_count = get_int_config(
        input_data,
        "search_candidate_count",
        DEFAULT_SEARCH_CANDIDATE_COUNT,
        10,
        200,
    )

    provider, model, dimension, vectors = embed_texts(input_data, [query.strip()])
    if vectors and dimension is None:
        dimension = len(vectors[0])
    with connect_database(storage_root) as connection:
        assert_index_profile(connection, provider, model, dimension or 0)

    candidate_limit = max(top_k_value * 10, configured_candidate_count)
    dense_candidates = (
        search_test_vectors(storage_root, vectors[0], candidate_limit)
        if provider == "test"
        else search_faiss_vectors(storage_root, vectors[0], candidate_limit)
    )
    dense_by_vector_id = {
        vector_id: {"rank": rank, "score": score}
        for rank, (vector_id, score) in enumerate(dense_candidates, start=1)
    }
    with connect_database(storage_root) as connection:
        lexical_candidates = search_lexical_vectors(connection, query.strip(), candidate_limit)
    lexical_by_vector_id = {
        vector_id: {"rank": rank, "score": score}
        for rank, (vector_id, score) in enumerate(lexical_candidates, start=1)
    }
    candidate_vector_ids = set(dense_by_vector_id).union(lexical_by_vector_id)

    scored_results: list[dict[str, Any]] = []
    seen_chunks: set[str] = set()
    with connect_database(storage_root) as connection:
        for vector_id in candidate_vector_ids:
            row = connection.execute(
                """
                SELECT rag_chunks.*, rag_documents.original_filename, rag_documents.security
                FROM rag_chunks
                INNER JOIN rag_documents
                  ON rag_documents.rag_document_id = rag_chunks.rag_document_id
                WHERE rag_chunks.faiss_vector_id = ?
                  AND rag_documents.status = 'indexed'
                """,
                (vector_id,),
            ).fetchone()
            if not row or row["chunk_id"] in seen_chunks:
                continue
            document_workspace_ids = workspace_ids_for_document(connection, row["rag_document_id"])
            scope = get_document_scope(document_workspace_ids, workspace_ids, include_global)
            if scope is None:
                continue
            if not document_matches_security(str(row["security"]), query_security):
                continue
            dense_candidate = dense_by_vector_id.get(vector_id)
            lexical_candidate = lexical_by_vector_id.get(vector_id)
            dense_rank = int(dense_candidate["rank"]) if dense_candidate else None
            dense_score = float(dense_candidate["score"]) if dense_candidate else None
            lexical_rank = int(lexical_candidate["rank"]) if lexical_candidate else None
            lexical_score = float(lexical_candidate["score"]) if lexical_candidate else None

            if lexical_rank is None and (dense_score is None or dense_score < dense_only_threshold):
                continue
            hybrid_score = rrf_score(dense_rank, DEFAULT_DENSE_RRF_WEIGHT) + rrf_score(
                lexical_rank,
                DEFAULT_LEXICAL_RRF_WEIGHT,
            )
            adjusted_score = hybrid_score + (workspace_score_bonus if scope == "workspace" else 0.0)
            seen_chunks.add(row["chunk_id"])
            scored_results.append(
                chunk_result_from_row(
                    connection,
                    row,
                    dense_score,
                    dense_rank,
                    lexical_score,
                    lexical_rank,
                    hybrid_score,
                    adjusted_score,
                    scope,
                )
            )

    final_results: list[dict[str, Any]] = []
    document_counts: dict[str, int] = {}
    sorted_results = sorted(scored_results, key=lambda item: item["adjustedScore"], reverse=True)
    for hybrid_rank, result in enumerate(sorted_results, start=1):
        result["hybridRank"] = hybrid_rank
        rag_document_id = str(result["ragDocumentId"])
        if document_counts.get(rag_document_id, 0) >= max_chunks_per_document:
            continue
        document_counts[rag_document_id] = document_counts.get(rag_document_id, 0) + 1
        final_results.append(result)
        if len(final_results) >= top_k_value:
            break
    return final_results


def runtime_check(_: dict[str, Any]) -> dict[str, Any]:
    return {"available": True, "executable": sys.executable, "version": sys.version.split()[0]}


def check_embedding_status(input_data: dict[str, Any]) -> dict[str, Any]:
    provider = create_embedding_provider(input_data)

    if provider.provider_name == "local" and isinstance(provider, LocalEmbeddingProvider):
        try:
            tags = get_json(f"{provider.ollama_base_url}/api/tags")
        except urllib.error.URLError as error:
            raise WorkerError("ollama_unavailable", "Ollama is not reachable.") from error
        except Exception as error:
            raise WorkerError("ollama_unavailable", "Ollama is not reachable.") from error

        models = tags.get("models")
        model_names = {
            model.get("name")
            for model in models
            if isinstance(model, dict) and isinstance(model.get("name"), str)
        } if isinstance(models, list) else set()
        accepted_model_names = {provider.model_name}
        if ":" not in provider.model_name:
            accepted_model_names.add(f"{provider.model_name}:latest")

        if model_names.isdisjoint(accepted_model_names):
            raise WorkerError(
                "embedding_model_unavailable",
                "Local embedding model is not available.",
            )

        vector = provider.embed_query("mimora embedding check")
        return {
            "available": True,
            "provider": "local",
            "model": provider.model_name,
            "endpoint": provider.ollama_base_url,
            "dimension": len(vector),
            "message": "Local embedding model is available.",
        }

    if provider.provider_name == "openai":
        vector = provider.embed_query("mimora embedding check")
        return {
            "available": True,
            "provider": "openai",
            "model": provider.model_name,
            "dimension": len(vector),
            "message": "OpenAI embedding model is available.",
        }

    vector = provider.embed_query("mimora embedding check")
    return {
        "available": True,
        "provider": provider.provider_name,
        "model": provider.model_name,
        "dimension": len(vector),
        "message": "Embedding provider is available.",
    }


def schedule_source_path(input_data: dict[str, Any]) -> Path:
    source_path = input_data.get("source_path")
    if not isinstance(source_path, str) or not source_path.strip():
        raise WorkerError("schedule_source_not_found", "Schedule source path is required.")

    path_value = Path(source_path).resolve()
    if path_value.suffix.lower() not in SUPPORTED_SCHEDULE_EXTENSIONS:
        raise WorkerError("unsupported_schedule_file", "Schedule source must be .xlsx or .xlsm.")
    if not path_value.exists() or not path_value.is_file():
        raise WorkerError("schedule_source_not_found", "Schedule source file was not found.")
    return path_value


def normalize_schedule_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("_x000D_", " ").replace("_x000d_", " ")
    return re.sub(r"[^0-9A-Za-z가-힣]+", "", text.strip().lower())


SCHEDULE_COLUMN_ALIASES = {
    "wbs": ["wbs", "wbsid", "작업코드", "작업id"],
    "level": ["level", "wbslevel", "레벨", "수준"],
    "name": ["taskname", "task", "name", "작업명", "작업", "태스크", "일정명"],
    "planned_start": ["plannedstart", "planstart", "start", "계획시작", "시작일", "시작예정", "착수예정"],
    "planned_finish": ["plannedfinish", "planfinish", "finish", "계획종료", "완료일", "완료예정", "종료예정"],
    "actual_start": ["actualstart", "실제시작", "실제시작일", "실착수"],
    "actual_finish": ["actualfinish", "실제종료", "실제완료일", "실제완료", "실완료"],
    "planned_workload": ["plannedworkload", "planworkload", "계획공수", "계획작업량", "일일계획작업량"],
    "actual_workload": ["actualworkload", "실적공수", "실제공수", "실적작업량", "실제총작업량", "실제작업량"],
    "planned_duration": ["plannedduration", "계획기간", "기간"],
    "actual_duration": ["actualduration", "실제기간", "실적기간", "실제총기간"],
    "resource": ["resource", "resources", "담당자", "담당", "작업자", "리소스"],
    "deliverable": ["deliverable", "output", "산출물", "결과물"],
    "planned_progress": ["plannedprogress", "planprogress", "계획진척", "계획진도", "계획율", "계획", "계획작업량진척률"],
    "actual_progress": ["actualprogress", "progress", "실적진척", "실제진척", "실적", "진척율", "진도율", "실적기성진척률"],
    "calendar": ["calendar", "캘린더", "달력"],
    "predecessors": ["predecessors", "predecessor", "dependencies", "dependency", "선행작업", "선행", "의존성"],
}


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_wbs(value: Any) -> str:
    text = cell_to_text(value)
    if not text:
        return ""
    text = re.sub(r"\s+", "", text)
    parts = text.split(".")
    if all(re.fullmatch(r"\d+", part or "") for part in parts):
        return ".".join(str(int(part)) for part in parts)
    return text


def parse_schedule_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if hasattr(value, "date") and hasattr(value, "isoformat"):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isnan(float(value)):
            return None
        serial = float(value)
        if 1 <= serial <= 60000:
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date().isoformat()
    text = cell_to_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10] if pattern != "%Y%m%d" else text[:8], pattern).date().isoformat()
        except ValueError:
            continue
    return None


def parse_schedule_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if math.isnan(float(value)):
            return None
        return float(value)
    text = cell_to_text(value).replace(",", "").replace("%", "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_schedule_progress(value: Any) -> float | None:
    number = parse_schedule_number(value)
    if number is None:
        return None
    return number / 100 if number > 1 else number


def parse_resource_assignments(value: Any) -> list[dict[str, Any]]:
    text = cell_to_text(value)
    if not text:
        return []
    assignments: list[dict[str, Any]] = []
    for part in re.split(r"[,;/\n]+", text):
        item = part.strip()
        if not item:
            continue
        match = re.match(r"^(.*?)\s*\[\s*([0-9]+(?:\.[0-9]+)?)\s*%\s*\]\s*$", item)
        if match:
            assignments.append({"name": match.group(1).strip(), "allocation": float(match.group(2)) / 100})
        else:
            assignments.append({"name": item, "allocation": 1.0})
    return assignments


def parse_dependency_references(value: Any, successor_wbs: str) -> list[dict[str, Any]]:
    text = cell_to_text(value)
    if not text:
        return []
    dependencies: list[dict[str, Any]] = []
    for part in re.split(r"[,;/\n]+", text):
        item = part.strip()
        if not item:
            continue
        match = re.match(
            r"^(?P<wbs>\d+(?:\.\d+)*)(?:\s*(?P<type>FS|SS|FF|SF))?(?:\s*(?P<lag>[+-]\s*\d+))?\s*(?:d|day|days|일|영업일)?$",
            item,
            re.IGNORECASE,
        )
        if not match:
            continue
        dependency_type = (match.group("type") or "FS").upper()
        lag_text = (match.group("lag") or "0").replace(" ", "")
        dependencies.append({
            "predecessor_wbs": normalize_wbs(match.group("wbs")),
            "successor_wbs": successor_wbs,
            "type": dependency_type,
            "lag_days": int(lag_text),
        })
    return dependencies


def build_schedule_dependencies(tasks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    task_wbs_values = {task.get("wbs") for task in tasks if task.get("wbs")}
    dependencies: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, int]] = set()
    for task in tasks:
        successor_wbs = task.get("wbs")
        if not successor_wbs:
            continue
        for dependency in parse_dependency_references(task.get("predecessorsRaw"), successor_wbs):
            if dependency["predecessor_wbs"] not in task_wbs_values:
                continue
            key = (
                dependency["predecessor_wbs"],
                dependency["successor_wbs"],
                dependency["type"],
                dependency["lag_days"],
            )
            if key in seen:
                continue
            seen.add(key)
            dependencies.append(dependency)
    return dependencies


def schedule_header_matches(normalized: str, aliases: set[str]) -> bool:
    if normalized in aliases:
        return True
    return any(alias and len(alias) >= 6 and alias in normalized for alias in aliases)


def detect_header_mapping(rows: list[tuple[Any, ...]], aliases: dict[str, list[str]], required: set[str]) -> tuple[int, dict[str, int]]:
    normalized_aliases = {
        field: {normalize_schedule_header(alias) for alias in values}
        for field, values in aliases.items()
    }
    best_row = -1
    best_mapping: dict[str, int] = {}
    best_score = -1
    for row_index, row in enumerate(rows[:30]):
        mapping: dict[str, int] = {}
        normalized_cells = [normalize_schedule_header(value) for value in row]
        for cell_index, normalized in enumerate(normalized_cells):
            if not normalized:
                continue
            for field, field_aliases in normalized_aliases.items():
                if field in mapping:
                    continue
                if schedule_header_matches(normalized, field_aliases):
                    mapping[field] = cell_index
        score = len(mapping)
        if score > best_score:
            best_score = score
            best_row = row_index
            best_mapping = mapping
    if not required.issubset(best_mapping.keys()):
        raise WorkerError("schedule_column_mapping_failed", "Schedule column mapping failed.")
    return best_row, best_mapping


def row_value(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def cell_merge_span(sheet: Any, row_number: int, column_number: int) -> tuple[int, int] | None:
    merged_ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", [])
    for merged_range in merged_ranges:
        if (
            merged_range.min_row <= row_number <= merged_range.max_row and
            merged_range.min_col <= column_number <= merged_range.max_col
        ):
            return merged_range.min_col - 1, merged_range.max_col - 1
    return None


def detect_task_name_columns(sheet: Any, header_row: int, mapping: dict[str, int]) -> list[int]:
    name_index = mapping.get("name")
    if name_index is None:
        return []

    merged_span = cell_merge_span(sheet, header_row + 1, name_index + 1)
    if merged_span:
        return list(range(merged_span[0], merged_span[1] + 1))

    later_mapped_indices = [
        index
        for field, index in mapping.items()
        if field != "name" and index > name_index
    ]
    if later_mapped_indices:
        return list(range(name_index, min(later_mapped_indices)))
    return [name_index]


def first_text_in_columns(row: tuple[Any, ...], column_indices: list[int]) -> str:
    for index in column_indices:
        if index < len(row):
            text = cell_to_text(row[index])
            if text:
                return text
    return ""


def schedule_row_has_signal(row: tuple[Any, ...], mapping: dict[str, int]) -> bool:
    signal_fields = [
        "wbs",
        "name",
        "planned_start",
        "planned_finish",
        "actual_start",
        "actual_finish",
        "planned_progress",
        "actual_progress",
        "resource",
    ]
    return any(cell_to_text(row_value(row, mapping, field)) for field in signal_fields)


def schedule_parse_warning(
    row_number: int,
    reason: str,
    row: tuple[Any, ...],
    mapping: dict[str, int],
    name_columns: list[int],
) -> dict[str, Any]:
    raw_title = first_text_in_columns(row, name_columns)
    raw_date = (
        cell_to_text(row_value(row, mapping, "planned_start"))
        or cell_to_text(row_value(row, mapping, "planned_finish"))
        or cell_to_text(row_value(row, mapping, "actual_start"))
        or cell_to_text(row_value(row, mapping, "actual_finish"))
        or None
    )

    return {
        "row": row_number,
        "reason": reason,
        "rawTask": cell_to_text(row_value(row, mapping, "wbs")) or None,
        "rawTitle": raw_title or None,
        "rawDate": raw_date,
    }


def infer_wbs_level(wbs: str, explicit_level: Any) -> int:
    explicit = parse_schedule_number(explicit_level)
    if explicit is not None and explicit >= 0:
        return int(explicit)
    if not wbs:
        return 0
    return len([part for part in re.split(r"[.\-_/]+", wbs) if part])


def infer_parent_wbs(wbs: str) -> str | None:
    if "." in wbs:
        parent = ".".join(wbs.split(".")[:-1]).strip()
        return parent or None
    return None


def parse_schedule_sheet(workbook: Any, workspace_id: str, source_path: Path) -> tuple[list[dict[str, Any]], dict[str, str], list[dict[str, Any]]]:
    if "Schedule" not in workbook.sheetnames:
        raise WorkerError("schedule_sheet_missing", "Schedule sheet is missing.")
    sheet = workbook["Schedule"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise WorkerError("schedule_parse_failed", "Schedule sheet is empty.")
    header_row, mapping = detect_header_mapping(rows, SCHEDULE_COLUMN_ALIASES, {"wbs", "name"})
    name_columns = detect_task_name_columns(sheet, header_row, mapping)
    tasks: list[dict[str, Any]] = []
    parse_warnings: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        wbs = normalize_wbs(row_value(row, mapping, "wbs"))
        name = first_text_in_columns(row, name_columns)
        if not wbs and not name:
            if schedule_row_has_signal(row, mapping):
                parse_warnings.append(
                    schedule_parse_warning(
                        row_number,
                        "missing_wbs_and_task_name",
                        row,
                        mapping,
                        name_columns,
                    )
                )
            continue
        if not name or name.lower() in {"summary", "total", "합계"}:
            if not name and schedule_row_has_signal(row, mapping):
                parse_warnings.append(
                    schedule_parse_warning(
                        row_number,
                        "missing_task_name",
                        row,
                        mapping,
                        name_columns,
                    )
                )
            continue
        task = {
            "taskId": f"{workspace_id}:{wbs or row_number}",
            "wbs": wbs or str(row_number),
            "level": infer_wbs_level(wbs, row_value(row, mapping, "level")),
            "name": name,
            "parentWbs": infer_parent_wbs(wbs),
            "isLeaf": True,
            "plannedStart": parse_schedule_date(row_value(row, mapping, "planned_start")),
            "plannedFinish": parse_schedule_date(row_value(row, mapping, "planned_finish")),
            "actualStart": parse_schedule_date(row_value(row, mapping, "actual_start")),
            "actualFinish": parse_schedule_date(row_value(row, mapping, "actual_finish")),
            "plannedWorkload": parse_schedule_number(row_value(row, mapping, "planned_workload")),
            "actualWorkload": parse_schedule_number(row_value(row, mapping, "actual_workload")),
            "plannedDuration": parse_schedule_number(row_value(row, mapping, "planned_duration")),
            "actualDuration": parse_schedule_number(row_value(row, mapping, "actual_duration")),
            "plannedProgress": parse_schedule_progress(row_value(row, mapping, "planned_progress")),
            "actualProgress": parse_schedule_progress(row_value(row, mapping, "actual_progress")),
            "resource": parse_resource_assignments(row_value(row, mapping, "resource")),
            "predecessorsRaw": cell_to_text(row_value(row, mapping, "predecessors")) or None,
            "deliverable": cell_to_text(row_value(row, mapping, "deliverable")) or None,
            "calendar": cell_to_text(row_value(row, mapping, "calendar")) or None,
        }
        tasks.append(task)
    parent_wbs_values = {task["parentWbs"] for task in tasks if task["parentWbs"]}
    for task in tasks:
        task["isLeaf"] = task["wbs"] not in parent_wbs_values
    column_mapping = {field: cell_to_text(rows[header_row][index]) for field, index in mapping.items()}
    if name_columns:
        column_mapping["nameColumns"] = ",".join(str(index + 1) for index in name_columns)
    return tasks, column_mapping, parse_warnings


def parse_calendar_sheet(workbook: Any) -> list[dict[str, Any]]:
    if "Calendar" not in workbook.sheetnames:
        return []
    sheet = workbook["Calendar"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    try:
        header_row, mapping = detect_header_mapping(
            rows,
            {
                "date": ["date", "일자", "날짜"],
                "type": ["type", "구분", "근무구분", "휴일구분"],
                "calendar": ["calendar", "캘린더", "달력"],
                "name": ["name", "명칭", "휴일명", "비고"],
            },
            {"date"},
        )
    except WorkerError:
        result: list[dict[str, Any]] = []
        seen_dates: set[str] = set()
        for row in rows:
            for date_index in (1, 4, 7):
                if date_index >= len(row):
                    continue
                date = parse_schedule_date(row[date_index])
                if not date or date in seen_dates:
                    continue
                seen_dates.add(date)
                result.append({
                    "date": date,
                    "type": "holiday",
                    "calendar": None,
                    "name": cell_to_text(row[date_index + 1]) if date_index + 1 < len(row) else None,
                })
        return result
    result: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        date = parse_schedule_date(row_value(row, mapping, "date"))
        if not date:
            continue
        type_text = cell_to_text(row_value(row, mapping, "type")).lower()
        day_type = "working"
        if "holiday" in type_text or "휴" in type_text:
            day_type = "holiday"
        elif "non" in type_text or "비" in type_text:
            day_type = "non-working"
        result.append({
            "date": date,
            "type": day_type,
            "calendar": cell_to_text(row_value(row, mapping, "calendar")) or None,
            "name": cell_to_text(row_value(row, mapping, "name")) or None,
        })
    return result


def parse_progress_sheet(workbook: Any) -> list[dict[str, Any]]:
    sheet_name = "Progress_Data" if "Progress_Data" in workbook.sheetnames else "Progress"
    if sheet_name not in workbook.sheetnames:
        return []
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if not rows:
        return []
    try:
        header_row, mapping = detect_header_mapping(
            rows,
            {
                "date": ["date", "일자", "날짜"],
                "planned_workload": ["plannedworkload", "계획공수", "일일계획작업량"],
                "cumulative_planned_workload": ["cumulativeplannedworkload", "누적계획공수", "계획작업량누적"],
                "planned_progress": ["plannedprogress", "계획진척", "계획진도", "계획작업량진척률"],
                "earned_value": ["earnedvalue", "ev", "획득가치", "실적기성"],
                "cumulative_earned_value": ["cumulativeearnedvalue", "누적획득가치", "실적기성누적"],
                "actual_progress": ["actualprogress", "실적진척", "실제진척", "실적기성진척률"],
                "actual_workload": ["actualworkload", "실적공수", "일일실제투입작업량"],
                "cumulative_actual_workload": ["cumulativeactualworkload", "누적실적공수", "실제투입작업량누적"],
            },
            {"date"},
        )
    except WorkerError:
        return []
    points: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        date = parse_schedule_date(row_value(row, mapping, "date"))
        if not date:
            continue
        points.append({
            "date": date,
            "plannedWorkload": parse_schedule_number(row_value(row, mapping, "planned_workload")),
            "cumulativePlannedWorkload": parse_schedule_number(row_value(row, mapping, "cumulative_planned_workload")),
            "plannedProgress": parse_schedule_progress(row_value(row, mapping, "planned_progress")),
            "earnedValue": parse_schedule_number(row_value(row, mapping, "earned_value")),
            "cumulativeEarnedValue": parse_schedule_number(row_value(row, mapping, "cumulative_earned_value")),
            "actualProgress": parse_schedule_progress(row_value(row, mapping, "actual_progress")),
            "actualWorkload": parse_schedule_number(row_value(row, mapping, "actual_workload")),
            "cumulativeActualWorkload": parse_schedule_number(row_value(row, mapping, "cumulative_actual_workload")),
        })
    return points


def parse_settings_sheet(workbook: Any) -> dict[str, Any]:
    if "Settings" not in workbook.sheetnames:
        return {}
    settings: dict[str, Any] = {}
    for row_index, row in enumerate(workbook["Settings"].iter_rows(values_only=True)):
        if not row or len(row) < 2:
            continue
        if row_index == 0 and normalize_schedule_header(row[1] if len(row) > 1 else None) == "item":
            continue
        key_index = 1 if len(row) > 2 and cell_to_text(row[1]) else 0
        value_index = 2 if key_index == 1 and len(row) > 2 else 1
        key = cell_to_text(row[key_index])
        if not key:
            continue
        value = row[value_index] if value_index < len(row) else None
        if hasattr(value, "isoformat"):
            settings[key] = parse_schedule_date(value)
        elif isinstance(value, (str, int, float, bool)) or value is None:
            settings[key] = value
        else:
            settings[key] = cell_to_text(value)
    return settings


def average_progress(tasks: list[dict[str, Any]], field: str) -> float | None:
    values = [task[field] for task in tasks if isinstance(task.get(field), (int, float))]
    if not values:
        return None
    return sum(float(value) for value in values) / len(values)


def is_active_task(task: dict[str, Any], as_of_date: str) -> bool:
    planned_start = task.get("plannedStart")
    planned_finish = task.get("plannedFinish")
    actual_finish = task.get("actualFinish")
    if actual_finish:
        return False
    if planned_start and planned_start <= as_of_date and (not planned_finish or planned_finish >= as_of_date):
        return True
    progress = task.get("actualProgress")
    return isinstance(progress, (int, float)) and 0 < float(progress) < 1


def is_delayed_task(task: dict[str, Any], as_of_date: str) -> bool:
    planned_finish = task.get("plannedFinish")
    actual_finish = task.get("actualFinish")
    actual_progress = task.get("actualProgress")
    if not planned_finish or planned_finish >= as_of_date:
        return False
    completed = bool(actual_finish) or (
        isinstance(actual_progress, (int, float)) and float(actual_progress) >= 1
    )
    return not completed


def schedule_summary(schedule: dict[str, Any], source: dict[str, Any], as_of_date: str | None = None) -> dict[str, Any]:
    now_date = as_of_date or datetime.now().date().isoformat()
    tasks = schedule.get("tasks", [])
    leaf_tasks = [task for task in tasks if task.get("isLeaf")]
    active_tasks = [task for task in tasks if is_active_task(task, now_date)]
    active_leaf_tasks = [task for task in leaf_tasks if is_active_task(task, now_date)]
    delayed_wbs_nodes = [task for task in tasks if is_delayed_task(task, now_date)]
    delayed_leaf_tasks = [task for task in leaf_tasks if is_delayed_task(task, now_date)]
    progress_series = schedule.get("progressSeries", [])
    last_progress = progress_series[-1] if progress_series else {}
    return {
        "workspaceId": schedule["workspaceId"],
        "filename": schedule["filename"],
        "modifiedAt": source["modifiedAt"],
        "lastParsedAt": source.get("lastParsedAt") or now_iso(),
        "parseStatus": source.get("parseStatus", "parsed"),
        "taskCount": len(tasks),
        "leafTaskCount": len(leaf_tasks),
        "activeTaskCount": len(active_leaf_tasks),
        "activeWbsNodeCount": len(active_tasks),
        "activeLeafTaskCount": len(active_leaf_tasks),
        "delayedTaskCount": len(delayed_leaf_tasks),
        "delayedWbsNodeCount": len(delayed_wbs_nodes),
        "plannedProgress": last_progress.get("plannedProgress") if last_progress else average_progress(tasks, "plannedProgress"),
        "actualProgress": last_progress.get("actualProgress") if last_progress else average_progress(tasks, "actualProgress"),
        "progressAsOfDate": now_date,
        "projectStart": schedule.get("projectStart"),
        "projectFinish": schedule.get("projectFinish"),
        "parsedSheets": schedule.get("parsedSheets", []),
        "skippedRows": len(schedule.get("parseWarnings", []) or []),
        "parseWarnings": schedule.get("parseWarnings", []) or [],
        "source": source,
    }


def parse_schedule(input_data: dict[str, Any]) -> dict[str, Any]:
    workspace_id = input_data.get("workspace_id")
    if not isinstance(workspace_id, str) or not workspace_id.strip():
        raise WorkerError("invalid_input", "workspace_id is required.")
    source_path = schedule_source_path(input_data)
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as error:
        raise WorkerError("schedule_parse_failed", "openpyxl is required to parse Schedule Excel files.") from error
    try:
        workbook = load_workbook(source_path, read_only=False, data_only=True, keep_vba=False)
        sheet_names = list(workbook.sheetnames)
        tasks, column_mapping, parse_warnings = parse_schedule_sheet(workbook, workspace_id, source_path)
        dependencies = build_schedule_dependencies(tasks)
        calendars = parse_calendar_sheet(workbook)
        progress_series = parse_progress_sheet(workbook)
        settings = parse_settings_sheet(workbook)
    except WorkerError:
        raise
    except Exception as error:
        raise WorkerError("schedule_parse_failed", "Schedule workbook parsing failed.") from error
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    planned_starts = [task["plannedStart"] for task in tasks if task.get("plannedStart")]
    planned_finishes = [task["plannedFinish"] for task in tasks if task.get("plannedFinish")]
    project_start = parse_schedule_date(settings.get("project start") or settings.get("project_start")) or (min(planned_starts) if planned_starts else None)
    project_finish = parse_schedule_date(settings.get("project finish") or settings.get("project_finish")) or (max(planned_finishes) if planned_finishes else None)
    parsed_sheets = [
        sheet_name
        for sheet_name in ["Schedule", "Calendar", "Progress", "Progress_Data", "Settings"]
        if sheet_name in sheet_names
    ]
    source = {
        "workspaceId": workspace_id,
        "sourcePath": str(source_path),
        "filename": source_path.name,
        "fileSize": source_path.stat().st_size,
        "modifiedAt": datetime.fromtimestamp(source_path.stat().st_mtime, timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "lastParsedAt": now_iso(),
        "parseStatus": "parsed",
    }
    schedule = {
        "workspaceId": workspace_id,
        "sourceFile": str(source_path),
        "filename": source_path.name,
        "projectStart": project_start,
        "projectFinish": project_finish,
        "tasks": tasks,
        "dependencies": dependencies,
        "calendars": calendars,
        "progressSeries": progress_series,
        "settings": settings,
        "parsedSheets": parsed_sheets,
        "columnMapping": column_mapping,
        "parseWarnings": parse_warnings,
    }
    return {
        "source": source,
        "schedule": schedule,
        "summary": schedule_summary(schedule, source),
        "fromCache": False,
    }


def schedule_register(input_data: dict[str, Any]) -> dict[str, Any]:
    workspace_id = input_data.get("workspace_id")
    if not isinstance(workspace_id, str) or not workspace_id.strip():
        raise WorkerError("invalid_input", "workspace_id is required.")
    source_path = schedule_source_path(input_data)
    stats = source_path.stat()
    return {
        "workspaceId": workspace_id,
        "sourcePath": str(source_path),
        "filename": source_path.name,
        "fileSize": stats.st_size,
        "modifiedAt": datetime.fromtimestamp(stats.st_mtime, timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "lastParsedAt": None,
        "parseStatus": "registered",
    }


def compact_task(task: dict[str, Any]) -> str:
    resources = ", ".join(item.get("name", "") for item in task.get("resource", []) if item.get("name"))
    parts = [
        f"WBS: {task.get('wbs')}",
        f"Task: {task.get('name')}",
        f"Planned Finish: {task.get('plannedFinish') or '-'}",
        f"Actual Finish: {task.get('actualFinish') or '-'}",
        f"Planned Progress: {format_percent(task.get('plannedProgress'))}",
        f"Actual Progress: {format_percent(task.get('actualProgress'))}",
    ]
    if resources:
        parts.append(f"Resource: {resources}")
    return " | ".join(parts)


def format_percent(value: Any) -> str:
    if not isinstance(value, (int, float)):
        return "-"
    return f"{round(float(value) * 100, 1)}%"


def schedule_context_text(kind: str, summary: dict[str, Any], tasks: list[dict[str, Any]]) -> str:
    lines = [
        "[SCHEDULE CONTEXT]",
        f"Source: {summary['filename']}",
        f"Last Parsed: {summary['lastParsedAt']}",
        f"Project Start: {summary.get('projectStart') or '-'}",
        f"Project Finish: {summary.get('projectFinish') or '-'}",
        "Project Progress",
        f"Planned: {format_percent(summary.get('plannedProgress'))}",
        f"Actual: {format_percent(summary.get('actualProgress'))}",
        f"Tasks: {summary['taskCount']}",
        f"Active Tasks: {summary['activeTaskCount']}",
        f"Delayed Tasks: {summary['delayedTaskCount']}",
        f"Query Result: {kind}",
    ]
    if tasks:
        lines.append("Tasks")
        for index, task in enumerate(tasks[:20], start=1):
            lines.append(f"{index}. {compact_task(task)}")
    else:
        lines.append("Tasks: No matching schedule tasks.")
    lines.append("[/SCHEDULE CONTEXT]")
    return "\n".join(lines)


def normalize_schedule_lookup_text(value: Any) -> str:
    return re.sub(r"\s+", "", cell_to_text(value)).lower()


def strip_query_noise(query: str) -> str:
    text = re.sub(r"[?？!！.。]+", " ", cell_to_text(query))
    for token in [
        "알려줘",
        "알려 주세요",
        "보여줘",
        "보여 주세요",
        "무엇인가",
        "무엇이야",
        "어떻게 되어 있어",
        "어떻게 돼",
    ]:
        text = text.replace(token, " ")
    return re.sub(r"\s+", " ", text).strip()


def extract_resource_lookup_target(query: str) -> str | None:
    text = strip_query_noise(query)
    match = re.search(
        r"(.+?)\s*(?:담당\s*작업|담당자|담당|resource|리소스)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return None
    target = match.group(1).strip()
    return target or None


def resource_matches_target(task: dict[str, Any], target: str) -> bool:
    normalized_target = normalize_schedule_lookup_text(target)
    if not normalized_target:
        return False
    return any(
        normalize_schedule_lookup_text(resource.get("name")) == normalized_target
        for resource in task.get("resource", [])
        if isinstance(resource, dict)
    )


def collect_schedule_resource_names(tasks: list[dict[str, Any]]) -> list[str]:
    names: dict[str, str] = {}
    for task in tasks:
        for resource in task.get("resource", []):
            if not isinstance(resource, dict):
                continue
            name = cell_to_text(resource.get("name"))
            normalized = normalize_schedule_lookup_text(name)
            if normalized and normalized not in names:
                names[normalized] = name
    return sorted(names.values(), key=len, reverse=True)


def find_resource_entity(query: str, tasks: list[dict[str, Any]]) -> str | None:
    normalized_query = normalize_schedule_lookup_text(query)
    for name in collect_schedule_resource_names(tasks):
        if normalize_schedule_lookup_text(name) in normalized_query:
            return name
    return None


def extract_task_lookup_target(query: str) -> str | None:
    text = strip_query_noise(query)
    text = re.sub(r"\b(?:schedule|task|wbs)\b", " ", text, flags=re.IGNORECASE)
    for token in [
        "일정은",
        "일정",
        "작업은",
        "작업",
        "진척은",
        "진척",
        "진도는",
        "진도",
        "계획은",
        "계획",
        "상태는",
        "상태",
    ]:
        text = text.replace(token, " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def find_task_entity(query: str, tasks: list[dict[str, Any]]) -> tuple[str, str] | tuple[None, None]:
    normalized_query = normalize_schedule_lookup_text(query)
    task_entities: list[tuple[str, str]] = []
    for task in tasks:
        wbs = cell_to_text(task.get("wbs"))
        name = cell_to_text(task.get("name"))
        if wbs:
            task_entities.append(("wbs", wbs))
        if name:
            task_entities.append(("task", name))
    task_entities.sort(key=lambda item: len(item[1]), reverse=True)
    for entity_type, value in task_entities:
        normalized_value = normalize_schedule_lookup_text(value)
        if normalized_value and normalized_value in normalized_query:
            return entity_type, value
    return None, None


def lookup_tasks_by_name_or_wbs(tasks: list[dict[str, Any]], target: str) -> list[dict[str, Any]]:
    normalized_target = normalize_schedule_lookup_text(target)
    if not normalized_target:
        return []
    exact = [
        task
        for task in tasks
        if normalize_schedule_lookup_text(task.get("wbs")) == normalized_target
        or normalize_schedule_lookup_text(task.get("name")) == normalized_target
    ]
    if exact:
        return exact
    return [
        task
        for task in tasks
        if normalized_target in normalize_schedule_lookup_text(task.get("wbs"))
        or normalized_target in normalize_schedule_lookup_text(task.get("name"))
    ]


def is_completed_task(task: dict[str, Any]) -> bool:
    actual_progress = task.get("actualProgress")
    return bool(task.get("actualFinish")) or (
        isinstance(actual_progress, (int, float)) and float(actual_progress) >= 1
    )


def is_open_schedule_task(task: dict[str, Any]) -> bool:
    return not is_completed_task(task)


def schedule_contains_any(query_text: str, tokens: list[str]) -> bool:
    return any(token in query_text for token in tokens)


def schedule_compact_text(query_text: str) -> str:
    return re.sub(r"\s+", "", query_text)


def schedule_requests_completed_history(query_text: str) -> bool:
    compact = schedule_compact_text(query_text)
    return (
        schedule_contains_any(query_text, ["완료 이력 포함", "완료된 일정도 포함", "완료된 작업도 포함", "완료도 포함", "include completed", "include finished"])
        or ("포함" in compact and schedule_contains_any(compact, ["완료된일정", "완료된작업", "완료이력"]))
    )


def schedule_requests_remaining(query_text: str) -> bool:
    compact = schedule_compact_text(query_text)
    return schedule_contains_any(query_text, [
        "남은",
        "남아있는",
        "해야 할",
        "해야할",
        "할 일",
        "할일",
        "완료된 일정 제외",
        "완료된 작업 제외",
        "remaining",
        "todo",
        "to do",
    ]) or schedule_contains_any(compact, [
        "완료된일정제외",
        "완료된작업제외",
        "남은것",
        "남은작업",
        "남은일정",
        "해야할일",
        "이번주해야할일",
    ])


def schedule_requests_this_week(query_text: str) -> bool:
    compact = schedule_compact_text(query_text)
    return schedule_contains_any(query_text, ["이번 주", "이번주", "this week"]) or "이번주" in compact


def schedule_requests_risk(query_text: str) -> bool:
    return schedule_contains_any(query_text, ["리스크", "위험", "risk"])


def task_overlaps_period(task: dict[str, Any], start_date: str, finish_date: str) -> bool:
    planned_start = task.get("plannedStart")
    planned_finish = task.get("plannedFinish")
    if planned_finish and planned_finish < start_date:
        return False
    if planned_start and planned_start > finish_date:
        return False
    return bool(planned_start or planned_finish)


def schedule_open_sort_key(task: dict[str, Any], as_of_date: str) -> tuple[int, str, str, str]:
    delayed_rank = 0 if is_delayed_task(task, as_of_date) else 1
    finish = cell_to_text(task.get("plannedFinish")) or "9999-12-31"
    start = cell_to_text(task.get("plannedStart")) or "9999-12-31"
    return (delayed_rank, finish, start, cell_to_text(task.get("wbs")))


def is_not_started_task(task: dict[str, Any]) -> bool:
    actual_progress = task.get("actualProgress")
    has_progress = isinstance(actual_progress, (int, float)) and float(actual_progress) > 0
    return not task.get("actualStart") and not has_progress and not task.get("actualFinish")


def schedule_task_status(task: dict[str, Any], as_of_date: str) -> str:
    if is_completed_task(task):
        return "completed"
    if is_delayed_task(task, as_of_date):
        return "delayed"
    if is_active_task(task, as_of_date):
        return "active"
    if is_not_started_task(task):
        return "not_started"
    return "in_progress"


def matching_resource_allocation(task: dict[str, Any], target: str | None) -> float | None:
    if not target:
        return None
    normalized_target = normalize_schedule_lookup_text(target)
    for resource in task.get("resource", []):
        if not isinstance(resource, dict):
            continue
        if normalize_schedule_lookup_text(resource.get("name")) == normalized_target:
            allocation = resource.get("allocation")
            return float(allocation) if isinstance(allocation, (int, float)) else None
    return None


def analyze_schedule_tasks(
    tasks: list[dict[str, Any]],
    as_of_date: str,
    resource_target: str | None = None,
) -> list[dict[str, Any]]:
    analyses: list[dict[str, Any]] = []
    for task in tasks:
        is_completed = is_completed_task(task)
        is_delayed = bool(task.get("isLeaf")) and is_delayed_task(task, as_of_date)
        is_active = is_active_task(task, as_of_date)
        analyses.append({
            "taskId": task.get("taskId"),
            "wbs": task.get("wbs"),
            "name": task.get("name"),
            "allocation": matching_resource_allocation(task, resource_target),
            "plannedStart": task.get("plannedStart"),
            "plannedFinish": task.get("plannedFinish"),
            "actualStart": task.get("actualStart"),
            "actualFinish": task.get("actualFinish"),
            "actualProgress": task.get("actualProgress"),
            "isActive": is_active,
            "isDelayed": is_delayed,
            "isCompleted": is_completed,
            "status": schedule_task_status(task, as_of_date),
        })
    return analyses


def summarize_task_analyses(analyses: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "totalTasks": len(analyses),
        "completedTasks": sum(1 for item in analyses if item.get("status") == "completed"),
        "activeTasks": sum(1 for item in analyses if item.get("status") in {"active", "in_progress"}),
        "delayedTasks": sum(1 for item in analyses if item.get("status") == "delayed"),
        "notStartedTasks": sum(1 for item in analyses if item.get("status") == "not_started"),
    }


def parse_iso_date(value: Any) -> datetime | None:
    text = cell_to_text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text[:10])
    except ValueError:
        return None


def date_text(value: datetime | None) -> str | None:
    return value.date().isoformat() if value else None


def calendar_day_types(schedule: dict[str, Any]) -> dict[str, str]:
    return {
        day["date"]: day.get("type", "working")
        for day in schedule.get("calendars", [])
        if isinstance(day, dict) and day.get("date")
    }


def is_working_day(day: datetime, schedule: dict[str, Any]) -> bool:
    calendar_types = calendar_day_types(schedule)
    day_type = calendar_types.get(day.date().isoformat())
    if day_type in {"holiday", "non-working"}:
        return False
    if day_type == "working":
        return True
    return day.weekday() < 5


def working_days_between(start: datetime, finish: datetime, schedule: dict[str, Any]) -> int:
    if finish < start:
        return -working_days_between(finish, start, schedule)
    count = 0
    current = start
    while current <= finish:
        if is_working_day(current, schedule):
            count += 1
        current += timedelta(days=1)
    return count


def add_working_days(start: datetime, days: int, schedule: dict[str, Any]) -> datetime:
    if days <= 1:
        return start
    current = start
    remaining = days - 1
    while remaining > 0:
        current += timedelta(days=1)
        if is_working_day(current, schedule):
            remaining -= 1
    return current


def elapsed_working_day_index(start: datetime, target: datetime, schedule: dict[str, Any]) -> int:
    return max(1, working_days_between(start, target, schedule))


def progress_points_until(schedule: dict[str, Any], as_of_date: str) -> list[dict[str, Any]]:
    points = [
        point
        for point in schedule.get("progressSeries", [])
        if isinstance(point, dict) and point.get("date") and point["date"] <= as_of_date
    ]
    return sorted(points, key=lambda point: point["date"])


def interpolate_earned_schedule_date(
    progress_series: list[dict[str, Any]],
    actual_progress: float,
    schedule: dict[str, Any],
    project_start: datetime,
) -> datetime | None:
    if not progress_series:
        return None
    previous = progress_series[0]
    first_progress = previous.get("plannedProgress")
    if isinstance(first_progress, (int, float)) and actual_progress <= float(first_progress):
        return parse_iso_date(previous.get("date"))
    for point in progress_series[1:]:
        previous_progress = previous.get("plannedProgress")
        current_progress = point.get("plannedProgress")
        if not isinstance(previous_progress, (int, float)) or not isinstance(current_progress, (int, float)):
            previous = point
            continue
        if float(current_progress) < actual_progress:
            previous = point
            continue
        previous_date = parse_iso_date(previous.get("date"))
        current_date = parse_iso_date(point.get("date"))
        if not previous_date or not current_date:
            return parse_iso_date(point.get("date"))
        denominator = float(current_progress) - float(previous_progress)
        if denominator <= 0:
            return current_date
        ratio = (actual_progress - float(previous_progress)) / denominator
        previous_index = elapsed_working_day_index(project_start, previous_date, schedule)
        current_index = elapsed_working_day_index(project_start, current_date, schedule)
        earned_index = max(1, round(previous_index + ((current_index - previous_index) * ratio)))
        return add_working_days(project_start, earned_index, schedule)
    return parse_iso_date(progress_series[-1].get("date"))


def calculate_schedule_performance(schedule: dict[str, Any], source: dict[str, Any], as_of_date: str) -> dict[str, Any]:
    summary = schedule_summary(schedule, source, as_of_date)
    warnings: list[str] = []
    if not schedule.get("calendars"):
        warnings.append("calendar_incomplete")
    points = progress_points_until(schedule, as_of_date) or schedule.get("progressSeries", [])
    last_progress = points[-1] if points else {}
    actual_progress = last_progress.get("actualProgress")
    planned_progress = summary.get("plannedProgress")
    project_start = parse_iso_date(schedule.get("projectStart"))
    planned_finish = parse_iso_date(schedule.get("projectFinish"))
    if not isinstance(actual_progress, (int, float)) or not project_start:
        return {
            "asOfDate": as_of_date,
            "plannedProgress": planned_progress,
            "actualProgress": actual_progress if isinstance(actual_progress, (int, float)) else None,
            "warnings": warnings + ["insufficient_progress_history"],
            "available": False,
        }
    earned_date = interpolate_earned_schedule_date(schedule.get("progressSeries", []), float(actual_progress), schedule, project_start)
    as_of = parse_iso_date(as_of_date) or parse_iso_date(last_progress.get("date")) or datetime.now()
    earned_working_days = working_days_between(project_start, earned_date, schedule) if earned_date else None
    actual_working_days = working_days_between(project_start, as_of, schedule)
    sv = (earned_working_days - actual_working_days) if earned_working_days is not None else None
    spi = (earned_working_days / actual_working_days) if earned_working_days is not None and actual_working_days > 0 else None
    planned_total_working_days = (
        working_days_between(project_start, planned_finish, schedule)
        if planned_finish
        else None
    )
    return {
        "available": True,
        "asOfDate": as_of_date,
        "plannedProgress": planned_progress,
        "actualProgress": float(actual_progress),
        "earnedScheduleDate": date_text(earned_date),
        "actualTimeDate": as_of.date().isoformat(),
        "earnedScheduleDays": earned_working_days,
        "actualTimeDays": actual_working_days,
        "scheduleVarianceDays": sv,
        "earnedWorkingDays": earned_working_days,
        "actualWorkingDays": actual_working_days,
        "scheduleVarianceWorkingDays": sv,
        "schedulePerformanceIndex": spi,
        "plannedFinish": date_text(planned_finish),
        "plannedTotalWorkingDays": planned_total_working_days,
        "warnings": warnings,
    }


def forecast_from_earned_schedule(performance: dict[str, Any], schedule: dict[str, Any]) -> dict[str, Any]:
    spi = performance.get("schedulePerformanceIndex")
    total_days = performance.get("plannedTotalWorkingDays")
    project_start = parse_iso_date(schedule.get("projectStart"))
    if not isinstance(spi, (int, float)) or spi <= 0 or not isinstance(total_days, int) or not project_start:
        return {
            "method": "earned_schedule",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "SPI(t) or planned working-day duration is unavailable.",
            "quality": "insufficient",
            "confidence": "low",
            "dataQuality": "insufficient",
            "assumptions": [
                "Requires SPI(t) and planned working-day duration.",
            ],
            "warning": "earned_schedule_unavailable",
        }
    estimated_days = max(1, math.ceil(total_days / float(spi)))
    estimated_finish = date_text(add_working_days(project_start, estimated_days, schedule))
    return {
        "method": "earned_schedule",
        "available": True,
        "estimate": estimated_finish,
        "estimatedFinish": estimated_finish,
        "basis": f"planned working days / SPI(t) = {total_days} / {float(spi):.4f}",
        "reason": "If cumulative SPI(t) stays at the current level, this is the extrapolated finish scenario.",
        "quality": "limited",
        "confidence": "low",
        "dataQuality": "limited",
        "assumptions": [
            f"Current SPI(t) = {float(spi):.4f}.",
            "Cumulative schedule performance remains unchanged.",
            "This is a performance scenario, not the primary operational estimate.",
        ],
    }


def forecast_from_recent_velocity(schedule: dict[str, Any], as_of_date: str) -> dict[str, Any]:
    points = progress_points_until(schedule, as_of_date) or schedule.get("progressSeries", [])
    if len(points) < 2:
        return {
            "method": "recent_velocity",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "Insufficient progress history.",
            "quality": "insufficient",
            "confidence": "low",
            "dataQuality": "insufficient",
            "assumptions": [
                "Requires at least two progress points in the recent window.",
            ],
            "warning": "insufficient_progress_history",
        }
    end = points[-1]
    end_date = parse_iso_date(end.get("date"))
    end_progress = end.get("actualProgress")
    if not end_date or not isinstance(end_progress, (int, float)):
        return {
            "method": "recent_velocity",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "Actual progress is unavailable at the end of the recent window.",
            "quality": "insufficient",
            "confidence": "low",
            "dataQuality": "insufficient",
            "assumptions": [
                "Requires actual progress at the end of the recent window.",
            ],
            "warning": "insufficient_progress_history",
        }
    window_start = end_date - timedelta(days=28)
    candidates = [point for point in points if parse_iso_date(point.get("date")) and parse_iso_date(point.get("date")) >= window_start]
    start = candidates[0] if candidates else points[0]
    start_date = parse_iso_date(start.get("date"))
    start_progress = start.get("actualProgress")
    if not start_date or not isinstance(start_progress, (int, float)):
        return {
            "method": "recent_velocity",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "Actual progress is unavailable at the start of the recent window.",
            "quality": "insufficient",
            "confidence": "low",
            "dataQuality": "insufficient",
            "assumptions": [
                "Requires actual progress at the start of the recent window.",
            ],
            "warning": "insufficient_progress_history",
        }
    elapsed_days = max(1, working_days_between(start_date, end_date, schedule))
    velocity = (float(end_progress) - float(start_progress)) / elapsed_days
    if velocity <= 0:
        return {
            "method": "recent_velocity",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "No positive actual progress in the recent 4-week window.",
            "basis": "No positive actual progress in the recent 4-week window.",
            "quality": "stale",
            "confidence": "low",
            "dataQuality": "stale",
            "assumptions": [
                "Recent progress velocity uses the latest 4-week window.",
                "A positive actual progress delta is required.",
            ],
            "warning": "recent_velocity_unavailable",
        }
    remaining_days = math.ceil(max(0, 1 - float(end_progress)) / velocity)
    as_of = parse_iso_date(as_of_date) or end_date
    estimated_finish = date_text(add_working_days(as_of, remaining_days, schedule))
    return {
        "method": "recent_velocity",
        "available": True,
        "estimate": estimated_finish,
        "estimatedFinish": estimated_finish,
        "basis": f"recent velocity {velocity:.6f} progress per working day",
        "reason": "Projected from the most recent 4-week actual progress velocity.",
        "quality": "recent_velocity",
        "confidence": "medium",
        "dataQuality": "recent_velocity",
        "assumptions": [
            "Recent 4-week progress velocity continues unchanged.",
        ],
    }


def forecast_from_remaining_tasks(schedule: dict[str, Any], as_of_date: str) -> dict[str, Any]:
    incomplete_leaf_tasks = [
        task
        for task in schedule.get("tasks", [])
        if task.get("isLeaf") and not is_completed_task(task)
    ]
    if not incomplete_leaf_tasks:
        return {
            "method": "remaining_tasks",
            "available": True,
            "estimate": as_of_date,
            "estimatedFinish": as_of_date,
            "basis": "All leaf tasks are complete.",
            "reason": "No remaining leaf tasks.",
            "quality": "complete",
            "confidence": "medium",
            "dataQuality": "complete_leaf_tasks",
            "assumptions": [
                "All leaf tasks are complete.",
            ],
        }
    remaining_durations = [
        max(1, math.ceil(float(task.get("plannedDuration")) * (1 - float(task.get("actualProgress") or 0))))
        for task in incomplete_leaf_tasks
        if isinstance(task.get("plannedDuration"), (int, float)) and float(task.get("plannedDuration")) > 0
    ]
    as_of = parse_iso_date(as_of_date)
    if not remaining_durations or not as_of:
        return {
            "method": "remaining_tasks",
            "available": False,
            "estimate": None,
            "estimatedFinish": None,
            "reason": "Remaining leaf task duration or as-of date is unavailable.",
            "quality": "insufficient",
            "confidence": "low",
            "dataQuality": "insufficient",
            "assumptions": [
                "Requires planned duration for incomplete leaf tasks.",
            ],
            "warning": "remaining_tasks_unavailable",
        }
    estimated = add_working_days(as_of, max(remaining_durations), schedule)
    estimated_finish = date_text(estimated)
    return {
        "method": "remaining_tasks",
        "available": True,
        "estimate": estimated_finish,
        "estimatedFinish": estimated_finish,
        "basis": f"max remaining leaf duration = {max(remaining_durations)} working days across {len(incomplete_leaf_tasks)} incomplete leaf tasks",
        "reason": "Heuristic based on remaining duration of incomplete leaf tasks; dependencies are not inferred.",
        "quality": "heuristic",
        "confidence": "low",
        "dataQuality": "heuristic",
        "assumptions": [
            "Remaining leaf tasks use planned duration multiplied by remaining progress.",
            "Dependency relationships unavailable.",
            "Resource leveling and dependency propagation are not applied.",
        ],
    }


def calculate_schedule_forecast(schedule: dict[str, Any], source: dict[str, Any], as_of_date: str) -> dict[str, Any]:
    performance = calculate_schedule_performance(schedule, source, as_of_date)
    methods = [
        forecast_from_earned_schedule(performance, schedule),
        forecast_from_recent_velocity(schedule, as_of_date),
        forecast_from_remaining_tasks(schedule, as_of_date),
    ]
    methods_by_name = {method.get("method"): method for method in methods}
    remaining_method = methods_by_name.get("remaining_tasks", {})
    earned_method = methods_by_name.get("earned_schedule", {})
    available_methods = [method for method in methods if method.get("available")]
    unavailable_methods = [method for method in methods if not method.get("available")]
    primary_methods = [
        method
        for method in [remaining_method]
        if method.get("available") and method.get("estimatedFinish")
    ]
    primary_finish_dates = [method["estimatedFinish"] for method in primary_methods]
    all_finish_dates = [method.get("estimatedFinish") for method in available_methods if method.get("estimatedFinish")]
    warnings: list[str] = []
    if not available_methods:
        warnings.append("forecast_unavailable")
    elif unavailable_methods:
        warnings.append("forecast_partially_unavailable")
    for method in unavailable_methods:
        warning = method.get("warning")
        if warning:
            warnings.append(str(warning))
    if not schedule.get("dependencies"):
        warnings.append("dependency_relationships_unavailable")
    primary_estimate = primary_finish_dates[0] if primary_finish_dates else None
    return {
        "asOfDate": as_of_date,
        "plannedFinish": performance.get("plannedFinish") or schedule.get("projectFinish"),
        "primaryEstimate": primary_estimate,
        "primaryMethod": "remaining_tasks" if primary_estimate else None,
        "primaryReason": (
            "Remaining-task heuristic is used as the primary operational estimate because dependency relationships are unavailable and recent velocity is unavailable."
            if primary_estimate
            else "No primary operational estimate is available."
        ),
        "primaryAssumptions": [
            "Remaining-task heuristic.",
            "Dependency relationships unavailable.",
            "Resource leveling and dependency propagation are not applied.",
        ] if primary_estimate else [],
        "primaryForecastRange": {
            "earliest": min(primary_finish_dates) if primary_finish_dates else None,
            "latest": max(primary_finish_dates) if primary_finish_dates else None,
        },
        "performanceScenario": {
            "method": "earned_schedule",
            "estimate": earned_method.get("estimate"),
            "available": bool(earned_method.get("available")),
            "quality": earned_method.get("quality") or earned_method.get("dataQuality"),
            "reason": earned_method.get("reason"),
            "assumptions": earned_method.get("assumptions") or [],
            "spi": performance.get("schedulePerformanceIndex"),
        },
        "methods": methods,
        "forecastRange": {
            "earliest": min(primary_finish_dates) if primary_finish_dates else None,
            "latest": max(primary_finish_dates) if primary_finish_dates else None,
        },
        "allScenarioRange": {
            "earliest": min(all_finish_dates) if all_finish_dates else None,
            "latest": max(all_finish_dates) if all_finish_dates else None,
        },
        "forecastQuality": "limited" if primary_estimate else "unavailable",
        "dependencyRelationshipsAvailable": bool(schedule.get("dependencies")),
        "warnings": list(dict.fromkeys(performance.get("warnings", []) + warnings)),
    }


def dependency_graph(schedule: dict[str, Any]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    predecessors: dict[str, list[dict[str, Any]]] = {}
    successors: dict[str, list[dict[str, Any]]] = {}
    for dependency in schedule.get("dependencies", []):
        predecessors.setdefault(dependency["successor_wbs"], []).append(dependency)
        successors.setdefault(dependency["predecessor_wbs"], []).append(dependency)
    return predecessors, successors


def dependency_chain(start_wbs: str, edges: dict[str, list[dict[str, Any]]], direction: str) -> tuple[list[dict[str, Any]], list[str]]:
    chain: list[dict[str, Any]] = []
    warnings: list[str] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def walk(wbs: str, depth: int) -> None:
        if wbs in visiting:
            warnings.append("dependency_cycle")
            return
        if wbs in visited:
            return
        visiting.add(wbs)
        for dependency in edges.get(wbs, []):
            next_wbs = dependency["successor_wbs"] if direction == "downstream" else dependency["predecessor_wbs"]
            chain.append({**dependency, "depth": depth})
            walk(next_wbs, depth + 1)
        visiting.remove(wbs)
        visited.add(wbs)

    walk(start_wbs, 1)
    return chain, warnings


def analyze_dependencies(schedule: dict[str, Any], target_wbs: str | None = None) -> dict[str, Any]:
    predecessors, successors = dependency_graph(schedule)
    warnings: list[str] = []
    if not schedule.get("dependencies"):
        warnings.append("dependency_not_available")
    upstream: list[dict[str, Any]] = []
    downstream: list[dict[str, Any]] = []
    if target_wbs:
        upstream, upstream_warnings = dependency_chain(target_wbs, predecessors, "upstream")
        downstream, downstream_warnings = dependency_chain(target_wbs, successors, "downstream")
        warnings.extend(upstream_warnings + downstream_warnings)
    return {
        "dependencyCount": len(schedule.get("dependencies", [])),
        "targetWbs": target_wbs,
        "predecessors": predecessors.get(target_wbs, []) if target_wbs else [],
        "successors": successors.get(target_wbs, []) if target_wbs else [],
        "upstreamChain": upstream,
        "downstreamChain": downstream,
        "warnings": list(dict.fromkeys(warnings)),
    }


def extract_delay_days(query: str) -> int | None:
    query_text = cell_to_text(query).lower()
    if re.search(r"(?:일\s*주일|일주일|한\s*주|1\s*주)", query_text):
        return 5
    week_match = re.search(r"(\d+)\s*주", query_text)
    if week_match:
        return int(week_match.group(1)) * 5
    working_day_match = re.search(
        r"(\d+)\s*(?:영업\s*일|영업일|working\s*days?|business\s*days?)",
        query_text,
        re.IGNORECASE,
    )
    if working_day_match:
        return int(working_day_match.group(1))
    day_match = re.search(r"(\d+)\s*(?:일|days?|day)", query_text, re.IGNORECASE)
    return int(day_match.group(1)) if day_match else None


def query_mentions_what_if_scenario(query: str) -> bool:
    query_text = cell_to_text(query).lower()
    return any(
        token in query_text
        for token in [
            "늦어지",
            "밀리",
            "지연되면",
            "지연된다면",
            "delay by",
            "delayed by",
            "what-if",
            "what if",
        ]
    )


def calculate_what_if(schedule: dict[str, Any], target_wbs: str | None, delay_days: int | None) -> dict[str, Any]:
    warnings: list[str] = []
    if not target_wbs or not delay_days:
        warnings.append("what_if_unavailable")
    dependency = analyze_dependencies(schedule, target_wbs)
    if dependency["dependencyCount"] == 0:
        warnings.append("what_if_dependency_missing")
    task_by_wbs = {task.get("wbs"): task for task in schedule.get("tasks", [])}
    target_task = task_by_wbs.get(target_wbs)
    planned_finish = parse_iso_date(target_task.get("plannedFinish")) if target_task else None
    shifted_finish = add_working_days(planned_finish, delay_days + 1, schedule) if planned_finish and delay_days else None
    project_finish = parse_iso_date(schedule.get("projectFinish"))
    simulated_finish_exceeds_project_finish = bool(
        shifted_finish and project_finish and shifted_finish > project_finish
    )
    if simulated_finish_exceeds_project_finish:
        warnings.append("simulated_task_finish_after_project_finish")
    dependency_propagation = "available" if dependency["dependencyCount"] > 0 else "unavailable"
    return {
        "intent": "what_if_task_delay",
        "task": target_task,
        "targetWbs": target_wbs,
        "delayWorkingDays": delay_days,
        "delayUnit": "working_day",
        "delayAssumption": "Schedule what-if delays are interpreted as working days.",
        "originalFinishBasis": "planned_finish",
        "targetOriginalFinish": date_text(planned_finish),
        "targetWhatIfFinish": date_text(shifted_finish),
        "simulatedFinish": date_text(shifted_finish),
        "dependencyPropagation": dependency_propagation,
        "directImpact": dependency.get("successors", []),
        "indirectImpact": dependency.get("downstreamChain", []),
        "projectOriginalFinish": schedule.get("projectFinish"),
        "projectWhatIfFinish": None if dependency["dependencyCount"] == 0 else date_text(shifted_finish),
        "projectFinishImpact": None if dependency["dependencyCount"] == 0 else {
            "projectOriginalFinish": schedule.get("projectFinish"),
            "projectWhatIfFinish": date_text(shifted_finish),
        },
        "projectFinishComparison": {
            "projectPlannedFinish": schedule.get("projectFinish"),
            "simulatedTaskFinish": date_text(shifted_finish),
            "simulatedTaskFinishExceedsProjectFinish": simulated_finish_exceeds_project_finish,
        },
        "warnings": list(dict.fromkeys(warnings + dependency.get("warnings", []))),
    }


def schedule_analysis_context_text(kind: str, analysis: dict[str, Any] | None) -> str:
    if not analysis:
        return ""
    if kind == "forecast":
        performance_scenario = analysis.get("performanceScenario")
        scenario = performance_scenario if isinstance(performance_scenario, dict) else {}
        recent_velocity_method = next(
            (
                method
                for method in analysis.get("methods", [])
                if isinstance(method, dict)
                and method.get("method") == "recent_velocity"
            ),
            None,
        )
        recent_velocity_estimate = (
            recent_velocity_method.get("estimate")
            if recent_velocity_method
            else None
        )
        return "\n".join([
            "[SCHEDULE ANALYSIS]",
            "Type: forecast",
            f"Primary Operational Estimate: {analysis.get('primaryEstimate') or 'unavailable'}",
            f"Primary Method: {analysis.get('primaryMethod') or 'unavailable'}",
            f"Primary Reason: {analysis.get('primaryReason') or '-'}",
            f"Earned Schedule Scenario: {scenario.get('estimate') or 'unavailable'}",
            f"Earned Schedule Scenario Quality: {scenario.get('quality') or '-'}",
            "Do not describe the Earned Schedule Scenario as a committed or primary finish date.",
            f"Recent Velocity: {recent_velocity_estimate or 'unavailable'}",
            json.dumps(analysis, ensure_ascii=False, indent=2),
            "[/SCHEDULE ANALYSIS]",
        ])
    if kind == "what_if":
        task = analysis.get("task")
        task_record = task if isinstance(task, dict) else {}
        return "\n".join([
            "[SCHEDULE ANALYSIS]",
            "Type: what_if_task_delay",
            f"Task: {task_record.get('name') or 'unavailable'}",
            f"WBS: {analysis.get('targetWbs') or 'unavailable'}",
            f"Original Finish: {analysis.get('targetOriginalFinish') or 'unavailable'}",
            f"Delay: {analysis.get('delayWorkingDays') or 'unavailable'} working days",
            "Delay Assumption: Schedule what-if delays are interpreted as working days.",
            f"Simulated Finish: {analysis.get('targetWhatIfFinish') or 'unavailable'}",
            f"Dependency Propagation: {analysis.get('dependencyPropagation') or 'unavailable'}",
            f"Project Finish Impact: {analysis.get('projectWhatIfFinish') or 'unavailable'}",
            "Source workbook is not modified by this simulation.",
            "Do not present this what-if simulation as a committed schedule change.",
            json.dumps(analysis, ensure_ascii=False, indent=2),
            "[/SCHEDULE ANALYSIS]",
        ])
    return "\n".join([
        "[SCHEDULE ANALYSIS]",
        f"Type: {kind}",
        json.dumps(analysis, ensure_ascii=False, indent=2),
        "[/SCHEDULE ANALYSIS]",
    ])


def query_requests_status(query: str) -> bool:
    query_text = query.lower()
    return any(
        token in query_text
        for token in [
            "진행",
            "진척",
            "진도",
            "상태",
            "지연",
            "지체",
            "늦",
            "밀",
            "잘",
            "status",
            "progress",
            "delay",
            "delayed",
            "active",
        ]
    )


def query_requests_dependency(query: str) -> bool:
    query_text = query.lower()
    return any(token in query_text for token in ["선행", "후행", "predecessor", "successor", "dependency", "dependencies"])


def query_requests_impact(query: str) -> bool:
    query_text = query.lower()
    return any(token in query_text for token in ["영향", "impact", "영향받", "영향 받아"])


def query_requests_earned_schedule(query: str) -> bool:
    query_text = query.lower()
    return any(token in query_text for token in ["일정 성과", "earned schedule", "spi", "sv(t)", "sv"])


def query_requests_schedule_performance(query: str) -> bool:
    query_text = query.lower()
    return query_requests_earned_schedule(query) or any(
        token in query_text
        for token in [
            "일정 성과",
            "성과",
            "계획보다",
            "얼마나 밀렸",
            "얼마나 지연",
            "spi",
            "sv(t)",
            "earned schedule",
        ]
    )


def query_requests_forecast(query: str) -> bool:
    query_text = query.lower()
    if any(
        token in query_text
        for token in [
            "언제 끝",
            "언제 완료",
            "예상",
            "전망",
            "추세",
            "계획 종료일",
            "지킬 수",
            "완료 가능",
        ]
    ):
        return True
    return any(token in query_text for token in ["언제 끝", "언제 완료", "예상", "전망", "추세", "forecast", "estimate"])


def query_requests_what_if(query: str) -> bool:
    return query_mentions_what_if_scenario(query) and extract_delay_days(query) is not None


def schedule_query(input_data: dict[str, Any]) -> dict[str, Any]:
    schedule = input_data.get("schedule")
    source = input_data.get("source")
    query = input_data.get("query")
    as_of_date = input_data.get("as_of_date") or datetime.now().date().isoformat()
    if not isinstance(schedule, dict) or not isinstance(source, dict) or not isinstance(query, str):
        raise WorkerError("invalid_input", "schedule, source, and query are required.")

    query_text = query.lower()
    tasks = schedule.get("tasks", [])
    selected = tasks
    kind = "summary"
    target = None
    detected_entity_type = None
    detected_entity = None
    task_analyses: list[dict[str, Any]] = []
    resource_status_summary = None
    advanced_analysis = None

    resource_entity = find_resource_entity(query, tasks)
    task_entity_type, task_entity = find_task_entity(query, tasks)
    what_if_requested = query_mentions_what_if_scenario(query)
    delay_days = extract_delay_days(query)
    include_completed_history = schedule_requests_completed_history(query_text)

    if task_entity and what_if_requested and delay_days is not None:
        target = task_entity
        detected_entity_type = task_entity_type
        detected_entity = task_entity
        selected = lookup_tasks_by_name_or_wbs(tasks, task_entity)
        kind = "what_if"
        advanced_analysis = calculate_what_if(
            schedule,
            selected[0].get("wbs") if selected else None,
            delay_days,
        )
    elif what_if_requested:
        kind = "unsupported"
        selected = []
        advanced_analysis = {
            "intent": "what_if_task_delay",
            "reason": "task_not_found" if not task_entity else "delay_days_missing",
            "target": task_entity,
            "delayWorkingDays": delay_days,
            "warnings": [
                "task_not_found" if not task_entity else "what_if_delay_missing"
            ],
        }
    elif task_entity and query_requests_impact(query):
        target = task_entity
        detected_entity_type = task_entity_type
        detected_entity = task_entity
        selected = lookup_tasks_by_name_or_wbs(tasks, task_entity)
        kind = "impact_analysis"
        advanced_analysis = analyze_dependencies(
            schedule,
            selected[0].get("wbs") if selected else None,
        )
    elif task_entity and query_requests_dependency(query):
        target = task_entity
        detected_entity_type = task_entity_type
        detected_entity = task_entity
        selected = lookup_tasks_by_name_or_wbs(tasks, task_entity)
        kind = "dependency_lookup"
        advanced_analysis = analyze_dependencies(
            schedule,
            selected[0].get("wbs") if selected else None,
        )
    elif resource_entity:
        target = resource_entity
        detected_entity_type = "resource"
        detected_entity = resource_entity
        selected = [task for task in tasks if resource_matches_target(task, resource_entity)]
        if query_requests_status(query):
            kind = "resource_status"
            task_analyses = analyze_schedule_tasks(selected, as_of_date, resource_entity)
            resource_status_summary = summarize_task_analyses(task_analyses)
        else:
            kind = "resource_lookup"
    elif task_entity:
        target = task_entity
        detected_entity_type = task_entity_type
        detected_entity = task_entity
        selected = lookup_tasks_by_name_or_wbs(tasks, task_entity)
        kind = "task_status" if query_requests_status(query) else "task_lookup"
        if kind == "task_status":
            task_analyses = analyze_schedule_tasks(selected, as_of_date)
    elif query_requests_schedule_performance(query):
        kind = "schedule_performance"
        selected = []
        advanced_analysis = calculate_schedule_performance(schedule, source, as_of_date)
    elif query_requests_forecast(query):
        kind = "forecast"
        selected = []
        advanced_analysis = calculate_schedule_forecast(schedule, source, as_of_date)
    elif schedule_requests_this_week(query_text) and schedule_requests_remaining(query_text):
        as_of = parse_iso_date(as_of_date)
        week_finish = date_text(as_of + timedelta(days=6)) if as_of else as_of_date
        selected = [
            task
            for task in tasks
            if task.get("isLeaf")
            and (include_completed_history or is_open_schedule_task(task))
            and task_overlaps_period(task, as_of_date, week_finish or as_of_date)
        ]
        selected = sorted(selected, key=lambda task: schedule_open_sort_key(task, as_of_date))
        kind = "remaining_tasks"
    elif schedule_requests_remaining(query_text):
        selected = [
            task
            for task in tasks
            if task.get("isLeaf") and (include_completed_history or is_open_schedule_task(task))
        ]
        selected = sorted(selected, key=lambda task: schedule_open_sort_key(task, as_of_date))
        kind = "remaining_tasks"
    elif schedule_requests_risk(query_text):
        selected = [
            task
            for task in tasks
            if task.get("isLeaf")
            and is_open_schedule_task(task)
            and (
                is_delayed_task(task, as_of_date)
                or is_active_task(task, as_of_date)
                or not task.get("plannedFinish")
            )
        ]
        selected = sorted(selected, key=lambda task: schedule_open_sort_key(task, as_of_date))
        kind = "delayed_tasks"
    elif any(token in query_text for token in ["지연", "지체", "delay", "delayed"]):
        selected = [
            task
            for task in tasks
            if task.get("isLeaf")
            and (include_completed_history or is_open_schedule_task(task))
            and is_delayed_task(task, as_of_date)
        ]
        selected = sorted(selected, key=lambda task: schedule_open_sort_key(task, as_of_date))
        kind = "delayed_tasks"
    elif any(token in query_text for token in ["진행", "active", "현재", "착수 중"]):
        selected = [
            task for task in tasks
            if task.get("isLeaf")
            and (include_completed_history or is_open_schedule_task(task))
            and is_active_task(task, as_of_date)
        ]
        selected = sorted(selected, key=lambda task: schedule_open_sort_key(task, as_of_date))
        kind = "active_tasks"
    elif any(token in query_text for token in ["완료", "finish", "finishing", "종료"]):
        selected = [
            task
            for task in tasks
            if task.get("plannedFinish") and task["plannedFinish"] >= as_of_date
        ]
        selected = sorted(selected, key=lambda task: task.get("plannedFinish") or "")[:20]
        kind = "finishing_between"
    elif any(token in query_text for token in ["착수", "start", "starting"]):
        selected = [
            task
            for task in tasks
            if task.get("plannedStart") and task["plannedStart"] >= as_of_date
        ]
        selected = sorted(selected, key=lambda task: task.get("plannedStart") or "")[:20]
        kind = "starting_between"
    elif any(token in query_text for token in ["담당", "resource", "리소스"]):
        target = extract_resource_lookup_target(query)
        selected = (
            [task for task in tasks if resource_matches_target(task, target)]
            if target
            else [task for task in tasks if task.get("resource")]
        )
        kind = "resource_lookup"
    elif "wbs" in query_text or "작업" in query_text or "일정" in query_text:
        target = extract_task_lookup_target(query)
        selected = lookup_tasks_by_name_or_wbs(tasks, target) if target else []
        kind = "task_lookup"

    summary = schedule_summary(schedule, source, as_of_date)
    if not task_analyses and kind in {"resource_lookup", "task_lookup"}:
        task_analyses = analyze_schedule_tasks(
            selected[:20],
            as_of_date,
            target if kind == "resource_lookup" else None,
        )
    return {
        "kind": kind,
        "target": target,
        "detectedEntityType": detected_entity_type,
        "detectedEntity": detected_entity,
        "detectedIntent": kind,
        "resourceStatusSummary": resource_status_summary,
        "taskAnalyses": task_analyses[:20],
        "advancedAnalysis": advanced_analysis,
        "workspaceId": schedule["workspaceId"],
        "filename": schedule["filename"],
        "lastParsedAt": summary["lastParsedAt"],
        "asOfDate": as_of_date,
        "summary": summary,
        "tasks": selected[:20],
        "contextText": (
            schedule_context_text(kind, summary, selected)
            + (
                "\n" + schedule_analysis_context_text(kind, advanced_analysis)
                if advanced_analysis
                else ""
            )
        ),
    }


def schedule_summary_command(input_data: dict[str, Any]) -> dict[str, Any]:
    schedule = input_data.get("schedule")
    source = input_data.get("source")
    if not isinstance(schedule, dict) or not isinstance(source, dict):
        raise WorkerError("invalid_input", "schedule and source are required.")
    return schedule_summary(schedule, source, input_data.get("as_of_date"))


def percent_text(value: Any) -> str:
    return f"{float(value) * 100:.2f}%" if isinstance(value, (int, float)) and math.isfinite(float(value)) else "-"


def report_date_text(value: Any) -> str:
    text = cell_to_text(value)
    return text[:10] if text else "-"


def add_report_paragraphs(document: Any, text: str) -> None:
    for line in normalize_text(text).splitlines():
        if line.strip():
            document.add_paragraph(line.strip())


def add_report_task_table(document: Any, tasks: list[dict[str, Any]]) -> None:
    if not tasks:
        document.add_paragraph("해당 항목 없음")
        return
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["WBS", "작업", "계획 기간", "실적", "상태"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for task in tasks:
        cells = table.add_row().cells
        cells[0].text = cell_to_text(task.get("wbs"))
        cells[1].text = cell_to_text(task.get("name"))
        cells[2].text = f"{report_date_text(task.get('planned_start'))} ~ {report_date_text(task.get('planned_finish'))}"
        cells[3].text = percent_text(task.get("actual_progress"))
        cells[4].text = cell_to_text(task.get("status")) or "-"


def add_report_source_list(document: Any, title: str, sources: list[dict[str, Any]]) -> None:
    document.add_heading(title, level=2)
    if not sources:
        document.add_paragraph("관련 근거 없음")
        return
    for source in sources:
        paragraph = document.add_paragraph(style=None)
        paragraph.add_run(cell_to_text(source.get("filename")) or "source").bold = True
        if source.get("rag_document_id"):
            paragraph.add_run(f" / {cell_to_text(source.get('rag_document_id'))}")
        excerpt = cell_to_text(source.get("excerpt"))
        if excerpt:
            document.add_paragraph(excerpt[:700])


def weekly_report_render(input_data: dict[str, Any]) -> dict[str, Any]:
    report_data = input_data.get("report_data")
    output_path_text = cell_to_text(input_data.get("output_path"))
    template_path_text = cell_to_text(input_data.get("template_path"))
    if not isinstance(report_data, dict) or not output_path_text:
        raise WorkerError("invalid_input", "report_data and output_path are required.")
    output_path = Path(output_path_text)
    if output_path.suffix.lower() != ".docx":
        raise WorkerError("output_write_failed", "Weekly report output must be a .docx file.")
    try:
        from docx import Document
    except ImportError as error:
        raise WorkerError("docx_render_failed", "python-docx is required to render DOCX reports.") from error

    try:
        if template_path_text:
            template_path = Path(template_path_text)
            if not template_path.exists():
                raise WorkerError("template_missing", "Weekly report template was not found.")
            document = Document(str(template_path))
        else:
            document = Document()

        document.add_heading("Mimora 주간보고서 초안", level=0)
        document.add_paragraph(f"프로젝트: {cell_to_text(report_data.get('workspace_name'))}")
        report_period = report_data.get("report_period") if isinstance(report_data.get("report_period"), dict) else {}
        document.add_paragraph(
            f"보고기간: {report_date_text(report_period.get('start'))} ~ {report_date_text(report_period.get('end'))}"
        )
        document.add_paragraph(f"생성시각: {cell_to_text(report_data.get('generated_at'))}")

        schedule = report_data.get("schedule") if isinstance(report_data.get("schedule"), dict) else {}
        document.add_heading("1. 프로젝트 개요", level=1)
        document.add_paragraph(f"Schedule Source: {cell_to_text(schedule.get('source_filename'))}")
        document.add_paragraph(f"기준일: {report_date_text(schedule.get('as_of_date'))}")

        document.add_heading("2. 금주 수행내용", level=1)
        this_week = report_data.get("this_week") if isinstance(report_data.get("this_week"), dict) else {}
        add_report_task_table(document, this_week.get("completed") if isinstance(this_week.get("completed"), list) else [])
        add_report_paragraphs(document, "\n".join(this_week.get("key_activities") if isinstance(this_week.get("key_activities"), list) else []))

        document.add_heading("3. 일정/진척 현황", level=1)
        document.add_paragraph(f"계획 진척률: {percent_text(schedule.get('planned_progress'))}")
        document.add_paragraph(f"실적 진척률: {percent_text(schedule.get('actual_progress'))}")
        document.add_paragraph(f"지연 Leaf Task: {cell_to_text(schedule.get('delayed_task_count'))}개")
        add_report_task_table(document, schedule.get("delayed_tasks") if isinstance(schedule.get("delayed_tasks"), list) else [])

        add_report_source_list(document, "4. 주요 이슈", report_data.get("issues") if isinstance(report_data.get("issues"), list) else [])
        add_report_source_list(document, "5. 주요 리스크", report_data.get("risks") if isinstance(report_data.get("risks"), list) else [])
        add_report_source_list(document, "6. 주요 의사결정", report_data.get("decisions") if isinstance(report_data.get("decisions"), list) else [])

        document.add_heading("7. 차주 계획", level=1)
        next_week = report_data.get("next_week") if isinstance(report_data.get("next_week"), dict) else {}
        add_report_task_table(document, next_week.get("planned_tasks") if isinstance(next_week.get("planned_tasks"), list) else [])

        document.add_heading("8. Forecast / 주의사항", level=1)
        forecast = schedule.get("forecast") if isinstance(schedule.get("forecast"), dict) else {}
        document.add_paragraph(f"현실적 운영 추정: {report_date_text(forecast.get('primaryEstimate'))}")
        scenario = forecast.get("performanceScenario") if isinstance(forecast.get("performanceScenario"), dict) else {}
        document.add_paragraph(f"누적 일정성과 시나리오: {report_date_text(scenario.get('estimate'))}")
        security = report_data.get("security") if isinstance(report_data.get("security"), dict) else {}
        if security.get("contains_sensitive_context"):
            document.add_paragraph("민감정보 포함: 생성 근거에 민감/비공개 컨텍스트가 포함되어 있습니다.")

        summary = cell_to_text(report_data.get("summary"))
        if summary:
            document.add_heading("요약", level=1)
            add_report_paragraphs(document, summary)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
    except WorkerError:
        raise
    except Exception as error:
        raise WorkerError("docx_render_failed", "Weekly report DOCX rendering failed.") from error

    return {
        "canceled": False,
        "outputPath": str(output_path),
        "templatePath": template_path_text or None,
    }


ISSUE_COLUMN_ALIASES = {
    "risk_analysis_id": ["위험id전사분석용", "위험ID/(전사분석용)"],
    "item_type": ["구분", "유형", "type", "itemtype", "riskissue"],
    "item_id": ["이슈id", "이슈ID", "관리번호", "번호", "no", "id", "issueid"],
    "priority": ["우선순위", "priority"],
    "issue_area": ["이슈영역", "영역", "issuearea", "area", "domain"],
    "issue_category": ["이슈분류", "분류", "issuecategory", "category", "classification"],
    "phase": ["단계", "phase", "stage"],
    "title": ["이슈사건", "이슈 사건", "이슈명", "제목", "위험이슈명", "위험/이슈명", "issuename", "issuetitle", "issue", "title", "subject"],
    "description": ["상세내용", "내용", "위험이슈내용", "위험/이슈 내용", "현상", "description"],
    "reported_date": ["이슈발생일", "발생일", "등록일", "occurreddate", "issuedate", "opendate", "createddate"],
    "risk_id": ["위험id", "위험ID", "리스크id", "riskid"],
    "probability": ["발생확률", "발생 확률", "확률", "probability", "likelihood"],
    "impact": ["영향도", "영향", "impact"],
    "severity": ["위험도", "심각도", "리스크등급", "risklevel", "severity"],
    "root_cause": ["원인", "근본원인", "rootcause", "cause"],
    "action_plan": ["대응방안", "대응계획", "조치계획", "actionplan", "responseplan", "mitigation"],
    "resolution": ["조치내용", "해결내용", "완료내용", "resolution"],
    "owner": ["조치담당자", "조치/담당자", "담당자", "owner", "assignee", "pic"],
    "target_date": ["조치예정일", "조치/예정일", "완료예정일", "목표완료일", "목표일", "duedate", "targetdate", "planfinish"],
    "organization_support_required": ["조직지원필요여부", "조직지원필요", "조직지원", "지원필요", "organizationsupportrequired", "supportrequired", "orgsupport"],
    "status": ["진행상황", "진행상태", "상태", "status", "progressstatus", "state"],
    "effort_mh": ["처리공수mh", "처리공수", "공수", "effortmh", "mh", "manhour"],
    "resolved_date": ["조치일자", "완료일", "조치완료일", "해결일", "종료일", "completeddate", "closeddate", "actiondate", "finishdate"],
    "remarks": ["비고", "메모", "remarks", "memo", "notes"],
}

ISSUE_REQUIRED_COLUMN_GROUPS = [{"title", "description"}, {"status"}]
ISSUE_RISK_HINT_FIELDS = {"risk_analysis_id", "risk_id", "probability", "impact", "severity"}


def issue_source_path(input_data: dict[str, Any]) -> Path:
    source_path = input_data.get("source_path")
    if not isinstance(source_path, str) or not source_path.strip():
        raise WorkerError("issue_source_not_found", "Issue source path is required.")
    path_value = Path(source_path).resolve()
    if path_value.suffix.lower() not in SUPPORTED_SCHEDULE_EXTENSIONS:
        raise WorkerError("unsupported_issue_file", "Issue source must be .xlsx or .xlsm.")
    if not path_value.exists() or not path_value.is_file():
        raise WorkerError("issue_source_not_found", "Issue source file was not found.")
    return path_value


def normalize_issue_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("_x000D_", " ").replace("_x000d_", " ")
    return re.sub(r"[^0-9A-Za-z가-힣]+", "", text.strip().lower())


def issue_header_matches(normalized: str, aliases: set[str]) -> bool:
    if normalized in aliases:
        return True
    return any(alias and len(alias) >= 6 and alias in normalized for alias in aliases)


def missing_issue_required_columns(mapping: dict[str, int]) -> list[str]:
    missing: list[str] = []
    for group in ISSUE_REQUIRED_COLUMN_GROUPS:
        if not any(field in mapping for field in group):
            missing.append("/".join(sorted(group)))
    return missing


def issue_detected_columns(rows: list[tuple[Any, ...]], header_row: int) -> list[str]:
    if header_row < 0 or header_row >= len(rows):
        return []
    return [
        cell_to_text(value)
        for value in rows[header_row]
        if cell_to_text(value)
    ]


def issue_mapping_failure_message(
    sheet_name: str | None,
    header_row: int,
    mapping: dict[str, int],
    rows: list[tuple[Any, ...]],
    missing: list[str],
) -> str:
    detected_columns = issue_detected_columns(rows, header_row)
    mapped_fields = ",".join(sorted(mapping.keys())) or "-"
    return (
        "Issue column mapping failed. "
        f"detected_sheet={sheet_name or '-'}; "
        f"detected_header_row={(header_row + 1) if header_row >= 0 else '-'}; "
        f"detected_columns={detected_columns}; "
        f"mapped_fields={mapped_fields}; "
        f"missing_required_columns={missing or '-'}"
    )


def detect_issue_header_mapping(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int], list[str]]:
    normalized_aliases = {
        field: {normalize_issue_header(alias) for alias in values}
        for field, values in ISSUE_COLUMN_ALIASES.items()
    }
    best_row = -1
    best_mapping: dict[str, int] = {}
    best_score = -1
    for row_index, row in enumerate(rows[:40]):
        mapping: dict[str, int] = {}
        normalized_cells = [normalize_issue_header(value) for value in row]
        for cell_index, normalized in enumerate(normalized_cells):
            if not normalized:
                continue
            for field, field_aliases in normalized_aliases.items():
                if field in mapping:
                    continue
                if issue_header_matches(normalized, field_aliases):
                    mapping[field] = cell_index
        score = len(mapping)
        if "title" in mapping or "description" in mapping:
            score += 2
        if "status" in mapping:
            score += 2
        if score > best_score:
            best_row = row_index
            best_mapping = mapping
            best_score = score
    return best_row, best_mapping, missing_issue_required_columns(best_mapping)


def parse_issue_date(value: Any) -> str | None:
    return parse_schedule_date(value)


def parse_issue_number(value: Any) -> float | None:
    return parse_schedule_number(value)


def parse_issue_bool(value: Any) -> bool | None:
    text = cell_to_text(value).strip().lower()
    if not text:
        return None
    if text in {"y", "yes", "true", "1", "필요", "예", "o", "○"}:
        return True
    if text in {"n", "no", "false", "0", "불필요", "아니오", "x"}:
        return False
    return True if "필요" in text or "support" in text else None


def normalize_issue_status(value: Any) -> str | None:
    text = cell_to_text(value)
    if not text:
        return None
    lowered = text.lower()
    if any(token in lowered for token in ["보류", "hold", "defer", "pending"]):
        return "on_hold"
    if any(token in lowered for token in ["진행중", "조치중", "doing", "progress", "in progress"]):
        return "in_progress"
    if any(token in lowered for token in ["종결", "closed"]):
        return "closed"
    if any(token in lowered for token in ["완료", "해결", "done", "complete", "resolved"]):
        return "resolved"
    if any(token in lowered for token in ["open", "신규", "접수", "미해결"]):
        return "open"
    return text


def is_completed_issue(issue: dict[str, Any]) -> bool:
    return issue.get("status") in {"resolved", "closed", "completed"} or bool(issue.get("completedDate") or issue.get("resolvedDate"))


def normalize_issue_item_type(value: Any, row: tuple[Any, ...], mapping: dict[str, int]) -> str:
    text = cell_to_text(value).lower()
    if any(token in text for token in ["risk", "위험", "리스크"]):
        return "risk"
    if any(token in text for token in ["issue", "이슈"]):
        return "issue"
    for field in ISSUE_RISK_HINT_FIELDS:
        if cell_to_text(row_value(row, mapping, field)):
            return "risk"
    return "issue"


def days_between_dates(start: str | None, finish: str | None) -> int | None:
    start_date = parse_iso_date(start)
    finish_date = parse_iso_date(finish)
    if not start_date or not finish_date:
        return None
    return (finish_date.date() - start_date.date()).days


def extract_latest_note_date(value: Any) -> str | None:
    text = cell_to_text(value)
    if not text:
        return None
    candidates: list[str] = []
    for match in re.finditer(r"\((\d{4}[./-]\d{1,2}[./-]\d{1,2})\)", text):
        parsed = parse_issue_date(match.group(1))
        if parsed:
            candidates.append(parsed)
    for match in re.finditer(r"\b(\d{4}[./-]\d{1,2}[./-]\d{1,2})\b", text):
        parsed = parse_issue_date(match.group(1))
        if parsed:
            candidates.append(parsed)
    return max(candidates) if candidates else None


def issue_risk_score(probability: float | None, impact: float | None) -> float | None:
    if probability is None and impact is None:
        return None
    p = probability if probability is not None else 1
    i = impact if impact is not None else 1
    return float(p) * float(i)


def issue_risk_flags(issue: dict[str, Any], as_of_date: str) -> list[str]:
    flags: list[str] = []
    metrics = issue.get("derivedMetrics") if isinstance(issue.get("derivedMetrics"), dict) else {}
    if not is_completed_issue(issue) and not issue.get("dueDate"):
        flags.append("missing_due")
    if isinstance(metrics.get("overdueDays"), int) and metrics["overdueDays"] > 0:
        flags.append("overdue")
    if isinstance(metrics.get("openAgeDays"), int) and metrics["openAgeDays"] >= 60:
        flags.append("long_open_60d")
    if issue.get("organizationSupportRequired") is True and not is_completed_issue(issue):
        flags.append("organization_support_required")
    if isinstance(metrics.get("daysSinceLastUpdate"), int) and metrics["daysSinceLastUpdate"] >= 14:
        flags.append("stale_update")
    if issue_risk_score(issue.get("probability"), issue.get("impact")) and issue_risk_score(issue.get("probability"), issue.get("impact")) >= 9:
        flags.append("high_risk_score")
    return flags


def build_issue_metrics(issue: dict[str, Any], as_of_date: str) -> dict[str, Any]:
    completed = is_completed_issue(issue)
    as_of = as_of_date
    open_age_days = None if completed else days_between_dates(issue.get("occurredDate"), as_of)
    issue_age_finish = issue.get("actionDate") or issue.get("resolvedDate") or issue.get("completedDate") if completed else as_of
    issue_age_days = days_between_dates(issue.get("occurredDate"), issue_age_finish)
    overdue_days = None
    if not completed and issue.get("dueDate"):
        overdue_days = days_between_dates(issue.get("dueDate"), as_of)
        if overdue_days is not None:
            overdue_days = max(0, overdue_days)
    days_since_last_update = None if completed else days_between_dates(issue.get("latestNoteDate") or issue.get("occurredDate"), as_of)
    resolution_days = days_between_dates(issue.get("occurredDate"), issue.get("completedDate")) if completed else None
    completed_late_days = None
    if completed and issue.get("dueDate") and issue.get("completedDate"):
        completed_late_days = max(0, days_between_dates(issue.get("dueDate"), issue.get("completedDate")) or 0)
    derived_risk_score = issue_risk_score(issue.get("probability"), issue.get("impact"))
    metrics = {
        "openAgeDays": open_age_days,
        "issueAgeDays": issue_age_days,
        "overdueDays": overdue_days,
        "daysSinceLastUpdate": days_since_last_update,
        "resolutionDays": resolution_days,
        "completedLateDays": completed_late_days,
        "derivedRiskScore": derived_risk_score,
        "riskFlags": [],
    }
    issue["derivedMetrics"] = metrics
    metrics["riskFlags"] = issue_risk_flags(issue, as_of_date)
    return metrics


def candidate_issue_sheets(workbook: Any) -> list[Any]:
    visible_sheets = [
        sheet for sheet in workbook.worksheets
        if getattr(sheet, "sheet_state", "visible") == "visible"
    ]
    return visible_sheets or list(workbook.worksheets)


def detect_issue_sheet(workbook: Any) -> tuple[Any, int, dict[str, int]]:
    best_sheet = None
    best_rows: list[tuple[Any, ...]] = []
    best_header_row = -1
    best_mapping: dict[str, int] = {}
    best_missing: list[str] = []
    best_score = -1
    for sheet in candidate_issue_sheets(workbook):
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_row, mapping, missing = detect_issue_header_mapping(rows)
        body_signal = 0
        if header_row >= 0:
            for row in rows[header_row + 1 : header_row + 16]:
                if any(cell_to_text(value) for value in row):
                    body_signal += 1
        score = len(mapping) * 10 + body_signal - (len(missing) * 100)
        if score > best_score:
            best_sheet = sheet
            best_rows = rows
            best_header_row = header_row
            best_mapping = mapping
            best_missing = missing
            best_score = score
    if best_sheet is None or best_missing:
        raise WorkerError(
            "issue_column_mapping_failed",
            issue_mapping_failure_message(
                getattr(best_sheet, "title", None),
                best_header_row,
                best_mapping,
                best_rows,
                best_missing,
            ),
        )
    return best_sheet, best_header_row, best_mapping


def issue_has_row_signal(row: tuple[Any, ...], mapping: dict[str, int]) -> bool:
    signal_fields = [
        "item_id",
        "title",
        "description",
        "status",
        "action_plan",
        "resolution",
        "remarks",
        "owner",
    ]
    return any(cell_to_text(row_value(row, mapping, field)) for field in signal_fields)


def parse_issue_sheet(
    workbook: Any,
    workspace_id: str,
    source_path: Path,
    source_hash: str,
    as_of_date: str,
) -> tuple[list[dict[str, Any]], dict[str, str], str, int, list[dict[str, Any]]]:
    sheet, header_row, mapping = detect_issue_sheet(workbook)
    sheet_name = sheet.title
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise WorkerError("issue_parse_failed", "Issue sheet is empty.")
    issues: list[dict[str, Any]] = []
    parse_warnings: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        title = cell_to_text(row_value(row, mapping, "title"))
        description = cell_to_text(row_value(row, mapping, "description")) or None
        raw_status = cell_to_text(row_value(row, mapping, "status")) or None
        if not issue_has_row_signal(row, mapping):
            parse_warnings.append({"row": row_number, "reason": "empty_or_unidentified_row"})
            continue
        if not title and not description:
            parse_warnings.append({"row": row_number, "reason": "missing_title_or_description"})
            continue
        action_plan = cell_to_text(row_value(row, mapping, "action_plan")) or None
        owner = cell_to_text(row_value(row, mapping, "owner")) or None
        remarks = cell_to_text(row_value(row, mapping, "remarks")) or None
        explicit_item_id = cell_to_text(row_value(row, mapping, "item_id"))
        item_type = normalize_issue_item_type(row_value(row, mapping, "item_type"), row, mapping)
        item_id = explicit_item_id or f"{workspace_id}:{item_type}:{row_number}"
        reported_date = parse_issue_date(row_value(row, mapping, "reported_date"))
        target_date = parse_issue_date(row_value(row, mapping, "target_date"))
        resolved_date = parse_issue_date(row_value(row, mapping, "resolved_date"))
        status = normalize_issue_status(raw_status)
        issue = {
            "itemId": item_id,
            "itemType": item_type,
            "issueId": item_id,
            "sourceRow": row_number,
            "issueArea": cell_to_text(row_value(row, mapping, "issue_area")) or None,
            "issueCategory": cell_to_text(row_value(row, mapping, "issue_category")) or None,
            "phase": cell_to_text(row_value(row, mapping, "phase")) or None,
            "title": title,
            "issueEvent": title,
            "description": description,
            "reportedDate": reported_date,
            "occurredDate": reported_date,
            "riskAnalysisId": cell_to_text(row_value(row, mapping, "risk_analysis_id")) or None,
            "riskId": cell_to_text(row_value(row, mapping, "risk_id")) or None,
            "probability": parse_issue_number(row_value(row, mapping, "probability")),
            "impact": parse_issue_number(row_value(row, mapping, "impact")),
            "severity": cell_to_text(row_value(row, mapping, "severity")) or None,
            "riskLevel": cell_to_text(row_value(row, mapping, "severity")) or None,
            "rootCause": cell_to_text(row_value(row, mapping, "root_cause")) or None,
            "actionPlan": action_plan,
            "responsePlan": action_plan,
            "resolution": cell_to_text(row_value(row, mapping, "resolution")) or None,
            "owner": owner,
            "actionOwner": owner,
            "targetDate": target_date,
            "dueDate": target_date,
            "organizationSupportRequired": parse_issue_bool(row_value(row, mapping, "organization_support_required")),
            "rawStatus": raw_status,
            "progressStatus": raw_status,
            "status": status,
            "canonicalStatus": status,
            "priority": cell_to_text(row_value(row, mapping, "priority")) or None,
            "effortMh": parse_issue_number(row_value(row, mapping, "effort_mh")),
            "actionDate": resolved_date,
            "resolvedDate": resolved_date,
            "completedDate": resolved_date,
            "remarks": remarks,
            "noteTimeline": remarks or action_plan,
            "latestNoteDate": extract_latest_note_date(remarks or action_plan),
            "sourceHash": source_hash,
            "dataQualityFlags": [],
            "derivedMetrics": {},
        }
        if not raw_status:
            issue["dataQualityFlags"].append("missing_status")
            parse_warnings.append({"row": row_number, "reason": "missing_status_partial_parse"})
        if not issue["reportedDate"]:
            issue["dataQualityFlags"].append("missing_reported_date")
        if not issue["dueDate"] and not is_completed_issue(issue):
            issue["dataQualityFlags"].append("missing_due_date")
        if not issue["owner"]:
            issue["dataQualityFlags"].append("missing_owner")
        build_issue_metrics(issue, as_of_date)
        issues.append(issue)
    column_mapping = {field: cell_to_text(rows[header_row][index]) for field, index in mapping.items()}
    return issues, column_mapping, sheet_name, header_row + 1, parse_warnings


def issue_summary(issue_book: dict[str, Any], document: dict[str, Any], as_of_date: str | None = None) -> dict[str, Any]:
    current_as_of = as_of_date or datetime.now().date().isoformat()
    issues = [issue for issue in issue_book.get("issues", []) if isinstance(issue, dict)]
    for issue in issues:
        build_issue_metrics(issue, current_as_of)
    open_issues = [issue for issue in issues if not is_completed_issue(issue)]
    completed_issues = [issue for issue in issues if is_completed_issue(issue)]
    issue_items = [issue for issue in issues if issue.get("itemType") != "risk"]
    risk_items = [issue for issue in issues if issue.get("itemType") == "risk"]
    raw_status_values = sorted({
        cell_to_text(issue.get("rawStatus"))
        for issue in issues
        if cell_to_text(issue.get("rawStatus"))
    })
    status_counts: dict[str, int] = {}
    raw_status_counts: dict[str, int] = {}
    for issue in issues:
        status = cell_to_text(issue.get("status")) or "unknown"
        raw_status = cell_to_text(issue.get("rawStatus")) or "unknown"
        status_counts[status] = status_counts.get(status, 0) + 1
        raw_status_counts[raw_status] = raw_status_counts.get(raw_status, 0) + 1
    resolution_days = [
        issue.get("derivedMetrics", {}).get("resolutionDays")
        for issue in completed_issues
        if isinstance(issue.get("derivedMetrics", {}).get("resolutionDays"), int)
    ]
    completed_late_values = [
        issue.get("derivedMetrics", {}).get("completedLateDays")
        for issue in completed_issues
        if issue.get("dueDate") and issue.get("completedDate")
    ]
    due_adherent = sum(1 for value in completed_late_values if isinstance(value, int) and value <= 0)
    overdue_open = [
        issue for issue in open_issues
        if isinstance(issue.get("derivedMetrics", {}).get("overdueDays"), int)
        and issue["derivedMetrics"]["overdueDays"] > 0
    ]
    stale = [
        issue for issue in open_issues
        if isinstance(issue.get("derivedMetrics", {}).get("daysSinceLastUpdate"), int)
        and issue["derivedMetrics"]["daysSinceLastUpdate"] >= 14
    ]
    return {
        "workspaceId": issue_book["workspaceId"],
        "filename": issue_book["filename"],
        "modifiedAt": document.get("modifiedAt"),
        "lastAnalyzedAt": document.get("lastAnalyzedAt") or now_iso(),
        "parseStatus": document.get("parseStatus", "parsed"),
        "totalIssues": len(issues),
        "issueCount": len(issue_items),
        "riskCount": len(risk_items),
        "openIssues": len(open_issues),
        "completedIssues": len(completed_issues),
        "resolvedIssues": len(completed_issues),
        "openInProgressIssues": sum(1 for issue in issues if issue.get("status") in {"open", "in_progress"}),
        "completionRate": (len(completed_issues) / len(issues)) if issues else None,
        "overdueOpenIssues": len(overdue_open),
        "longOpenIssues30d": sum(1 for issue in open_issues if (issue.get("derivedMetrics", {}).get("openAgeDays") or 0) >= 30),
        "longOpenIssues60d": sum(1 for issue in open_issues if (issue.get("derivedMetrics", {}).get("openAgeDays") or 0) >= 60),
        "longOpenIssues90d": sum(1 for issue in open_issues if (issue.get("derivedMetrics", {}).get("openAgeDays") or 0) >= 90),
        "missingDueOpenIssues": sum(1 for issue in open_issues if not issue.get("dueDate")),
        "staleUpdateIssues": len(stale),
        "organizationSupportRequiredOpen": sum(1 for issue in open_issues if issue.get("organizationSupportRequired") is True),
        "averageResolutionDays": (sum(resolution_days) / len(resolution_days)) if resolution_days else None,
        "medianResolutionDays": median_number(resolution_days),
        "dueDateAdherenceRate": (due_adherent / len(completed_late_values)) if completed_late_values else None,
        "statusCounts": status_counts,
        "rawStatusCounts": raw_status_counts,
        "rawStatusValues": raw_status_values,
        "skippedRows": len(issue_book.get("parseWarnings", []) or []),
        "parseWarnings": issue_book.get("parseWarnings", []) or [],
        "source": document,
    }


def median_number(values: list[int | float]) -> float | None:
    if not values:
        return None
    sorted_values = sorted(float(value) for value in values)
    mid = len(sorted_values) // 2
    if len(sorted_values) % 2:
        return sorted_values[mid]
    return (sorted_values[mid - 1] + sorted_values[mid]) / 2


def parse_issue(input_data: dict[str, Any]) -> dict[str, Any]:
    workspace_id = input_data.get("workspace_id")
    if not isinstance(workspace_id, str) or not workspace_id.strip():
        raise WorkerError("invalid_workspace", "workspace_id is required.")
    source_path = issue_source_path(input_data)
    source_hash = cell_to_text(input_data.get("source_hash")) or sha256_file(source_path)
    as_of_date = input_data.get("as_of_date") if isinstance(input_data.get("as_of_date"), str) else datetime.now().date().isoformat()
    try:
        import openpyxl
    except ImportError as error:
        raise WorkerError("issue_parse_failed", "openpyxl is required to parse Issue Excel files.") from error
    try:
        workbook = openpyxl.load_workbook(source_path, data_only=True, read_only=False)
        issues, column_mapping, sheet_name, header_row, parse_warnings = parse_issue_sheet(workbook, workspace_id, source_path, source_hash, as_of_date)
    except WorkerError:
        raise
    except Exception as error:
        raise WorkerError("issue_parse_failed", "Issue workbook parsing failed.") from error
    stats = source_path.stat()
    analyzed_at = now_iso()
    document = {
        "id": cell_to_text(input_data.get("source_document_id")) or f"issue-{workspace_id}",
        "workspaceId": workspace_id,
        "originalFileName": source_path.name,
        "managedFilePath": str(source_path),
        "sourceHash": source_hash,
        "registeredAt": analyzed_at,
        "registeredBy": "local-user",
        "status": "active",
        "disconnectedAt": None,
        "lastAnalyzedAt": analyzed_at,
        "parserType": "issue_excel",
        "security": "internal",
        "notes": None,
        "fileSize": stats.st_size,
        "modifiedAt": datetime.fromtimestamp(stats.st_mtime, timezone.utc).isoformat().replace("+00:00", "Z"),
        "parseStatus": "parsed",
    }
    issue_book = {
        "workspaceId": workspace_id,
        "sourceFile": str(source_path),
        "filename": source_path.name,
        "parsedSheets": [sheet_name],
        "headerRow": header_row,
        "detectedColumns": list(column_mapping.values()),
        "columnMapping": column_mapping,
        "parseWarnings": parse_warnings,
        "issues": issues,
    }
    return {
        "document": document,
        "issueBook": issue_book,
        "summary": issue_summary(issue_book, document, as_of_date),
        "fromCache": False,
    }


def issue_summary_command(input_data: dict[str, Any]) -> dict[str, Any]:
    issue_book = input_data.get("issue_book")
    document = input_data.get("document")
    if not isinstance(issue_book, dict) or not isinstance(document, dict):
        raise WorkerError("invalid_input", "issue_book and document are required.")
    return issue_summary(issue_book, document, input_data.get("as_of_date"))


def normalize_issue_lookup_text(value: Any) -> str:
    return re.sub(r"\s+", " ", cell_to_text(value).lower()).strip()


def tokenize_issue_text(value: Any) -> list[str]:
    text = normalize_issue_lookup_text(value)
    tokens = re.findall(r"[0-9a-zA-Z가-힣]{2,}", text)
    stopwords = {"issue", "risk", "task", "open", "closed", "완료", "진행", "이슈", "조치"}
    return [token for token in tokens if token not in stopwords]


def issue_matches_query(issue: dict[str, Any], query: str) -> bool:
    normalized_query = normalize_issue_lookup_text(query)
    haystack = " ".join(
        normalize_issue_lookup_text(issue.get(field))
        for field in ["itemId", "issueId", "itemType", "title", "description", "issueArea", "issueCategory", "phase", "owner", "riskId", "rawStatus", "status"]
    )
    return any(token in haystack for token in tokenize_issue_text(normalized_query))


def issue_owner_summaries(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_owner: dict[str, list[dict[str, Any]]] = {}
    for issue in issues:
        owner = cell_to_text(issue.get("owner")) or "Unassigned"
        by_owner.setdefault(owner, []).append(issue)
    summaries: list[dict[str, Any]] = []
    for owner, owner_issues in by_owner.items():
        open_items = [issue for issue in owner_issues if not is_completed_issue(issue)]
        overdue_items = [
            issue for issue in open_items
            if (issue.get("derivedMetrics", {}).get("overdueDays") or 0) > 0
        ]
        long_open_items = [
            issue for issue in open_items
            if (issue.get("derivedMetrics", {}).get("openAgeDays") or 0) >= 30
        ]
        missing_due_items = [
            issue for issue in open_items
            if not issue.get("dueDate")
        ]
        stale_update_items = [
            issue for issue in open_items
            if (issue.get("derivedMetrics", {}).get("daysSinceLastUpdate") or 0) >= 14
        ]
        resolution_values = [
            issue.get("derivedMetrics", {}).get("resolutionDays")
            for issue in owner_issues
            if isinstance(issue.get("derivedMetrics", {}).get("resolutionDays"), int)
        ]
        bottleneck_score = (
            (len(overdue_items) * 4)
            + (len(missing_due_items) * 3)
            + (len(stale_update_items) * 2)
            + (len(long_open_items) * 2)
            + len(open_items)
        )
        summaries.append({
            "owner": owner,
            "totalAssigned": len(owner_issues),
            "openAssigned": len(open_items),
            "overdueAssigned": len(overdue_items),
            "missingDueAssigned": len(missing_due_items),
            "staleUpdateAssigned": len(stale_update_items),
            "longOpenAssigned": len(long_open_items),
            "averageResolutionDays": (sum(resolution_values) / len(resolution_values)) if resolution_values else None,
            "bottleneckScore": bottleneck_score,
        })
    return sorted(summaries, key=lambda item: (-item["bottleneckScore"], item["owner"]))[:20]


def issue_keyword_clusters(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    clusters: dict[str, list[dict[str, Any]]] = {}
    for issue in issues:
        tokens = tokenize_issue_text(issue.get("title"))
        for token in tokens:
            if len(token) < 3:
                continue
            clusters.setdefault(token, []).append(issue)
    result = []
    for key, cluster_issues in clusters.items():
        unique_ids = list(dict.fromkeys(cell_to_text(issue.get("issueId")) for issue in cluster_issues))
        if len(unique_ids) < 2:
            continue
        result.append({
            "key": key,
            "count": len(unique_ids),
            "issueIds": unique_ids[:20],
            "sampleTitles": [cell_to_text(issue.get("title")) for issue in cluster_issues[:3]],
        })
    return sorted(result, key=lambda item: (-item["count"], item["key"]))[:20]


def issue_context_text(kind: str, summary: dict[str, Any], issues: list[dict[str, Any]], assignees: list[dict[str, Any]] | None = None, clusters: list[dict[str, Any]] | None = None, query_plan: dict[str, Any] | None = None) -> str:
    lines = [
        "[ISSUE SOURCE OF TRUTH]",
        f"Type: {kind}",
        f"Workspace ID: {summary.get('workspaceId')}",
        f"Source: {summary.get('filename')}",
        f"Total Items: {summary.get('totalIssues')}",
        f"Issue Count: {summary.get('issueCount')}",
        f"Risk Count: {summary.get('riskCount')}",
        f"Open/In Progress Items: {summary.get('openInProgressIssues')}",
        f"Resolved/Closed Items: {summary.get('resolvedIssues')}",
        f"Status Counts: {summary.get('statusCounts')}",
        f"Raw Status Values: {summary.get('rawStatusValues')}",
        f"Overdue Open Issues: {summary.get('overdueOpenIssues')}",
        f"Organization Support Required Open: {summary.get('organizationSupportRequiredOpen')}",
        f"Skipped Rows: {summary.get('skippedRows')}",
    ]
    if query_plan:
        lines.extend([
            f"Detected Intent: {query_plan.get('detectedIntent')}",
            f"Temporal Scope: {query_plan.get('temporalScope')}",
            f"Status Filter: {query_plan.get('statusFilter')}",
            f"Additional Filters: {query_plan.get('additionalFilters')}",
            f"Sort Order: {query_plan.get('sortOrder')}",
        ])
    if assignees:
        lines.append("Assignee Bottlenecks:")
        for assignee in assignees[:10]:
            lines.append(
                f"- {assignee.get('owner')}: open={assignee.get('openAssigned')} overdue={assignee.get('overdueAssigned')} missingDue={assignee.get('missingDueAssigned')} stale={assignee.get('staleUpdateAssigned')} score={assignee.get('bottleneckScore')}"
            )
    if clusters:
        lines.append("Repeated Issue Clusters:")
        for cluster in clusters[:10]:
            lines.append(f"- {cluster.get('key')}: {cluster.get('count')} issues")
    if issues:
        lines.append("Issues:")
        for issue in issues[:20]:
            metrics = issue.get("derivedMetrics", {}) if isinstance(issue.get("derivedMetrics"), dict) else {}
            lines.append(
                f"- Row {issue.get('sourceRow')} | {issue.get('itemType') or 'issue'} | {issue.get('itemId') or issue.get('issueId')} | {issue.get('title')} | owner={issue.get('owner') or '-'} | rawStatus={issue.get('rawStatus') or '-'} | status={issue.get('status') or '-'} | target={issue.get('targetDate') or issue.get('dueDate') or '-'} | overdueDays={metrics.get('overdueDays')}"
            )
    else:
        lines.append("Issues: No matching issues.")
    lines.extend([
        "Instruction: These are authoritative current Issue values from the connected Excel source.",
        "Instruction: Use RAG only for past resolved case references, not for current issue status.",
        "[/ISSUE SOURCE OF TRUTH]",
    ])
    return "\n".join(lines)


CURRENT_OPEN_STATUSES = {"open", "in_progress", "on_hold"}
RESOLVED_STATUSES = {"resolved", "closed", "completed"}
ISSUE_QUERY_INTENTS = {
    "issue_list",
    "issue_lookup",
    "open_issues",
    "unresolved_issues",
    "owner_bottleneck",
    "overdue_issues",
    "organization_support",
    "risk_candidates",
    "resolved_cases",
    "issue_status",
}
ISSUE_QUERY_STATUS_SCOPES = {"unresolved", "resolved", "all"}
ISSUE_QUERY_FILTER_FIELDS = {
    "owner",
    "issue_area",
    "issue_category",
    "phase",
    "organization_support_required",
    "overdue_only",
}
ISSUE_QUERY_SORT_VALUES = {
    "overdue_desc",
    "missing_due_desc",
    "issue_age_desc",
    "open_count_desc",
    "latest_note_date_asc_null_first",
    "organization_support_desc",
    "target_date_urgency_asc",
    "status_desc",
    "severity_priority_desc",
    "source_row_asc",
}
ISSUE_QUERY_GROUP_BY_VALUES = {None, "action_owner", "issue_area", "issue_category", "phase", "status"}
ISSUE_QUERY_STATUS_TO_TEMPORAL = {
    "unresolved": "CURRENT_OPEN",
    "resolved": "RESOLVED_ONLY",
    "all": "HISTORICAL_ALL",
}


def issue_contains_any(query_text: str, tokens: list[str]) -> bool:
    return any(token in query_text for token in tokens)


def issue_compact_text(query_text: str) -> str:
    return re.sub(r"\s+", "", query_text)


def has_unresolved_issue_signal(query_text: str) -> bool:
    compact = issue_compact_text(query_text)
    return issue_contains_any(query_text, [
        "미해결",
        "진행 중",
        "진행중",
        "열린 이슈",
        "오픈 이슈",
        "남은 이슈",
        "남아있는 이슈",
        "open",
        "unresolved",
    ]) or issue_contains_any(compact, [
        "해결안된",
        "해결안된이슈",
        "해결이안된",
        "안끝난",
        "진행중",
        "진행중인",
        "열린이슈",
        "오픈이슈",
        "남은이슈",
        "남아있는이슈",
    ])


def has_completed_history_include_signal(query_text: str) -> bool:
    compact = issue_compact_text(query_text)
    return (
        issue_contains_any(query_text, ["완료 이력 포함", "완료된 이슈도 포함", "완료도 포함", "resolved history", "include resolved", "include completed"])
        or ("포함" in compact and issue_contains_any(compact, ["완료된", "해결된", "완료이력"]))
    )


def issue_status_scope_from_text(query_text: str) -> str:
    if has_completed_history_include_signal(query_text):
        return "all"
    if issue_contains_any(query_text, ["해결된", "완료된", "어떻게 해결", "조치가 완료", "해결 사례", "resolved", "closed", "completed"]):
        return "resolved"
    if issue_contains_any(query_text, ["과거", "이력", "이력이 있는", "필요했던", "발생했던", "전체", "모든", "historical", "history", "all"]):
        return "all"
    if has_unresolved_issue_signal(query_text):
        return "unresolved"
    if issue_contains_any(query_text, ["누가", "담당", "맡", "owner", "assignee", "bottleneck", "병목"]):
        return "unresolved"
    if issue_contains_any(query_text, ["조직지원", "조직 지원", "조직의 지원", "지원"]) and issue_contains_any(query_text, ["필요한", "필요"]):
        return "unresolved"
    if issue_contains_any(query_text, [
        "현재",
        "진행 중",
        "진행중",
        "해결이 필요한",
        "위험한 이슈",
        "지연된 이슈",
        "오래 지연",
        "조치가 필요한",
        "조직지원이 필요한",
        "조직지원 필요한",
        "지원이 필요한",
        "지원 필요한",
        "지원 필요",
        "조직의 지원이 필요한",
        "open",
        "current",
        "overdue",
        "late",
    ]):
        return "unresolved"
    return "all"


def issue_temporal_scope(query_text: str) -> str:
    return ISSUE_QUERY_STATUS_TO_TEMPORAL[issue_status_scope_from_text(query_text)]


def create_issue_query_filters() -> dict[str, Any]:
    return {
        "owner": None,
        "issue_area": None,
        "issue_category": None,
        "phase": None,
        "organization_support_required": None,
        "overdue_only": False,
    }


def issue_intent_from_text(query_text: str, status_scope: str) -> str:
    has_owner_intent = issue_contains_any(query_text, ["bottleneck", "owner", "assignee", "누가", "담당", "병목", "맡"])
    has_support_intent = issue_contains_any(query_text, ["support", "조직지원", "조직 지원", "조직의 지원", "지원"])
    has_risk_intent = issue_contains_any(query_text, ["risk", "위험", "리스크"])
    has_overdue_intent = issue_contains_any(query_text, ["overdue", "late", "지연", "기한", "늦"])
    has_status_intent = issue_contains_any(query_text, ["상태", "진행상황", "진행 상황", "status"])

    if has_owner_intent:
        return "owner_bottleneck"
    if has_completed_history_include_signal(query_text):
        return "open_issues"
    if has_unresolved_issue_signal(query_text):
        return "open_issues"
    if has_support_intent:
        return "organization_support"
    if has_risk_intent:
        return "risk_candidates"
    if has_overdue_intent:
        return "overdue_issues"
    if status_scope == "resolved":
        return "resolved_cases"
    if has_status_intent:
        return "issue_status"
    return "issue_lookup" if tokenize_issue_text(query_text) else "issue_list"


def issue_sort_from_intent(intent: str) -> list[str]:
    if intent == "owner_bottleneck":
        return ["overdue_desc", "missing_due_desc", "issue_age_desc", "latest_note_date_asc_null_first", "open_count_desc"]
    if intent in {"open_issues", "unresolved_issues"}:
        return ["overdue_desc", "missing_due_desc", "issue_age_desc", "latest_note_date_asc_null_first"]
    if intent == "risk_candidates":
        return [
            "overdue_desc",
            "issue_age_desc",
            "organization_support_desc",
            "target_date_urgency_asc",
            "status_desc",
            "severity_priority_desc",
        ]
    if intent == "organization_support":
        return ["overdue_desc", "issue_age_desc", "status_desc"]
    if intent == "overdue_issues":
        return ["overdue_desc", "issue_age_desc"]
    if intent == "resolved_cases":
        return ["source_row_asc"]
    return ["source_row_asc"]


def issue_group_by_from_intent(intent: str) -> str | None:
    return "action_owner" if intent == "owner_bottleneck" else None


def interpret_issue_query_locally(query: str) -> dict[str, Any]:
    query_text = normalize_issue_lookup_text(query)
    status_scope = issue_status_scope_from_text(query_text)
    intent = issue_intent_from_text(query_text, status_scope)
    filters = create_issue_query_filters()
    if intent == "organization_support":
        filters["organization_support_required"] = True
    if intent == "overdue_issues":
        filters["overdue_only"] = True
    return {
        "intent": intent,
        "status_scope": status_scope,
        "filters": filters,
        "group_by": issue_group_by_from_intent(intent),
        "sort": issue_sort_from_intent(intent),
        "limit": 20,
    }


def validate_issue_query_filter(raw_filters: Any) -> dict[str, Any]:
    filters = create_issue_query_filters()
    if not isinstance(raw_filters, dict):
        return filters
    for key, value in raw_filters.items():
        if key not in ISSUE_QUERY_FILTER_FIELDS:
            continue
        if key in {"owner", "issue_area", "issue_category", "phase"}:
            filters[key] = cell_to_text(value) or None
        elif key == "organization_support_required":
            filters[key] = value if isinstance(value, bool) else None
        elif key == "overdue_only":
            filters[key] = value is True
    return filters


def validate_issue_query(value: Any, fallback_query: str) -> dict[str, Any]:
    fallback = interpret_issue_query_locally(fallback_query)
    if not isinstance(value, dict):
        return fallback
    intent = value.get("intent") if value.get("intent") in ISSUE_QUERY_INTENTS else fallback["intent"]
    status_scope = value.get("status_scope") if value.get("status_scope") in ISSUE_QUERY_STATUS_SCOPES else fallback["status_scope"]
    if fallback["intent"] in {"open_issues", "unresolved_issues"} and intent in {"issue_lookup", "issue_list"}:
        intent = fallback["intent"]
    if intent == "owner_bottleneck" and status_scope == "all" and fallback["status_scope"] == "unresolved":
        status_scope = "unresolved"
    filters = validate_issue_query_filter(value.get("filters"))
    raw_group_by = value.get("group_by")
    group_by = raw_group_by if raw_group_by in ISSUE_QUERY_GROUP_BY_VALUES else issue_group_by_from_intent(intent)
    raw_sort = value.get("sort")
    sort = [
        item for item in raw_sort
        if isinstance(item, str) and item in ISSUE_QUERY_SORT_VALUES
    ] if isinstance(raw_sort, list) else []
    limit = value.get("limit")
    if not isinstance(limit, int) or limit < 1 or limit > 100:
        limit = 20
    return {
        "intent": intent,
        "status_scope": status_scope,
        "filters": filters,
        "group_by": group_by,
        "sort": sort or issue_sort_from_intent(intent),
        "limit": limit,
    }


def issue_status_scope_label(scope: str) -> str:
    if scope == "unresolved":
        return "진행중/미해결"
    if scope == "resolved":
        return "해결/종결(resolved, closed)"
    if scope == "all":
        return "전체 이력"
    if scope == "CURRENT_OPEN":
        return "진행중/미해결"
    if scope == "RESOLVED_ONLY":
        return "해결/종결(resolved, closed)"
    return "전체 이력"


def issue_in_temporal_scope(issue: dict[str, Any], scope: str) -> bool:
    status = cell_to_text(issue.get("status"))
    if scope == "CURRENT_OPEN":
        return status in CURRENT_OPEN_STATUSES and not is_completed_issue(issue)
    if scope == "RESOLVED_ONLY":
        return status in RESOLVED_STATUSES or is_completed_issue(issue)
    return True


def scoped_issues(issues: list[dict[str, Any]], scope: str) -> list[dict[str, Any]]:
    temporal_scope = ISSUE_QUERY_STATUS_TO_TEMPORAL.get(scope, scope)
    return [issue for issue in issues if issue_in_temporal_scope(issue, temporal_scope)]


def issue_filter_text_matches(value: Any, expected: str | None) -> bool:
    if not expected:
        return True
    return normalize_issue_lookup_text(expected) in normalize_issue_lookup_text(value)


def apply_issue_query_filters(issues: list[dict[str, Any]], issue_query: dict[str, Any]) -> list[dict[str, Any]]:
    filtered = scoped_issues(issues, cell_to_text(issue_query.get("status_scope")) or "all")
    filters = issue_query.get("filters") if isinstance(issue_query.get("filters"), dict) else {}
    owner = filters.get("owner")
    issue_area = filters.get("issue_area")
    issue_category = filters.get("issue_category")
    phase = filters.get("phase")
    if owner:
        filtered = [issue for issue in filtered if issue_filter_text_matches(issue.get("owner") or issue.get("actionOwner"), owner)]
    if issue_area:
        filtered = [issue for issue in filtered if issue_filter_text_matches(issue.get("issueArea"), issue_area)]
    if issue_category:
        filtered = [issue for issue in filtered if issue_filter_text_matches(issue.get("issueCategory"), issue_category)]
    if phase:
        filtered = [issue for issue in filtered if issue_filter_text_matches(issue.get("phase"), phase)]
    if filters.get("organization_support_required") is True:
        filtered = [issue for issue in filtered if issue.get("organizationSupportRequired") is True]
    if filters.get("organization_support_required") is False:
        filtered = [issue for issue in filtered if issue.get("organizationSupportRequired") is False]
    if filters.get("overdue_only") is True:
        filtered = [issue for issue in filtered if (issue.get("derivedMetrics", {}).get("overdueDays") or 0) > 0]
    return filtered


def issue_priority_signal(value: Any) -> int:
    text = cell_to_text(value).lower()
    if not text:
        return 0
    if parse_issue_number(text) is not None:
        return int(parse_issue_number(text) or 0)
    if any(token in text for token in ["critical", "urgent", "high", "상", "높", "긴급", "심각"]):
        return 3
    if any(token in text for token in ["medium", "중"]):
        return 2
    if any(token in text for token in ["low", "하", "낮"]):
        return 1
    return 0


def issue_target_urgency(issue: dict[str, Any], as_of_date: str) -> int:
    target = parse_iso_date(issue.get("targetDate") or issue.get("dueDate"))
    as_of = parse_iso_date(as_of_date)
    if not target or not as_of:
        return 0
    days_until = (target.date() - as_of.date()).days
    return max(0, 365 - days_until)


def issue_risk_sort_key(issue: dict[str, Any], as_of_date: str) -> tuple[int, int, int, int, int, int, int]:
    metrics = issue.get("derivedMetrics", {}) if isinstance(issue.get("derivedMetrics"), dict) else {}
    overdue_days = metrics.get("overdueDays") if isinstance(metrics.get("overdueDays"), int) else 0
    issue_age_days = metrics.get("issueAgeDays") if isinstance(metrics.get("issueAgeDays"), int) else metrics.get("openAgeDays")
    status_weight = {"on_hold": 3, "in_progress": 2, "open": 1}.get(cell_to_text(issue.get("status")), 0)
    signal = max(issue_priority_signal(issue.get("severity")), issue_priority_signal(issue.get("priority")), issue_priority_signal(issue.get("riskLevel")))
    return (
        1 if overdue_days > 0 else 0,
        overdue_days,
        int(issue_age_days or 0),
        1 if issue.get("organizationSupportRequired") is True else 0,
        issue_target_urgency(issue, as_of_date),
        status_weight,
        signal,
    )


def issue_open_sort_key(issue: dict[str, Any]) -> tuple[int, int, int, int, str, int]:
    metrics = issue.get("derivedMetrics", {}) if isinstance(issue.get("derivedMetrics"), dict) else {}
    overdue_days = metrics.get("overdueDays") if isinstance(metrics.get("overdueDays"), int) else 0
    open_age_days = metrics.get("openAgeDays") if isinstance(metrics.get("openAgeDays"), int) else metrics.get("issueAgeDays")
    missing_due = 1 if not issue.get("dueDate") and not is_completed_issue(issue) else 0
    latest_note_date = cell_to_text(issue.get("latestNoteDate"))
    return (
        1 if is_completed_issue(issue) else 0,
        -overdue_days,
        -missing_due,
        -int(open_age_days or 0),
        latest_note_date,
        int(issue.get("sourceRow") or 0),
    )


def issue_support_sort_key(issue: dict[str, Any]) -> tuple[int, int, int]:
    metrics = issue.get("derivedMetrics", {}) if isinstance(issue.get("derivedMetrics"), dict) else {}
    overdue_days = metrics.get("overdueDays") if isinstance(metrics.get("overdueDays"), int) else 0
    issue_age_days = metrics.get("issueAgeDays") if isinstance(metrics.get("issueAgeDays"), int) else metrics.get("openAgeDays")
    status_weight = {"on_hold": 3, "in_progress": 2, "open": 1, "resolved": 0, "closed": 0}.get(cell_to_text(issue.get("status")), 0)
    return (
        overdue_days,
        int(issue_age_days or 0),
        status_weight,
    )


def issue_query_plan(scope: str, detected_intent: str, additional_filters: list[str] | None = None, sort_order: list[str] | None = None) -> dict[str, Any]:
    return {
        "temporalScope": ISSUE_QUERY_STATUS_TO_TEMPORAL.get(scope, scope),
        "detectedIntent": detected_intent,
        "statusFilter": issue_status_scope_label(scope),
        "additionalFilters": additional_filters or [],
        "sortOrder": sort_order or [],
    }


def format_issue_intent_label(intent: str) -> str:
    labels = {
        "issue_list": "이슈 목록",
        "issue_lookup": "키워드 검색",
        "open_issues": "미해결 이슈",
        "unresolved_issues": "미해결 이슈",
        "owner_bottleneck": "담당자 병목",
        "overdue_issues": "지연 이슈",
        "organization_support": "조직지원 필요 이슈",
        "risk_candidates": "위험 이슈",
        "resolved_cases": "해결/완료 이슈",
        "issue_status": "이슈 상태",
    }
    return labels.get(intent, intent)


def format_issue_sort_label(sort_key: str) -> str:
    labels = {
        "overdue_desc": "Overdue",
        "missing_due_desc": "Missing due",
        "issue_age_desc": "Age",
        "open_count_desc": "Open Count",
        "latest_note_date_asc_null_first": "Latest note date asc/null first",
        "organization_support_desc": "조직지원",
        "target_date_urgency_asc": "Target date urgency",
        "status_desc": "Status",
        "severity_priority_desc": "Severity/Priority signal",
        "source_row_asc": "Source row",
    }
    return labels.get(sort_key, sort_key)


def issue_query(input_data: dict[str, Any]) -> dict[str, Any]:
    issue_book = input_data.get("issue_book")
    document = input_data.get("document")
    query = input_data.get("query")
    as_of_date = input_data.get("as_of_date") or datetime.now().date().isoformat()
    if not isinstance(issue_book, dict) or not isinstance(document, dict) or not isinstance(query, str):
        raise WorkerError("invalid_input", "issue_book, document, and query are required.")
    issues = [issue for issue in issue_book.get("issues", []) if isinstance(issue, dict)]
    for issue in issues:
        build_issue_metrics(issue, as_of_date)
    query_text = query.lower()
    structured_query = validate_issue_query(
        input_data.get("issue_query"),
        query,
    )
    filtered_issues = apply_issue_query_filters(issues, structured_query)
    intent = cell_to_text(structured_query.get("intent")) or "issue_list"
    status_scope = cell_to_text(structured_query.get("status_scope")) or "all"
    kind = {
        "issue_list": "summary",
        "issue_lookup": "issue_lookup",
        "open_issues": "open_issues",
        "unresolved_issues": "open_issues",
        "owner_bottleneck": "owner_bottlenecks",
        "overdue_issues": "overdue_issues",
        "organization_support": "organization_support",
        "risk_candidates": "risk_issues",
        "resolved_cases": "completed_issues",
        "issue_status": "open_issues" if status_scope == "unresolved" else "summary",
    }.get(intent, "summary")
    target = None
    selected = filtered_issues[: structured_query["limit"]]
    assignees = None
    clusters = None
    additional_filters = []
    filters = structured_query.get("filters") if isinstance(structured_query.get("filters"), dict) else {}
    for key in sorted(ISSUE_QUERY_FILTER_FIELDS):
        value = filters.get(key)
        if value is not None and value is not False:
            additional_filters.append(f"{key}={value}")
    sort_labels = [format_issue_sort_label(value) for value in structured_query.get("sort", [])]
    detected_intent = format_issue_intent_label(intent)
    plan = issue_query_plan(status_scope, detected_intent, additional_filters, sort_labels)
    if intent in {"open_issues", "unresolved_issues", "owner_bottleneck"}:
        plan["statusFilter"] = "진행중/미해결 + 완료 이력" if status_scope == "all" else "진행중/미해결"
    plan["structuredQuery"] = structured_query
    plan["filteredIssueCount"] = len(filtered_issues)

    if intent == "owner_bottleneck":
        assignees = issue_owner_summaries(filtered_issues)
        selected = sorted(
            filtered_issues,
            key=issue_open_sort_key,
        )[: structured_query["limit"]]
    elif intent in {"open_issues", "unresolved_issues"}:
        selected = sorted(
            filtered_issues,
            key=issue_open_sort_key,
        )[: structured_query["limit"]]
    elif intent == "risk_candidates":
        selected = sorted(
            filtered_issues,
            key=lambda issue: issue_risk_sort_key(issue, as_of_date),
            reverse=True,
        )[: structured_query["limit"]]
    elif intent == "organization_support":
        selected = sorted(
            filtered_issues,
            key=issue_support_sort_key,
            reverse=True,
        )[: structured_query["limit"]]
    elif intent == "overdue_issues":
        selected = sorted(
            filtered_issues,
            key=lambda issue: (issue.get("derivedMetrics", {}).get("overdueDays") or 0, issue.get("derivedMetrics", {}).get("issueAgeDays") or 0),
            reverse=True,
        )[: structured_query["limit"]]
    elif intent == "issue_lookup":
        matches = [issue for issue in filtered_issues if issue_matches_query(issue, query)]
        if matches:
            selected = matches[: structured_query["limit"]]
            target = query
    elif issue_contains_any(query_text, ["repeat", "similar", "cluster", "반복", "유사"]):
        kind = "repeated_issues"
        clusters = issue_keyword_clusters(filtered_issues)
        selected = []
    summary = issue_summary(issue_book, document, as_of_date)
    return {
        "kind": kind,
        "target": target,
        "workspaceId": issue_book["workspaceId"],
        "filename": issue_book["filename"],
        "lastAnalyzedAt": summary["lastAnalyzedAt"],
        "asOfDate": as_of_date,
        "summary": summary,
        "issues": selected[:20],
        "assignees": assignees,
        "clusters": clusters,
        "queryPlan": plan,
        "structuredQuery": structured_query,
        "contextText": issue_context_text(kind, summary, selected, assignees, clusters, plan),
    }


COMMANDS = {
    "runtime-check": runtime_check,
    "embedding-check": check_embedding_status,
    "rag-import": import_document,
    "rag-list": list_documents,
    "rag-delete": delete_document,
    "rag-replace": replace_document,
    "rag-index": index_document,
    "rag-search": search_documents,
    "schedule-register": schedule_register,
    "schedule-parse": parse_schedule,
    "schedule-summary": schedule_summary_command,
    "schedule-query": schedule_query,
    "issue-parse": parse_issue,
    "issue-summary": issue_summary_command,
    "issue-query": issue_query,
    "weekly-report-render": weekly_report_render,
}


def main() -> None:
    parser = argparse.ArgumentParser(description="Mimora local worker")
    parser.add_argument("command", choices=COMMANDS.keys())
    args = parser.parse_args()
    try:
        input_data = read_input()
        write_success(COMMANDS[args.command](input_data))
    except WorkerError as error:
        write_error(error)
    except Exception as error:
        print(f"Unhandled worker error: {error}", file=sys.stderr)
        write_error(WorkerError("worker_failed", "Python worker failed."))


if __name__ == "__main__":
    main()
