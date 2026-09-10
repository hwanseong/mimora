from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

import mimora_worker

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None


def excel_serial(date_text: str) -> int:
    from datetime import date

    year, month, day = [int(part) for part in date_text.split("-")]
    return (date(year, month, day) - date(1899, 12, 30)).days


class ScheduleWorkerTests(unittest.TestCase):
    def setUp(self) -> None:
        if Workbook is None:
            self.skipTest("openpyxl is not installed")
        self.root = tempfile.TemporaryDirectory()
        self.source_path = Path(self.root.name) / "XLGantt_sample.xlsm"
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "Schedule"
        schedule.merge_cells("E3:N4")
        headers = {
            "B3": "task_group",
            "C3": "wbs_level",
            "D3": "WBS",
            "E3": "작업*",
            "O3": "비고",
            "P3": "시작일*",
            "Q3": "완료일*",
            "R3": "Calendar",
            "S3": "총 작업량",
            "T3": "계획 작업량",
            "U3": "총 기간",
            "V3": "계획 기간",
            "W3": "실제 시작일",
            "X3": "실제 완료일",
            "Y3": "실제 총작업량",
            "Z3": "실제 총기간",
            "AB3": "담당",
            "AC3": "산출물",
            "AD3": "계획",
            "AE3": "실적*",
        }
        for cell, value in headers.items():
            schedule[cell] = value

        rows = [
            (5, 1, "1", "E", "프로젝트 관리", "2026-03-05", "2026-08-06", None, None, None, 26.6, 10.2, None, 1, 0.3834586466),
            (6, 2, "1.1", "F", "프로젝트 착수 및 계획수립", "2026-03-05", "2026-03-19", "2026-03-05", "2026-03-19", "이지은", 10, 10, None, 1, 1),
            (7, 3, "1.1.1", "G", "프로젝트,착수", "2026-03-05", "2026-03-05", "2026-03-05", "2026-03-05", "이지은,김분석", 2, 2, None, 1, 1),
            (8, 3, "1.1.2", "G", "프로젝트 범위정의", "2026-03-05", "2026-03-19", "2026-03-05", "2026-03-19", None, 8, 8, None, 1, 1),
            (9, 4, "1.1.2.1", "H", "프로젝트 계획 수립", "2026-03-05", "2026-03-16", "2026-03-05", "2026-03-16", "이지은", 8, 8, None, 1, 1),
            (10, 4, "1.1.2.2", "H", "착수보고(Kick off)", "2026-03-19", "2026-03-19", "2026-03-19", "2026-03-19", "이지은", 0, 0, None, 1, 1),
            (22, 1, "2", "E", "구축", "2026-03-12", "2026-08-09", "2026-03-12", None, None, 227.1, 56.27, None, 1, 0.2169528842),
            (35, 2, "2.3", "F", "개발", "2026-05-01", "2026-07-10", "2026-05-18", None, None, 157, 7.7, None, 1, 0.049044586),
            (36, 3, "2.3.1", "G", "프로그램 개발", "2026-05-01", "2026-06-26", "2026-05-18", None, None, 141, 7.7, None, 1, 0.054609929),
            (37, 4, "2.3.1.1", "H", "단위테스트시나리오작성", "2026-05-01", "2026-05-05", None, None, None, 4, 0, None, 1, 0),
            (38, 4, "2.3.1.2", "H", "프로그램코딩 및 단위테스트", "2026-05-01", "2026-06-26", "2026-05-18", None, None, 137, 7.7, None, 1, 0.056204379),
            (39, 5, "2.3.1.2.1", "I", "서버 개발", "2026-05-04", "2026-06-22", "2026-05-18", None, None, 32, 7.7, None, 1, 0.240625),
            (40, 6, "2.3.1.2.1.1", "J", "프로그램A", "2026-05-04", "2026-05-15", None, None, None, 8, 0, None, 1, 0),
            (41, 6, "2.3.1.2.1.2", "J", "프로그램B", "2026-05-18", "2026-06-08", "2026-05-18", None, "박설계[120%],이지은[50%]", 14, 7.7, None, 1, 0.55),
            (42, 6, "2.3.1.2.1.3", "J", "프로그램C", "2026-06-09", "2026-06-22", None, None, None, 10, 0, None, 1, 0),
            (55, 4, "2.4.2.3", "H", "인수테스트실시", "2026-08-03", "2026-08-06", None, None, None, 2, 0, None, 1, 0),
            (56, 4, "2.4.2.4", "H", "시스템오픈", "2026-08-09", "2026-08-09", None, None, None, 0, 0, None, 1, 0),
        ]
        for (
            row,
            level,
            wbs,
            name_column,
            name,
            planned_start,
            planned_finish,
            actual_start,
            actual_finish,
            resource,
            planned_workload,
            actual_workload,
            deliverable,
            planned_progress,
            actual_progress,
        ) in rows:
            schedule[f"C{row}"] = level
            if level <= 3:
                schedule[f"B{row}"] = "G"
            schedule[f"D{row}"] = wbs
            schedule[f"{name_column}{row}"] = name
            schedule[f"P{row}"] = excel_serial(planned_start)
            schedule[f"Q{row}"] = excel_serial(planned_finish)
            schedule[f"W{row}"] = excel_serial(actual_start) if actual_start else None
            schedule[f"X{row}"] = excel_serial(actual_finish) if actual_finish else None
            schedule[f"T{row}"] = planned_workload
            schedule[f"Y{row}"] = actual_workload
            schedule[f"AB{row}"] = resource
            schedule[f"AC{row}"] = deliverable
            schedule[f"AD{row}"] = planned_progress
            schedule[f"AE{row}"] = actual_progress

        calendar = workbook.create_sheet("Calendar")
        calendar.append(["Date", "Type", "Calendar", "Name"])
        calendar.append([excel_serial("2026-03-07"), "holiday", "Default", "Weekend"])
        progress = workbook.create_sheet("Progress_Data")
        progress.append([
            "날짜",
            "주차",
            "일일 계획작업량(Man-Day)",
            "계획작업량 누적(Man-Day)",
            "계획작업량 진척률",
            "일일 실적기성(Man-Day)",
            "실적기성 누적(Man-Day)",
            "실적기성 진척률",
        ])
        progress.append([excel_serial("2026-08-09"), 32, 3, 253.7, 1, 0, 59.47, 0.2344107213])
        settings = workbook.create_sheet("Settings")
        settings.sheet_state = "veryHidden"
        settings.append(["ITEM_GROUP", "ITEM", "VALUE1"])
        settings.append(["CM", "PROGRAM_NAME", "XLGantt"])
        settings.append(["CM", "VERSION", "6.1.0"])
        workbook.save(self.source_path)

    def tearDown(self) -> None:
        self.root.cleanup()

    def test_parse_xlgantt_schedule_to_canonical_model(self) -> None:
        result = mimora_worker.parse_schedule(
            {
                "workspace_id": "WS-2026-0001",
                "source_path": str(self.source_path),
            }
        )
        schedule = result["schedule"]
        tasks = schedule["tasks"]
        by_wbs = {task["wbs"]: task for task in tasks}

        self.assertGreater(len(tasks), 2)
        self.assertEqual(result["summary"]["taskCount"], 17)
        self.assertIn("1", by_wbs)
        self.assertIn("2", by_wbs)
        self.assertIn("2.3.1.2.1.1", by_wbs)
        self.assertNotEqual(by_wbs["1"]["name"], "G")
        self.assertEqual(by_wbs["2.3.1"]["parentWbs"], "2.3")
        self.assertFalse(by_wbs["2.3.1"]["isLeaf"])
        self.assertTrue(by_wbs["2.3.1.2.1.1"]["isLeaf"])
        self.assertEqual(by_wbs["1"]["plannedStart"], "2026-03-05")
        self.assertEqual(by_wbs["1.1"]["actualFinish"], "2026-03-19")
        self.assertEqual(by_wbs["2.3.1.2.1.2"]["resource"][0]["name"], "박설계")
        self.assertEqual(by_wbs["2.3.1.2.1.2"]["resource"][0]["allocation"], 1.2)
        self.assertEqual(by_wbs["2.3.1.2.1.2"]["resource"][1]["allocation"], 0.5)
        self.assertEqual(schedule["progressSeries"][0]["plannedProgress"], 1)
        self.assertAlmostEqual(schedule["progressSeries"][0]["actualProgress"], 0.2344107213)
        self.assertEqual(result["summary"]["plannedProgress"], 1)
        self.assertAlmostEqual(result["summary"]["actualProgress"], 0.2344107213)
        self.assertEqual(result["summary"]["activeTaskCount"], 1)
        self.assertEqual(result["summary"]["activeWbsNodeCount"], 7)
        self.assertEqual(result["summary"]["delayedTaskCount"], 6)
        self.assertEqual(result["summary"]["delayedWbsNodeCount"], 12)

    def test_schedule_query_returns_context(self) -> None:
        parsed = mimora_worker.parse_schedule(
            {
                "workspace_id": "WS-2026-0001",
                "source_path": str(self.source_path),
            }
        )
        result = mimora_worker.schedule_query(
            {
                "source": parsed["source"],
                "schedule": parsed["schedule"],
                "query": "현재 지연 작업은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "delayed_tasks")
        self.assertGreaterEqual(len(result["tasks"]), 1)
        delayed_wbs = {task["wbs"] for task in result["tasks"]}
        self.assertNotIn("1.1.1", delayed_wbs)
        self.assertNotIn("1.1.2.1", delayed_wbs)
        self.assertNotIn("1.1.2.2", delayed_wbs)
        self.assertNotIn("2.1.1", delayed_wbs)
        self.assertNotIn("2.1.2.1", delayed_wbs)
        self.assertNotIn("2.2.1", delayed_wbs)
        self.assertNotIn("2.2.2.1", delayed_wbs)
        self.assertNotIn("2.2.2.2", delayed_wbs)
        self.assertIn("[SCHEDULE CONTEXT]", result["contextText"])

    def test_optional_real_workbook_regression(self) -> None:
        real_path = Path(
            os.environ.get(
                "MIMORA_REAL_XLGANTT",
                r"D:\vault_test\Schedule\WS-2026-0001_Schedule.xlsm",
            )
        )
        if not real_path.exists():
            self.skipTest("real XLGantt workbook is not available")

        before = real_path.stat().st_mtime_ns
        result = mimora_worker.parse_schedule(
            {
                "workspace_id": "WS-2026-0001",
                "source_path": str(real_path),
            }
        )
        after = real_path.stat().st_mtime_ns

        self.assertEqual(before, after)
        self.assertEqual(result["summary"]["taskCount"], 51)
        self.assertEqual(result["summary"]["leafTaskCount"], 33)
        self.assertEqual(result["summary"]["activeTaskCount"], 2)
        self.assertEqual(result["summary"]["activeWbsNodeCount"], 13)
        self.assertEqual(result["summary"]["delayedTaskCount"], 24)
        self.assertEqual(result["summary"]["delayedWbsNodeCount"], 40)
        self.assertEqual(result["summary"]["plannedProgress"], 1)
        self.assertAlmostEqual(result["summary"]["actualProgress"], 0.2344107213)
        by_wbs = {task["wbs"]: task for task in result["schedule"]["tasks"]}
        self.assertIn("1", by_wbs)
        self.assertIn("2", by_wbs)
        self.assertIn("2.3.1.2.1.1", by_wbs)
        self.assertEqual(by_wbs["2.3.1.2.2"]["parentWbs"], "2.3.1.2")

        resource_result = mimora_worker.schedule_query(
            {
                "source": result["source"],
                "schedule": result["schedule"],
                "query": "박설계 담당 작업은?",
                "as_of_date": "2026-09-10",
            }
        )
        self.assertEqual(resource_result["kind"], "resource_lookup")
        self.assertEqual(resource_result["target"], "박설계")
        self.assertEqual([task["wbs"] for task in resource_result["tasks"]], ["2.1.2.2"])
        self.assertEqual(resource_result["tasks"][0]["name"], "데이터모델링")
        self.assertEqual(resource_result["tasks"][0]["resource"][0]["name"], "박설계")
        self.assertEqual(resource_result["tasks"][0]["resource"][0]["allocation"], 1.2)

        program_b_result = mimora_worker.schedule_query(
            {
                "source": result["source"],
                "schedule": result["schedule"],
                "query": "프로그램B 일정은?",
                "as_of_date": "2026-09-10",
            }
        )
        self.assertEqual(program_b_result["kind"], "task_lookup")
        self.assertEqual(len(program_b_result["tasks"]), 1)
        program_b = program_b_result["tasks"][0]
        self.assertEqual(program_b["wbs"], "2.3.1.2.1.2")
        self.assertEqual(program_b["plannedStart"], "2026-05-21")
        self.assertEqual(program_b["plannedFinish"], "2026-06-11")
        self.assertEqual(program_b["actualStart"], "2026-05-21")
        self.assertIsNone(program_b["actualFinish"])
        self.assertEqual(program_b["actualProgress"], 0.55)

        pm_result = mimora_worker.schedule_query(
            {
                "source": result["source"],
                "schedule": result["schedule"],
                "query": "박피엠 작업은 잘 진행되고 있나?",
                "as_of_date": "2026-09-10",
            }
        )
        self.assertEqual(pm_result["kind"], "resource_status")
        self.assertEqual(pm_result["detectedEntityType"], "resource")
        self.assertEqual(pm_result["detectedEntity"], "박피엠")
        self.assertEqual([task["wbs"] for task in pm_result["tasks"]], ["1.2.1", "1.2.2", "1.2.3"])
        self.assertEqual(pm_result["resourceStatusSummary"]["totalTasks"], 3)
        self.assertEqual(pm_result["resourceStatusSummary"]["completedTasks"], 1)
        self.assertEqual(pm_result["resourceStatusSummary"]["activeTasks"], 0)
        self.assertEqual(pm_result["resourceStatusSummary"]["delayedTasks"], 2)
        self.assertEqual(pm_result["resourceStatusSummary"]["notStartedTasks"], 0)


def make_query_task(
    wbs: str,
    name: str,
    resource: list[dict[str, object]] | None = None,
    planned_start: str | None = None,
    planned_finish: str | None = None,
    actual_start: str | None = None,
    actual_finish: str | None = None,
    actual_progress: float | None = None,
) -> dict[str, object]:
    return {
        "taskId": f"WS-2026-0001:{wbs}",
        "wbs": wbs,
        "level": len(wbs.split(".")),
        "name": name,
        "parentWbs": ".".join(wbs.split(".")[:-1]) if "." in wbs else None,
        "isLeaf": True,
        "plannedStart": planned_start,
        "plannedFinish": planned_finish,
        "actualStart": actual_start,
        "actualFinish": actual_finish,
        "plannedWorkload": None,
        "actualWorkload": None,
        "plannedDuration": None,
        "actualDuration": None,
        "plannedProgress": 1,
        "actualProgress": actual_progress,
        "resource": resource or [],
        "deliverable": None,
        "calendar": None,
    }


class ScheduleQueryFilterTests(unittest.TestCase):
    def make_schedule(self) -> dict[str, object]:
        return {
            "workspaceId": "WS-2026-0001",
            "filename": "WS-2026-0001_Schedule.xlsm",
            "projectStart": "2026-03-05",
            "projectFinish": "2026-08-09",
            "tasks": [
                make_query_task(
                    "1.1.1",
                    "프로젝트 착수",
                    [{"name": "이지은", "allocation": 1.0}],
                    "2026-03-05",
                    "2026-03-05",
                    "2026-03-05",
                    "2026-03-05",
                    1,
                ),
                make_query_task(
                    "2.1.2.2",
                    "데이터모델링",
                    [
                        {"name": "박설계", "allocation": 1.2},
                        {"name": "이지은", "allocation": 0.5},
                    ],
                    "2026-03-26",
                    "2026-03-30",
                    "2026-03-26",
                    None,
                    0.7,
                ),
                make_query_task(
                    "2.2.1",
                    "UI 설계",
                    [{"name": "박피엠", "allocation": 0.5}],
                    "2026-03-31",
                    "2026-04-10",
                    "2026-03-31",
                    "2026-04-10",
                    1,
                ),
                make_query_task(
                    "2.3.2.1",
                    "통합 테스트 준비",
                    [{"name": "박피엠", "allocation": 1.0}],
                    "2026-09-01",
                    "2026-09-30",
                    "2026-09-01",
                    None,
                    0.4,
                ),
                make_query_task(
                    "2.3.2.2",
                    "결함 조치 추적",
                    [{"name": "박피엠", "allocation": 1.0}],
                    "2026-08-01",
                    "2026-08-20",
                    "2026-08-01",
                    None,
                    0.6,
                ),
                make_query_task(
                    "2.4.1",
                    "운영 전환 리허설",
                    [{"name": "박피엠", "allocation": 0.25}],
                    "2026-09-15",
                    "2026-09-18",
                    None,
                    None,
                    0,
                ),
                make_query_task(
                    "2.3.1.2.1.2",
                    "프로그램B",
                    [{"name": "김개발", "allocation": 1.0}],
                    "2026-05-21",
                    "2026-06-11",
                    "2026-05-21",
                    None,
                    0.55,
                ),
            ],
            "progressSeries": [
                {
                    "date": "2026-09-10",
                    "plannedProgress": 1,
                    "actualProgress": 0.2344107213,
                }
            ],
            "parsedSheets": ["Schedule"],
        }

    def make_source(self) -> dict[str, object]:
        return {
            "workspaceId": "WS-2026-0001",
            "sourcePath": r"D:\vault_test\Schedule\WS-2026-0001_Schedule.xlsm",
            "filename": "WS-2026-0001_Schedule.xlsm",
            "fileSize": 1,
            "modifiedAt": "2026-09-10T00:00:00Z",
            "lastParsedAt": "2026-09-10T00:00:00Z",
            "parseStatus": "parsed",
        }

    def test_parse_resource_assignments_with_allocation(self) -> None:
        parsed = mimora_worker.parse_resource_assignments("박설계[120%], 이지은[50%]")

        self.assertEqual(parsed[0]["name"], "박설계")
        self.assertEqual(parsed[0]["allocation"], 1.2)
        self.assertEqual(parsed[1]["name"], "이지은")
        self.assertEqual(parsed[1]["allocation"], 0.5)

    def test_normalize_schedule_header_removes_excel_line_break_marker(self) -> None:
        self.assertEqual(
            mimora_worker.normalize_schedule_header("실제_x000D_\n시작일"),
            "실제시작일",
        )

    def test_resource_lookup_filters_exact_resource_name(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "박설계 담당 작업은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "resource_lookup")
        self.assertEqual(result["target"], "박설계")
        self.assertEqual(len(result["tasks"]), 1)
        self.assertEqual(result["tasks"][0]["wbs"], "2.1.2.2")
        self.assertEqual(result["tasks"][0]["name"], "데이터모델링")
        self.assertEqual(result["tasks"][0]["resource"][0]["allocation"], 1.2)
        self.assertNotIn("1.1.1", {task["wbs"] for task in result["tasks"]})

    def test_task_lookup_program_b_regression(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "프로그램B 일정은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "task_lookup")
        self.assertEqual(len(result["tasks"]), 1)
        task = result["tasks"][0]
        self.assertEqual(task["wbs"], "2.3.1.2.1.2")
        self.assertEqual(task["name"], "프로그램B")
        self.assertEqual(task["plannedStart"], "2026-05-21")
        self.assertEqual(task["plannedFinish"], "2026-06-11")
        self.assertEqual(task["actualStart"], "2026-05-21")
        self.assertIsNone(task["actualFinish"])
        self.assertEqual(task["actualProgress"], 0.55)

    def test_resource_entity_progress_query_routes_to_resource_status(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "박피엠 작업은 잘 진행되고 있나?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "resource_status")
        self.assertEqual(result["detectedEntityType"], "resource")
        self.assertEqual(result["detectedEntity"], "박피엠")
        self.assertEqual(result["detectedIntent"], "resource_status")
        self.assertEqual(
            {task["wbs"] for task in result["tasks"]},
            {"2.2.1", "2.3.2.1", "2.3.2.2", "2.4.1"},
        )
        self.assertNotIn("1.1.1", {task["wbs"] for task in result["tasks"]})
        self.assertEqual(result["resourceStatusSummary"]["totalTasks"], 4)
        self.assertEqual(result["resourceStatusSummary"]["completedTasks"], 1)
        self.assertEqual(result["resourceStatusSummary"]["activeTasks"], 1)
        self.assertEqual(result["resourceStatusSummary"]["delayedTasks"], 1)
        self.assertEqual(result["resourceStatusSummary"]["notStartedTasks"], 1)

    def test_task_entity_progress_query_routes_to_task_status(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "프로그램B는 잘 진행되고 있나?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "task_status")
        self.assertEqual(result["detectedEntityType"], "task")
        self.assertEqual(result["detectedEntity"], "프로그램B")
        self.assertEqual(len(result["tasks"]), 1)
        self.assertEqual(result["tasks"][0]["wbs"], "2.3.1.2.1.2")
        self.assertEqual(result["taskAnalyses"][0]["status"], "delayed")

    def test_general_active_query_keeps_existing_routing_without_entity(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "현재 진행 중인 작업은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "active_tasks")

    def test_general_delayed_query_keeps_existing_routing_without_entity(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "현재 지연 작업은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "delayed_tasks")

    def test_dependency_reference_parser_supports_dependency_types(self) -> None:
        parsed = mimora_worker.parse_dependency_references(
            "1.1FS, 1.2SS+2, 1.3FF-1, 1.4SF",
            "2.1",
        )

        self.assertEqual([item["type"] for item in parsed], ["FS", "SS", "FF", "SF"])
        self.assertEqual(parsed[1]["lag_days"], 2)
        self.assertEqual(parsed[2]["lag_days"], -1)

    def test_dependency_lookup_reports_unavailable_without_source_dependencies(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "프로그램B 선행 작업은?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "dependency_lookup")
        self.assertEqual(result["advancedAnalysis"]["dependencyCount"], 0)
        self.assertIn("dependency_not_available", result["advancedAnalysis"]["warnings"])

    def test_earned_schedule_query_returns_performance(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "현재 일정 성과는?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "schedule_performance")
        self.assertEqual(result["advancedAnalysis"]["actualProgress"], 0.2344107213)
        self.assertIn("schedulePerformanceIndex", result["advancedAnalysis"])
        self.assertIn("earnedScheduleDays", result["advancedAnalysis"])
        self.assertIn("actualTimeDays", result["advancedAnalysis"])

    def test_schedule_performance_query_routes_plan_variance(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "계획보다 얼마나 밀렸어?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "schedule_performance")
        self.assertIn("scheduleVarianceDays", result["advancedAnalysis"])

    def test_forecast_query_returns_forecast_methods(self) -> None:
        schedule = self.make_schedule()
        for task in schedule["tasks"]:
            if task.get("isLeaf") and not mimora_worker.is_completed_task(task):
                task["plannedDuration"] = 10
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": schedule,
                "query": "현재 추세면 프로젝트 언제 끝나?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "forecast")
        self.assertEqual(len(result["advancedAnalysis"]["methods"]), 3)
        self.assertTrue(all("available" in method for method in result["advancedAnalysis"]["methods"]))
        self.assertTrue(all("reason" in method for method in result["advancedAnalysis"]["methods"]))
        self.assertTrue(all("quality" in method for method in result["advancedAnalysis"]["methods"]))
        self.assertTrue(all("assumptions" in method for method in result["advancedAnalysis"]["methods"]))
        self.assertEqual(result["advancedAnalysis"]["primaryMethod"], "remaining_tasks")
        self.assertIsNotNone(result["advancedAnalysis"]["primaryEstimate"])
        self.assertEqual(
            result["advancedAnalysis"]["forecastRange"]["latest"],
            result["advancedAnalysis"]["primaryEstimate"],
        )
        self.assertIn("forecast_partially_unavailable", result["advancedAnalysis"]["warnings"])
        self.assertIn("recent_velocity_unavailable", result["advancedAnalysis"]["warnings"])
        self.assertNotIn("forecast_unavailable", result["advancedAnalysis"]["warnings"])
        self.assertEqual(
            result["advancedAnalysis"]["performanceScenario"]["method"],
            "earned_schedule",
        )

    def test_forecast_query_routes_plan_finish_question(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "계획 종료일을 지킬 수 있어?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "forecast")
        self.assertEqual(len(result["advancedAnalysis"]["methods"]), 3)

    def test_what_if_without_dependencies_does_not_invent_project_finish(self) -> None:
        result = mimora_worker.schedule_query(
            {
                "source": self.make_source(),
                "schedule": self.make_schedule(),
                "query": "프로그램B가 5영업일 늦어지면?",
                "as_of_date": "2026-09-10",
            }
        )

        self.assertEqual(result["kind"], "what_if")
        self.assertEqual(result["advancedAnalysis"]["delayWorkingDays"], 5)
        self.assertIn("what_if_dependency_missing", result["advancedAnalysis"]["warnings"])
        self.assertIsNone(result["advancedAnalysis"]["projectWhatIfFinish"])


class ScheduleRealWorkbookForecastRegressionTests(unittest.TestCase):
    def setUp(self) -> None:
        if Workbook is None:
            self.skipTest("openpyxl is not installed")
        self.source_path = Path(r"D:\vault_test\Schedule\WS-2026-0001_Schedule.xlsm")
        if not self.source_path.exists():
            self.skipTest("WS-2026-0001 sample workbook is not available")

    def test_real_workbook_schedule_performance_and_forecast(self) -> None:
        parsed = mimora_worker.parse_schedule(
            {
                "workspace_id": "WS-2026-0001",
                "source_path": str(self.source_path),
            }
        )
        schedule = parsed["schedule"]
        source = parsed["source"]

        performance = mimora_worker.calculate_schedule_performance(
            schedule,
            source,
            "2026-09-10",
        )
        forecast = mimora_worker.calculate_schedule_forecast(
            schedule,
            source,
            "2026-09-10",
        )

        self.assertEqual(performance["plannedProgress"], 1)
        self.assertAlmostEqual(performance["actualProgress"], 0.2344107213243989)
        self.assertEqual(performance["earnedScheduleDate"], "2026-04-13")
        self.assertEqual(performance["actualTimeDays"], 131)
        self.assertEqual(performance["scheduleVarianceDays"], -103)
        self.assertAlmostEqual(performance["schedulePerformanceIndex"], 0.21, places=2)

        self.assertEqual(forecast["primaryEstimate"], "2026-11-06")
        self.assertEqual(forecast["forecastRange"]["latest"], "2026-11-06")
        self.assertEqual(forecast["performanceScenario"]["estimate"], "2028-03-07")
        self.assertIn("forecast_partially_unavailable", forecast["warnings"])
        self.assertIn("recent_velocity_unavailable", forecast["warnings"])
        self.assertNotIn("forecast_unavailable", forecast["warnings"])
        recent_velocity = next(
            method
            for method in forecast["methods"]
            if method["method"] == "recent_velocity"
        )
        self.assertFalse(recent_velocity["available"])
        self.assertEqual(recent_velocity["quality"], "stale")


if __name__ == "__main__":
    unittest.main()
